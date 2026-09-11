#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""`prepare_materials.py` 契约测试（H0 数据隔离 + 重建法三条纪律）。

自洽：只用本技能 `scripts/` 内的脚本造临时 fixture，可随技能一起安装运行。
每条用例都对应一次**真实失效**，防止回退：

1. 缓存值优先 —— 源格"公式 + 缓存值"时，重建后必须落**缓存值**（否则后续 data_only
   读回 None，把"未重算"读成"数据缺失"，曾误报"评估结果列全空"）。
2. H0 —— 人工隐藏的 sheet / 行 / 列 / 折叠分组整体排除：其内容值**不得出现在任何产物**里。
3. 同名消歧 —— 定稿与送审稿同名时，工作版必须两份都在、来源可区分、不得静默覆盖。
4. 「值不可得」登记 —— 无缓存值的公式格单独计数，不混进"数据缺失"。
5. 源材料目录只读 —— 运行后源目录文件集合不变。
"""
from __future__ import annotations

import importlib.util
import json
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent


def _load_module():
    spec = importlib.util.spec_from_file_location(
        "prepare_materials_for_test", SCRIPTS_DIR / "prepare_materials.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _with_cached_value(path: Path, cell_ref: str, formula: str, cached: str) -> None:
    """把 openpyxl 写出的公式格补上 Excel 才会写的 `<v>` 缓存值。

    直接造"公式 + 缓存值"的源格，才能复现线上那种"源有意义的值、重建后却读回 None"的场景。
    """
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                text = data.decode("utf-8")
                needle = f'<f>{formula}</f>'
                assert needle in text, f"{item.filename}: 未找到 {needle}"
                text = text.replace(needle, f"{needle}<v>{cached}</v>", 1)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


class PrepareMaterialsContractTest(unittest.TestCase):
    def setUp(self):
        import openpyxl  # noqa: F401  （缺失则整体 skip）
        self.module = _load_module()
        self.tmp = tempfile.TemporaryDirectory()
        self.case = Path(self.tmp.name)
        self.src = self.case / "材料-源"
        self.src.mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self):
        return self.module.prepare(
            str(self.case), str(self.src),
            str(self.case / "提取"), str(self.case / "工作版"),
        )

    def _inventory(self):
        return json.loads((self.case / "材料盘点.json").read_text(encoding="utf-8"))

    # ---- 1 缓存值优先 ----------------------------------------------------
    def test_cached_value_wins_over_formula_string(self):
        """源格 = 公式 + 缓存值 → 重建后必须能 data_only 读回该值。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        f = d / "明细表.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "评估结果"
        wb.active["A2"] = "=1+1"
        wb.save(f)
        _with_cached_value(f, "A2", "1+1", "42")

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/明细表.xlsx", data_only=True)
        self.assertEqual(42, rebuilt.active["A2"].value)
        self.assertEqual(
            0,
            self._inventory()[0]["workbook"]["sheets"][0]["valueUnavailable"],
            "有缓存值的格不应被记成「值不可得」",
        )

    def test_formula_without_cached_value_is_registered_as_value_unavailable(self):
        """无缓存值的公式格：保留公式串（不丢结构），并单独登记「值不可得」。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        wb.active["A1"] = "=SUM(B1:B2)"
        wb.save(d / "外部引用.xlsx")

        self._run()

        rec = self._inventory()[0]
        self.assertEqual(1, rec["workbook"]["sheets"][0]["valueUnavailable"])
        self.assertTrue(rec["workbook"]["valueUnavailable"])
        rebuilt = openpyxl.load_workbook(self.case / "工作版/外部引用.xlsx", data_only=False)
        self.assertEqual("=SUM(B1:B2)", rebuilt.active["A1"].value)

    # ---- 2 H0 隐藏区 -----------------------------------------------------
    def test_hidden_content_never_reaches_any_artifact(self):
        """隐藏 sheet / 行 / 列的内容值不得出现在工作版或盘点中（H0）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        f = d / "含隐藏.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "可见表"
        # A1/C1 可见（对照组），A2 隐藏行、B1 隐藏列、隐藏 sheet 各放一个哨兵值
        ws["A1"], ws["A2"], ws["B1"], ws["C1"] = "保留", 999999, 888888, 777777
        ws.row_dimensions[2].hidden = True           # 隐藏行
        ws.column_dimensions["B"].hidden = True      # 隐藏列
        hid = wb.create_sheet("隐藏表")
        hid["A1"] = 666666                           # 隐藏 sheet 的内容值
        hid.sheet_state = "hidden"                   # 必须显式设为隐藏（create_sheet 默认可见）
        wb.save(f)

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/含隐藏.xlsx", data_only=True)
        values = [c.value for row in rebuilt["可见表"].iter_rows() for c in row if c.value is not None]
        self.assertIn("保留", values)
        self.assertIn(777777, values, "可见列的值必须保留（对照组）")
        for sentinel in (999999, 888888, 666666):
            self.assertNotIn(sentinel, values, f"隐藏值 {sentinel} 泄漏到工作版")
        meta = self._inventory()[0]["workbook"]["hiddenMeta"]
        self.assertEqual(["隐藏表"], meta["hiddenSheets"], "隐藏 sheet 的**名字**必须登记（忽略清单）")
        self.assertEqual([2], meta["hiddenRows"]["可见表"], "隐藏行段位必须登记")
        self.assertEqual(["B"], meta["hiddenCols"]["可见表"], "隐藏列段位必须登记")
        blob = (self.case / "材料盘点.json").read_text(encoding="utf-8")
        for sentinel in ("999999", "888888", "666666"):
            self.assertNotIn(sentinel, blob, f"隐藏区内容值 {sentinel} 不得进入盘点")

    def test_collapsed_outline_group_counts_as_hidden(self):
        """折叠分组内的行视同隐藏（Excel 折叠后不可见）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        f = d / "分组.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["A2"], ws["A3"], ws["A4"] = "头", 12345, 23456, "尾"
        ws.row_dimensions[2].outlineLevel = 1        # 组内行
        ws.row_dimensions[3].outlineLevel = 1        # 组内行
        ws.row_dimensions[4].collapsed = True        # 汇总行（默认在下方）= 组已折叠
        wb.save(f)

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/分组.xlsx", data_only=True)
        values = [c.value for row in rebuilt.active.iter_rows() for c in row if c.value is not None]
        self.assertEqual(["头", "尾"], values, "折叠组内的行必须排除")

    # ---- 3 同名消歧 ------------------------------------------------------
    def test_same_named_versions_do_not_overwrite_each_other(self):
        """定稿与送审稿同名 → 工作版两份都在，来源可区分。"""
        import openpyxl
        for stage, val in (("定稿", "出租"), ("送审稿", "自用")):
            d = self.src / stage
            d.mkdir()
            wb = openpyxl.Workbook()
            wb.active["A1"] = val
            wb.save(d / "13-评估明细表-20231231.xlsx")

        self._run()

        work = self.case / "工作版"
        names = sorted(p.name for p in work.glob("*.xlsx"))
        self.assertEqual(2, len(names), names)
        got = {
            openpyxl.load_workbook(work / n, data_only=True).active["A1"].value
            for n in names
        }
        self.assertEqual({"出租", "自用"}, got, "两份同名材料不得互相覆盖")
        records = self._inventory()
        self.assertEqual({"定稿", "送审稿"}, {r["stage"] for r in records})
        self.assertTrue(any(r.get("nameCollision") for r in records), "同名冲突须在盘点中标出")

    # ---- 4 源目录只读 ----------------------------------------------------
    def test_source_materials_directory_is_not_polluted(self):
        """运行后源材料目录的文件集合不得变化（不得落派生文件）。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        wb.save(d / "a.xlsx")
        before = sorted(str(p.relative_to(self.src)) for p in self.src.rglob("*"))

        self._run()

        after = sorted(str(p.relative_to(self.src)) for p in self.src.rglob("*"))
        self.assertEqual(before, after)

    # ---- 4b 缺依赖不得静默空返 -------------------------------------------
    def test_docx_without_python_docx_never_silently_returns_empty(self):
        """缺 python-docx 时必须回退（或明确报因），不得静默产出空文本。"""
        import sys
        import zipfile as zf
        d = self.src / "定稿"
        d.mkdir()
        # 手工造最小 .docx（zip + 两个必需部件）
        f = d / "说明.docx"
        with zf.ZipFile(f, "w") as z:
            z.writestr("[Content_Types].xml",
                       '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
                       'package/2006/content-types"><Default Extension="xml" ContentType='
                       '"application/xml"/></Types>')
            z.writestr("word/document.xml",
                       '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
                       'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>评估说明正文</w:t>'
                       '</w:r></w:p></w:body></w:document>')

        saved = sys.modules.get("docx")
        sys.modules["docx"] = None          # 模拟未安装 python-docx
        try:
            self._run()
        finally:
            if saved is None:
                sys.modules.pop("docx", None)
            else:
                sys.modules["docx"] = saved

        rec = self._inventory()[0]
        if rec["readable"]:
            text = (self.case / "提取/说明.docx.txt").read_text(encoding="utf-8")
            self.assertTrue(text.strip(), "回退成功时不得产出空文本")
            self.assertIn("回退", rec.get("note", ""))
        else:
            self.assertTrue(rec.get("note"), "读不到必须写明原因，不得静默")

    # ---- 5 格式语义 ------------------------------------------------------
    def test_number_format_is_preserved(self):
        """重建须保留 number_format，避免数值/日期语义被改变。"""
        import openpyxl
        d = self.src / "定稿"
        d.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 0.075
        ws["A1"].number_format = "0.00%"
        wb.save(d / "格式.xlsx")

        self._run()

        rebuilt = openpyxl.load_workbook(self.case / "工作版/格式.xlsx")
        self.assertEqual("0.00%", rebuilt.active["A1"].number_format)


if __name__ == "__main__":
    unittest.main(verbosity=2)

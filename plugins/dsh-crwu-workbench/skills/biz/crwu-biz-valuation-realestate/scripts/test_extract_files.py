#!/usr/bin/env python3
"""`extract_files.py` 的契约测试（报告评估多类型资料提取）。

钉住的是**降级语义**，不是"能不能解析某个格式"：

1. 输出必须是固定形状的合法 JSON（上层按字段读，缺字段会一路变成 `undefined`）；
2. 缺库 / 旧版二进制 Office / 无文本层 PDF / 压缩包 —— 一律进 `capabilityGaps`，
   并且**绝不产出 `extractedFields`**（"读不出来"与"读到了空"必须有区别）；
3. `extractedFields` 只做**标签锚定**：没有标签的行不产生字段，值保持原文；
4. 名字与扩展名不认识时也说得出话（不是抛异常）；
5. 单份文件读不出来不影响脚本进程本身（退出码 0 + 结果文件存在）；
6. **P1：扫描 PDF 按页报告** —— 逐页 `hasTextLayer` / `chars`，无文本层的页码进**一条**
   可执行的缺口（`第 2–3 页无文本层：…，请人工核对这几页的坐落/面积/权属`），
   并且本轮的 PDF 渲染 / OCR / 视觉能力是**当场探到的**（`capabilities`），不是写死的。

图像 / Word / Excel 的具体解析依赖 openpyxl / python-docx / Pillow；缺失时**跳过**那几条
而不是失败 —— 那是运行时的能力差异，不是脚本的缺陷。PDF 的逐页事实同理依赖
pypdf / PyPDF2 / pdfminer：三者都没有时**跳过逐页用例**（缺库分支由另一条用例覆盖）。
"""

import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zlib

# 导入被测模块只为读纯函数（分类 / 文案），不要顺势在技能目录里留下 `__pycache__`。
sys.dont_write_bytecode = True

SCRIPT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "extract_files.py")


def load_script_module():
    """把脚本当模块载入（只为调用纯函数；脚本的 CLI 契约由 subprocess 用例覆盖）。"""
    spec = importlib.util.spec_from_file_location("crwu_extract_files_under_test", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _pdf_stream(content: bytes) -> bytes:
    return b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"


# 一个"页面内容就是一张内嵌位图"的页标记（扫描件的典型形态）。
IMAGE_PAGE = "@image"

# pypdf **接口**替身：本机没有任何解释器同时装齐 pypdf 与 Pillow（缺 Pillow 时
# `page.images` 直接抛错），导出通道的成功分支只能用替身跑通。它只是库接口，**不是渲染器**。
FAKE_PYPDF = '''\
class _Image:
    def __init__(self, name, data):
        self.name, self.data = name, data


class _Page:
    def __init__(self, text, images):
        self._text, self._images = text, images

    def extract_text(self):
        return self._text

    @property
    def images(self):
        return self._images


class PdfReader:
    """两页：第 1 页有文本层，第 2 页只有一张内嵌位图。"""

    def __init__(self, path):
        self.pages = [
            _Page("ZuoLuo: No.1 Example Road", []),
            _Page("", [_Image("page2.png", b"\\x89PNG-fake-embedded-bytes")]),
        ]
'''


def build_pdf(page_specs) -> bytes:
    """手写一份最小 PDF：每页要么一个文本对象，要么空内容流，要么一张内嵌位图。

    只用标准库拼字节 —— 这样"扫描件"这个测试前提不依赖任何 PDF 库；能不能**读**出来
    才是被 pypdf / PyPDF2 / pdfminer 决定的那一半。文本只能用 ASCII（简单字体 / WinAnsi），
    中文正文需要 CID 字体，不是这个夹具要覆盖的东西。
    """
    objects: dict = {}

    def add(body: bytes) -> int:
        number = len(objects) + 1
        objects[number] = body
        return number

    catalog_num = add(b"")
    pages_num = add(b"")
    font_num = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    image_num = 0
    if any(spec == IMAGE_PAGE for spec in page_specs):
        payload = zlib.compress(bytes([220, 20, 20]) * 4)  # 2×2 DeviceRGB 8bpc
        image_num = add(
            b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceRGB "
            b"/BitsPerComponent 8 /Filter /FlateDecode /Length %d >>\nstream\n" % len(payload)
            + payload + b"\nendstream"
        )
    kids = []
    for spec in page_specs:
        if spec == IMAGE_PAGE:
            content = b"q 200 0 0 200 100 400 cm /Im1 Do Q"
            xobject = b"/XObject << /Im1 %d 0 R >> " % image_num
        elif spec is None:
            content, xobject = b"", b""
        else:
            content = b"BT /F1 12 Tf 72 720 Td (" + str(spec).encode("latin-1") + b") Tj ET"
            xobject = b""
        content_num = add(_pdf_stream(content))
        page_num = add(
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 %d 0 R >> %s>> /Contents %d 0 R >>"
            % (pages_num, font_num, xobject, content_num)
        )
        kids.append(b"%d 0 R" % page_num)
    objects[pages_num] = b"<< /Type /Pages /Kids [" + b" ".join(kids) + b"] /Count %d >>" % len(kids)
    objects[catalog_num] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_num

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1) + b"0000000000 65535 f \n"
    for number in sorted(objects):
        out += b"%010d 00000 n \n" % offsets[number]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def run_script(source_path: str, workdir: str, env: dict = None, python_args=()) -> dict:
    """跑一次脚本，返回解析后的 JSON（结果文件必须存在）。"""
    out_path = os.path.join(workdir, "out.json")
    environment = dict(os.environ)
    environment.update(env or {})
    completed = subprocess.run(
        [sys.executable, *python_args, SCRIPT, source_path, out_path],
        capture_output=True, text=True, timeout=120, check=False, env=environment,
    )
    if not os.path.isfile(out_path):
        raise AssertionError(f"脚本没有写出结果文件：stdout={completed.stdout!r} stderr={completed.stderr!r}")
    with open(out_path, "r", encoding="utf-8") as handle:
        return json.load(handle)


SHAPE = ("ok", "error", "fileName", "fileType", "text", "tables", "images",
         "extractedFields", "warnings", "capabilityGaps",
         # P1 新增：逐页事实、页码合计、能力声明与探针证据（只增不改旧键）。
         "pages", "pageCount", "pagesWithText", "pagesWithoutText",
         "capabilities", "capabilityProbe")

PAGE_KEYS = ("index", "hasTextLayer", "chars")
CAPABILITY_KEYS = ("pdfText", "pdfRender", "imageOcr", "visionFallback")
VISION_FALLBACKS = ("none", "session-images")


class ShapeAssertions:
    """结果形状断言（两个测试类共用；**不继承** TestCase，免得被当成一个测试类）。"""

    def assertShape(self, payload: dict) -> None:
        for key in SHAPE:
            self.assertIn(key, payload, f"结果缺字段 {key}")
        self.assertIsInstance(payload["ok"], bool)
        self.assertIsInstance(payload["tables"], list)
        self.assertIsInstance(payload["images"], list)
        self.assertIsInstance(payload["extractedFields"], dict)
        self.assertIsInstance(payload["warnings"], list)
        self.assertIsInstance(payload["capabilityGaps"], list)
        for key, value in payload["extractedFields"].items():
            self.assertIsInstance(value, str, f"字段 {key} 的值必须是字符串")
        # 逐页事实与能力声明的形状：每条结果都要有，且值域受限（上层按它们渲染）。
        self.assertIsInstance(payload["pages"], list, "pages 必须是列表")
        self.assertIsInstance(payload["pageCount"], int)
        self.assertIsInstance(payload["pagesWithText"], int)
        self.assertIsInstance(payload["pagesWithoutText"], int)
        self.assertEqual(payload["pagesWithText"] + payload["pagesWithoutText"], len(payload["pages"]))
        for page in payload["pages"]:
            for key in PAGE_KEYS:
                self.assertIn(key, page, f"逐页事实缺字段 {key}")
            self.assertIsInstance(page["index"], int)
            self.assertIsInstance(page["hasTextLayer"], bool)
            self.assertIsInstance(page["chars"], int)
            if page["hasTextLayer"]:
                self.assertGreater(page["chars"], 0, "有文本层就该有字数")
            else:
                self.assertEqual(page["chars"], 0, "无文本层的页不许报字数")
        capabilities = payload["capabilities"]
        self.assertIsInstance(capabilities, dict)
        for key in CAPABILITY_KEYS:
            self.assertIn(key, capabilities, f"能力声明缺字段 {key}")
        self.assertIsInstance(capabilities["pdfText"], bool)
        self.assertIsInstance(capabilities["pdfRender"], bool)
        self.assertIsInstance(capabilities["imageOcr"], bool)
        self.assertIn(capabilities["visionFallback"], VISION_FALLBACKS)
        self.assertIsInstance(payload["capabilityProbe"], dict)


class ExtractContractTest(ShapeAssertions, unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="crwu-extract-")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write(self, name: str, content: str) -> str:
        path = os.path.join(self.dir, name)
        with open(path, "w", encoding="utf-8") as handle:
            handle.write(content)
        return path

    # ── 文本 ────────────────────────────────────────────────────────────────

    def test_text_file_yields_original_text_and_label_anchored_fields(self):
        path = self.write("a.txt", "坐落：某市某区某路 1 号\n用途：住宅\n建筑面积：120.5 ㎡\n")
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertTrue(payload["ok"], payload)
        self.assertEqual(payload["fileType"], "text")
        self.assertIn("某路 1 号", payload["text"])
        self.assertEqual(payload["extractedFields"]["坐落"], "某市某区某路 1 号")
        self.assertEqual(payload["extractedFields"]["用途"], "住宅")
        self.assertEqual(payload["extractedFields"]["建筑面积"], "120.5 ㎡")

    def test_fields_are_only_label_anchored_never_inferred(self):
        # 没有标签的正文一个字段都不许出现（"1 号"不是坐落，"120 平"不是建筑面积）。
        path = self.write("b.txt", "这是一段没有标签的描述：某路 1 号，大约 120 平方米，1998 年建成。\n")
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertEqual(payload["extractedFields"], {}, "不许从无标签正文里推断字段")

    def test_gbk_text_is_decoded_and_reported(self):
        path = os.path.join(self.dir, "c.txt")
        with open(path, "wb") as handle:
            handle.write("坐落：某市老城区 3 号\n".encode("gb18030"))
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertTrue(payload["ok"], payload)
        self.assertEqual(payload["extractedFields"]["坐落"], "某市老城区 3 号")
        self.assertTrue(any("解码" in item for item in payload["warnings"]), payload["warnings"])

    # ── 只报缺口的类型 ──────────────────────────────────────────────────────

    def test_empty_file_is_an_error_not_a_gap(self):
        path = self.write("empty.txt", "")
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertFalse(payload["ok"])
        self.assertIn("空文件", payload["error"])
        self.assertEqual(payload["extractedFields"], {})

    def test_legacy_binary_office_reports_a_gap_without_fields(self):
        for name, hint in (("老报告.doc", "antiword"), ("老表.xls", "xlrd")):
            path = self.write(name, "not really a binary office file")
            payload = run_script(path, self.dir)
            self.assertShape(payload)
            self.assertFalse(payload["ok"], name)
            self.assertEqual(payload["extractedFields"], {}, name)
            self.assertTrue(any(hint in gap for gap in payload["capabilityGaps"]), payload["capabilityGaps"])

    def test_archive_is_not_expanded_in_this_version(self):
        path = self.write("资料.zip", "PK\x03\x04 not really a zip")
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertFalse(payload["ok"])
        self.assertEqual(payload["extractedFields"], {})
        self.assertTrue(any("压缩包" in gap for gap in payload["capabilityGaps"]), payload["capabilityGaps"])

    def test_unknown_extension_says_so_instead_of_raising(self):
        path = self.write("调研纪要.wps", "内容")
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertFalse(payload["ok"])
        self.assertTrue(any("不认识的文件类型" in gap for gap in payload["capabilityGaps"]), payload["capabilityGaps"])

    def test_image_is_registered_but_content_is_a_gap(self):
        try:
            from PIL import Image  # type: ignore
        except ImportError:
            self.skipTest("本机没有 Pillow：图像条目跳过（脚本的缺口分支已由其它用例覆盖）")
        path = os.path.join(self.dir, "现场.png")
        Image.new("RGB", (32, 16), "white").save(path)
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertTrue(payload["ok"], payload)
        self.assertEqual(len(payload["images"]), 1)
        self.assertIn("32×16", payload["images"][0]["description"])
        self.assertEqual(payload["extractedFields"], {}, "图片不许编造字段")
        self.assertTrue(any("OCR" in gap for gap in payload["capabilityGaps"]), payload["capabilityGaps"])

    def test_pdf_without_a_reader_reports_the_missing_library(self):
        # 无论运行时有没有 PDF 库，都不许出现"编出来的字段"：
        # 有库时可能读到空正文（→ 无文本层缺口），没库时直接报缺库。
        path = os.path.join(self.dir, "报告.pdf")
        with open(path, "wb") as handle:
            handle.write(b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n")
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertEqual(payload["extractedFields"], {})
        self.assertFalse(payload["ok"])
        self.assertTrue(payload["capabilityGaps"], "必须给出至少一条能力缺口")

    def test_missing_input_file_is_reported(self):
        payload = run_script(os.path.join(self.dir, "不存在.txt"), self.dir)
        self.assertShape(payload)
        self.assertFalse(payload["ok"])
        self.assertIn("不存在", payload["error"])


class ExtractOfficeTest(unittest.TestCase):
    """Word / Excel 的真实解析（缺库时跳过）。"""

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="crwu-extract-office-")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_docx_paragraphs_tables_and_fields(self):
        try:
            import docx  # type: ignore
        except ImportError:
            self.skipTest("本机没有 python-docx")
        path = os.path.join(self.dir, "说明.docx")
        document = docx.Document()
        document.add_paragraph("房屋性质：商品房")
        document.add_paragraph("建筑结构：钢筋混凝土")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "总层数"
        table.cell(0, 1).text = "18"
        document.save(path)

        payload = run_script(path, self.dir)
        self.assertTrue(payload["ok"], payload)
        self.assertEqual(payload["fileType"], "word")
        self.assertEqual(payload["extractedFields"]["房屋性质"], "商品房")
        self.assertEqual(payload["extractedFields"]["建筑结构"], "钢筋混凝土")
        self.assertEqual(payload["extractedFields"]["总楼层"], "18", "表格里的标签锚定也要生效")
        self.assertEqual(len(payload["tables"]), 1)

    def test_xlsx_sheets_become_tables_and_fields(self):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        path = os.path.join(self.dir, "计算表.xlsx")
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "基础信息表"
        sheet.append(["坐落", "某市某区某路 2 号"])
        sheet.append(["用途", "商业"])
        sheet.append(["建筑面积", 88.0])
        workbook.save(path)

        payload = run_script(path, self.dir)
        self.assertTrue(payload["ok"], payload)
        self.assertEqual(payload["fileType"], "excel")
        self.assertEqual(payload["extractedFields"]["坐落"], "某市某区某路 2 号")
        self.assertEqual(payload["extractedFields"]["用途"], "商业")
        self.assertEqual(payload["extractedFields"]["建筑面积"], "88")
        self.assertEqual(payload["tables"][0]["name"], "基础信息表")
        self.assertIn("基础信息表", payload["tables"][0]["name"])


def independent_pdf_library():
    """独立于脚本的一次探测（不能拿脚本自己的答案当"证据"）。"""
    for name in ("pypdf", "PyPDF2", "pdfminer.high_level"):
        try:
            if importlib.util.find_spec(name) is not None:
                return name
        except (ImportError, ValueError):
            continue
    return ""


class ExtractPdfPageFactsTest(ShapeAssertions, unittest.TestCase):
    """P1：**逐页**文本层事实、能力声明与页码缺口文案。

    这一组用例钉住的是"这份报告到底读到了哪几页、哪几页读不到、为什么读不到"。
    关键分界：`hasTextLayer=false` 的页**留空**（不产生字段），并由**一条带页码区间**的
    缺口交给人工 —— 而不是一句笼统的"PDF 无文本层，需要 OCR"。
    """

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="crwu-extract-pdf-")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write_pdf(self, name: str, page_specs) -> str:
        path = os.path.join(self.dir, name)
        with open(path, "wb") as handle:
            handle.write(build_pdf(page_specs))
        return path

    # ── 能力探针 ────────────────────────────────────────────────────────────

    def test_capability_probe_is_measured_not_asserted(self):
        """`capabilities` 必须等于**当场**探到的运行时事实（不许写死一个说法）。"""
        module = load_script_module()
        facts, evidence = module.probe_capabilities()

        self.assertEqual(set(CAPABILITY_KEYS) - {"visionFallback"}, set(facts) - {"visionFallback"})
        self.assertEqual(facts["pdfText"], independent_pdf_library() != "")
        self.assertEqual(facts["pdfText"], evidence["pdfTextLibrary"] != "")
        self.assertIn(facts["visionFallback"], VISION_FALLBACKS)
        self.assertEqual(facts["visionFallback"], "none", "探针本身不认识任何产物：缺省就是 none")

        # Check the installed capabilities without assuming a particular deployment.
        rasterizers = [name for name in ("pdftoppm", "pdftocairo", "mutool", "gs") if shutil.which(name)]
        try:
            rasterizable = importlib.util.find_spec("fitz") is not None
        except (ImportError, ValueError):
            rasterizable = False
        if not rasterizers and not rasterizable:
            self.assertFalse(facts["pdfRender"], "没有渲染器就不许声明能渲染")
            self.assertEqual(evidence["rasterizer"], "")
        ocr_engines = [name for name in ("easyocr", "paddleocr", "tesserocr", "pytesseract")
                       if importlib.util.find_spec(name) is not None]
        if not ocr_engines and not shutil.which("tesseract"):
            self.assertFalse(facts["imageOcr"], "没有 OCR 引擎就不许声明能识别")
            self.assertEqual(evidence["ocrEngine"], "")

    # ── 纯函数：页码区间与缺口文案 ──────────────────────────────────────────

    def test_page_range_text_compresses_consecutive_pages(self):
        module = load_script_module()
        self.assertEqual(module.page_range_text([7]), "7")
        self.assertEqual(module.page_range_text([3, 4, 5]), "3–5")
        self.assertEqual(module.page_range_text([3, 5, 6, 9, 10, 11]), "3、5–6、9–11")
        self.assertEqual(module.page_range_text([]), "")

    def test_gap_wording_without_a_renderer_matches_the_product_wording(self):
        module = load_script_module()
        result = module.empty_result("估值报告.pdf", "pdf")
        result["capabilities"]["pdfRender"] = False
        result["capabilityProbe"]["pdfImageExtractor"] = "pypdf"
        # 有内嵌图提取库、没有渲染器时，缺口就是产品文案里那一条。
        self.assertEqual(
            module.no_text_layer_gap([3, 4, 5], result),
            "第 3–5 页无文本层：本部署没有 PDF 页面渲染能力，请人工核对这几页的坐落/面积/权属",
        )
        # 两个通道都没有时，把"为什么拿不到图像"也说清楚（仍然只说事实）。
        result["capabilityProbe"]["pdfImageExtractor"] = ""
        self.assertEqual(
            module.no_text_layer_gap([3, 4, 5], result),
            "第 3–5 页无文本层：本部署没有 PDF 页面渲染能力，也没有可用的 PDF 内嵌图像提取库，"
            "请人工核对这几页的坐落/面积/权属",
        )
        # 已经导出页面图像时，缺口要指认产物（而不是说"什么都没有"）。
        result["images"] = [{"path": "x/第3页-图1.png", "description": "内容未识别"}]
        gap = module.no_text_layer_gap([3], result)
        self.assertIn("已导出内嵌页面图像 1 张", gap)
        self.assertIn("OCR", gap)
        self.assertIn("坐落/面积/权属", gap)
        self.assertNotIn("视觉识别能力", gap, "图像已经交给会话了，不许说本部署没有视觉通道")

    def test_gap_distinguishes_a_detected_renderer_from_implemented_rendering(self):
        module = load_script_module()
        result = module.empty_result("报告.pdf", "pdf")
        result["capabilities"]["pdfRender"] = True
        result["capabilityProbe"]["pdfImageExtractor"] = "pypdf"
        gap = module.no_text_layer_gap([2], result, "未取得可导出的 PDF 内嵌图像")
        self.assertIn("本脚本未接入页面渲染", gap)
        self.assertNotIn("没有 PDF 页面渲染能力", gap)
        self.assertIn("未取得可导出的 PDF 内嵌图像", gap)

    # ── 端到端：混合页 PDF ──────────────────────────────────────────────────

    def test_mixed_pdf_reports_per_page_facts_and_a_page_range_gap(self):
        if independent_pdf_library() == "":
            self.skipTest("本机没有 PDF 文本库（pypdf / PyPDF2 / pdfminer）：逐页事实无从产生")
        path = self.write_pdf("估值报告.pdf", ["Heli Text Page One", None, None])

        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertEqual(payload["pageCount"], 3)
        self.assertEqual([page["index"] for page in payload["pages"]], [1, 2, 3])
        self.assertTrue(payload["pages"][0]["hasTextLayer"], payload["pages"])
        self.assertEqual(payload["pages"][0]["chars"], len("HeliTextPageOne"))
        for page in payload["pages"][1:]:
            self.assertFalse(page["hasTextLayer"], page)
            self.assertEqual(page["chars"], 0, "无文本层的页不许报字数")
        self.assertEqual((payload["pagesWithText"], payload["pagesWithoutText"]), (1, 2))
        self.assertTrue(payload["capabilities"]["pdfText"], payload["capabilities"])

        # 读到的页照常有正文，读不到的页**一个字都不许编**。
        self.assertIn("Heli Text Page One", payload["text"])
        self.assertTrue(payload["ok"], payload)

        # 一条可执行的缺口，带页码区间与人工动作。
        gap_lines = [gap for gap in payload["capabilityGaps"] if "无文本层" in gap]
        self.assertEqual(len(gap_lines), 1, payload["capabilityGaps"])
        self.assertIn("第 2–3 页无文本层", gap_lines[0])
        self.assertIn("请人工核对这几页的坐落/面积/权属", gap_lines[0])

        # 内嵌页面图像：能导出就登记（并如实写"内容未识别"）；不能导出就说清为什么 —— 两者都不许沉默。
        if payload["images"]:
            for image in payload["images"]:
                self.assertTrue(image["path"], image)
                self.assertIn("内容未识别", image["description"])
            self.assertEqual(payload["capabilities"]["visionFallback"], "session-images")
        else:
            self.assertTrue(
                any(("渲染" in gap) or ("内嵌图像" in gap) for gap in payload["capabilityGaps"]),
                f"导出不了页面图像就必须说明原因：{payload['capabilityGaps']}",
            )
            self.assertEqual(payload["capabilities"]["visionFallback"], "none")

    def test_fully_scanned_pdf_names_every_page_and_never_invents_fields(self):
        if independent_pdf_library() == "":
            self.skipTest("本机没有 PDF 文本库：这一条走缺库分支（由脚本缺库用例覆盖）")
        path = self.write_pdf("扫描件.pdf", [None, None])

        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertEqual([page["hasTextLayer"] for page in payload["pages"]], [False, False])
        self.assertEqual(payload["pagesWithoutText"], 2)
        self.assertEqual(payload["extractedFields"], {}, "整份都是图，一个字段都不许出现")
        self.assertFalse(payload["ok"], "整份没有文本层就不算读出来了")
        self.assertTrue(
            any("第 1–2 页无文本层" in gap for gap in payload["capabilityGaps"]),
            payload["capabilityGaps"],
        )

    def test_embedded_page_images_are_exported_and_registered_by_the_library(self):
        """Embedded images can be exported without an installed renderer or OCR engine."""
        stub_dir = os.path.join(self.dir, "stub")
        os.makedirs(stub_dir, exist_ok=True)
        with open(os.path.join(stub_dir, "pypdf.py"), "w", encoding="utf-8") as handle:
            handle.write(FAKE_PYPDF)
        path = self.write_pdf("扫描件.pdf", [None, None])

        payload = run_script(
            path, self.dir, env={"PYTHONPATH": stub_dir, "PATH": ""}, python_args=("-S",),
        )
        self.assertShape(payload)
        self.assertEqual(payload["capabilityProbe"]["pdfTextLibrary"], "pypdf")
        self.assertFalse(payload["capabilities"]["pdfRender"], "替身不是渲染器，结论仍是 False")
        self.assertFalse(payload["capabilities"]["imageOcr"])
        self.assertEqual(payload["pagesWithoutText"], 1)
        self.assertEqual(len(payload["images"]), 1, payload["images"])

        image = payload["images"][0]
        self.assertIn("第2页-图1.png", image["path"])
        self.assertIn("第2页图片#1", image["description"], "定位符要能对回原文件里的那一页")
        self.assertIn("内容未识别", image["description"])
        self.assertNotIn("没有 OCR / 视觉识别能力", image["description"])
        gap = next(gap for gap in payload["capabilityGaps"] if "无文本层" in gap)
        self.assertIn("第 2 页无文本层", gap)
        self.assertIn("已导出内嵌页面图像 1 张", gap)
        self.assertEqual(payload["capabilities"]["visionFallback"], "session-images")
        self.assertNotIn("视觉识别能力", gap, "图像已交给会话，缺口不许自相矛盾地说没有视觉通道")
        # path 相对**结果 JSON 所在目录**（这里就是 workdir）：图像必须真的落在那里。
        landed = os.path.join(self.dir, image["path"])
        self.assertTrue(os.path.isfile(landed), f"图像没有落盘：{landed}")
        with open(landed, "rb") as handle:
            self.assertEqual(handle.read(), b"\x89PNG-fake-embedded-bytes")

    def test_image_only_page_reports_why_page_images_could_not_be_read(self):
        """页面带内嵌位图时：要么真的导出并登记（内容未识别），要么说清为什么导不出来。"""
        if independent_pdf_library() == "":
            self.skipTest("本机没有 PDF 文本库：内嵌图通道无从尝试")
        path = self.write_pdf("扫描件.pdf", [IMAGE_PAGE, None])

        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertEqual(payload["pagesWithoutText"], 2, payload["pages"])
        self.assertEqual(payload["extractedFields"], {}, "只读图不许编出字段")
        # Exported images are usable material; their contents remain unrecognized.
        self.assertEqual(payload["ok"], bool(payload["images"]), payload)
        gap = next(gap for gap in payload["capabilityGaps"] if "无文本层" in gap)
        self.assertIn("第 1–2 页无文本层", gap)
        if payload["images"]:
            # Every exported image must remain traceable without implying recognition.
            for image in payload["images"]:
                self.assertTrue(image["path"])
                self.assertIn("内容未识别", image["description"])
                landed = os.path.join(self.dir, image["path"])
                self.assertTrue(os.path.isfile(landed), f"图像没有落盘：{landed}")
            self.assertIn("已导出内嵌页面图像", gap)
            self.assertEqual(payload["capabilities"]["visionFallback"], "session-images")
        else:
            self.assertEqual(payload["ok"], False, "没有文本 / 表格 / 图像时不许算读出来了")
            self.assertTrue(
                ("内嵌图像" in gap) or ("渲染" in gap),
                f"导不出内嵌图就必须说清原因：{gap}",
            )
            self.assertEqual(payload["capabilities"]["visionFallback"], "none")

    def test_pdf_with_a_text_library_still_declares_what_it_cannot_do(self):
        """Reading PDF text preserves the measured render and OCR capabilities."""
        module = load_script_module()
        facts, _evidence = module.probe_capabilities()
        path = os.path.join(self.dir, "报告.pdf")
        with open(path, "wb") as handle:
            handle.write(build_pdf([None]))
        payload = run_script(path, self.dir)
        self.assertShape(payload)
        self.assertEqual(payload["capabilities"]["pdfRender"], facts["pdfRender"])
        self.assertEqual(payload["capabilities"]["imageOcr"], facts["imageOcr"])


if __name__ == "__main__":
    unittest.main(verbosity=2)

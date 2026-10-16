#!/usr/bin/env python3
"""Contract tests for generic workbook build/check (no bundled residential template, no legacy assets).

Pins the generic channel of `build_workbook.py` / `check_workbook_contract.py`:

1. Template and descriptor must be explicit (or an adjacent sidecar); no default fallback and no
   built-in seven-sheet profile.
2. Input cells are declared by the template's own descriptor, not guessed from fill colours.
3. All declared input cells are cleared before filling, so template samples cannot leak.
4. Formula / result cells and source template bytes stay untouched.
5. Generated output writes its own sidecar with source provenance that survives recalculation.
6. The checker only uses the descriptor; without one it refuses instead of falling back.

Fixtures are self-built workbooks in temp dirs (three arbitrary sheets), never bundled business assets.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

import test_report_data as fixtures

HERE = os.path.dirname(os.path.abspath(__file__))
BUILD = os.path.join(HERE, "build_workbook.py")
CHECK = os.path.join(HERE, "check_workbook_contract.py")


def run(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, script, *args],
        capture_output=True, text=True, timeout=180, check=False,
    )


def sha256_of(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(65536), b""):
            digest.update(block)
    return "sha256:" + digest.hexdigest()


def combined(completed: subprocess.CompletedProcess) -> str:
    return (completed.stdout or "") + (completed.stderr or "")


class GenericWorkbookTest(unittest.TestCase):
    """Self-built template plus explicit descriptor: neither build nor check may guess."""

    def setUp(self) -> None:
        try:
            import openpyxl  # type: ignore  # noqa: F401
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        self.directory = tempfile.mkdtemp(prefix="crwu-generic-workbook-")
        self.template = self._make_template()
        self.values = os.path.join(self.directory, "填入值.json")
        self.out = os.path.join(self.directory, "底稿.xlsx")
        self.descriptor = os.path.join(self.directory, "模板契约.json")

    def tearDown(self) -> None:
        import shutil

        shutil.rmtree(self.directory, ignore_errors=True)

    # ── fixtures ────────────────────────────────────────────────────────────────

    def _make_template(self) -> str:
        import openpyxl  # type: ignore

        path = os.path.join(self.directory, "简易模板.xlsx")
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for name in ("输入", "测算", "结果"):
            workbook.create_sheet(name)
        workbook["输入"]["B1"] = "本次填表单位：元"
        # Preset sample value: cleared when not provided so it cannot leak into the output.
        workbook["输入"]["B2"] = 120.5
        workbook["输入"]["B3"] = None
        workbook["测算"]["C3"] = "=ROUND(B2*100,0)"
        workbook["结果"]["B5"] = "=测算!C3"
        workbook.save(path)
        workbook.close()
        return path

    def _contract_payload(self, **patch) -> dict:
        payload = {
            "templateId": "简易测算底稿",
            "templateSha256": sha256_of(self.template),
            "sheets": [
                {"name": "输入", "role": "input"},
                {"name": "测算", "role": "formula"},
                {"name": "结果", "role": "result"},
            ],
            "inputCells": ["输入!B2", "输入!B3"],
            "resultCells": ["结果!B5"],
            "unitCells": ["输入!B1"],
        }
        payload.update(patch)
        return payload

    def _write_contract(self, **patch) -> str:
        with open(self.descriptor, "w", encoding="utf-8") as handle:
            json.dump(self._contract_payload(**patch), handle, ensure_ascii=False)
        return self.descriptor

    def _write_values(self, cells: dict) -> str:
        with open(self.values, "w", encoding="utf-8") as handle:
            json.dump({"cells": cells}, handle, ensure_ascii=False)
        return self.values

    def _build(self, *extra: str, cells: dict | None = None) -> subprocess.CompletedProcess:
        if cells is not None:
            self._write_values(cells)
        return run(
            BUILD,
            "--values", self.values,
            "--out", self.out,
            "--template", self.template,
            *extra,
        )

    # ── neutral checker fixtures (no builder involved) ─────────────────────

    def _make_neutral_workbook(self, name: str = "中性模板.xlsx", extra_cells=()) -> str:
        """One self-created workbook with a valid generic shape plus optional defect cells."""

        import openpyxl  # type: ignore

        path = os.path.join(self.directory, name)
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for sheet in ("输入", "测算", "结果"):
            workbook.create_sheet(sheet)
        workbook["输入"]["B1"] = "本次填表单位：元"
        workbook["输入"]["B2"] = None
        workbook["测算"]["C3"] = "=ROUND(B2*100,0)"
        workbook["结果"]["B5"] = "=测算!C3"
        for sheet, coord, value in extra_cells:
            workbook[sheet][coord] = value
        workbook.save(path)
        workbook.close()
        return path

    def _write_neutral_contract(self, workbook_path: str, **patch) -> str:
        payload = {
            "templateId": "中性底稿",
            "templateSha256": sha256_of(workbook_path),
            "sheets": [
                {"name": "输入", "role": "input"},
                {"name": "测算", "role": "formula"},
                {"name": "结果", "role": "result"},
            ],
            "inputCells": ["输入!B2"],
            "resultCells": ["结果!B5"],
            "unitCells": ["输入!B1"],
        }
        payload.update(patch)
        descriptor = workbook_path + ".contract.json"
        with open(descriptor, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False)
        return descriptor

    # ── explicit template / descriptor ────────────────────────────────────

    def test_missing_template_is_required_not_defaulted(self):
        self._write_contract()
        self._write_values({"输入!B3": 7})
        completed = run(
            BUILD, "--values", self.values, "--out", self.out, "--contract", self.descriptor,
        )
        self.assertNotEqual(completed.returncode, 0, combined(completed))
        self.assertFalse(os.path.exists(self.out), "没有模板时不许产出底稿")
        text = combined(completed)
        self.assertRegex(text, r"--template|必须.*模板|缺少.*模板")
        self.assertNotIn("缺少必需表", text, "不许退回内置七表 profile")

    def test_missing_descriptor_is_rejected_not_legacy_fallback(self):
        self._write_values({"输入!B3": 7})
        completed = self._build()  # neither --contract nor an adjacent sidecar
        self.assertNotEqual(completed.returncode, 0, combined(completed))
        self.assertFalse(os.path.exists(self.out), "没有契约时不许产出底稿")
        text = combined(completed)
        self.assertNotIn("缺少必需表", text, "不许退回内置七表 profile")
        self.assertRegex(text, r"契约|填写说明|描述")

    def test_explicit_descriptor_builds_arbitrary_sheet_structure(self):
        self._write_contract()
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))
        self.assertTrue(os.path.isfile(self.out))
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(self.out)
        try:
            self.assertEqual(sorted(workbook.sheetnames), ["测算", "结果", "输入"])
        finally:
            workbook.close()

    # ── input clearing / formula protection / source template unchanged ────────────────────────────────────

    def test_unprovided_prefilled_input_is_cleared_and_provided_value_kept(self):
        self._write_contract()
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(self.out)
        try:
            self.assertIsNone(workbook["输入"]["B2"].value,
                              "descriptor 声明的输入位未提供时，模板样例值必须被清空")
            self.assertEqual(workbook["输入"]["B3"].value, 7, "提供的值要保留")
        finally:
            workbook.close()

    def test_formula_and_result_cells_are_unchanged(self):
        self._write_contract()
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(self.out, data_only=False)
        try:
            self.assertEqual(workbook["测算"]["C3"].value, "=ROUND(B2*100,0)")
            self.assertEqual(workbook["结果"]["B5"].value, "=测算!C3")
        finally:
            workbook.close()

    def test_source_template_bytes_are_unchanged(self):
        self._write_contract()
        before = sha256_of(self.template)
        with open(self.template, "rb") as handle:
            before_bytes = handle.read()
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))
        self.assertEqual(sha256_of(self.template), before)
        with open(self.template, "rb") as handle:
            self.assertEqual(handle.read(), before_bytes,
                             "生成只读模板：源模板一个字节都不许变")

    # ── rejection paths ────────────────────────────────────────────────────────────

    def test_input_formula_is_rejected_without_output(self):
        self._write_contract()
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": "=1+1"})
        self.assertNotEqual(completed.returncode, 0)
        self.assertFalse(os.path.exists(self.out))
        self.assertIn("输入区不许写公式", combined(completed))

    def test_cell_outside_declared_inputs_is_rejected_without_output(self):
        self._write_contract()
        completed = self._build("--contract", self.descriptor, cells={"输入!B9": 1})
        self.assertNotEqual(completed.returncode, 0)
        self.assertFalse(os.path.exists(self.out))
        self.assertRegex(combined(completed), r"不在.*输入区|输入区")

    def test_input_result_conflict_is_rejected_without_output(self):
        # One cell declared as both input and result: contradictory, never resolved silently.
        self._write_contract(resultCells=["结果!B5", "输入!B2"])
        completed = self._build("--contract", self.descriptor, cells={"输入!B2": 1})
        self.assertNotEqual(completed.returncode, 0)
        self.assertFalse(os.path.exists(self.out))
        self.assertRegex(combined(completed), r"公式区|结果位")

    # ── output sidecar and standalone checking ─────────────────────────────────────────────

    def test_output_sidecar_keeps_source_provenance_not_mutable_output_bytes(self):
        self._write_contract()
        source_digest = sha256_of(self.template)
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))

        sidecar = self.out + ".contract.json"
        self.assertTrue(os.path.isfile(sidecar), "生成件必须带自己的 sidecar 供独立校验发现")
        with open(sidecar, encoding="utf-8") as handle:
            descriptor = json.load(handle)
        self.assertEqual(descriptor.get("sourceTemplateSha256"), source_digest,
                         "sidecar 要如实标注**来源模板**指纹")
        output_digest = sha256_of(self.out)
        self.assertNotEqual(source_digest, output_digest,
                            "填充后的产物字节必然与源模板不同（夹具前提）")
        self.assertNotIn("templateSha256", descriptor,
                         "sidecar 不得把产物字节哈希写死：真正重算 / 重存一次就对不上了")

    def test_output_sidecar_survives_recalculation_style_save(self):
        self._write_contract()
        source_digest = sha256_of(self.template)
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))
        before = sha256_of(self.out)

        # A recalc engine re-saves the workbook with cached values: bytes change, structure does not.
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(self.out, data_only=False)
        workbook["结果"]["B5"] = 61200
        workbook.save(self.out)
        workbook.close()
        self.assertNotEqual(sha256_of(self.out), before, "夹具前提：重存后字节必须变")

        verified = run(CHECK, self.out)
        self.assertEqual(verified.returncode, 0, combined(verified))
        with open(self.out + ".contract.json", encoding="utf-8") as handle:
            descriptor = json.load(handle)
        self.assertEqual(descriptor.get("sourceTemplateSha256"), source_digest,
                         "重存不得把来源模板指纹改写成产物字节")

    def test_output_sidecar_is_discovered_and_passes_standalone_check(self):
        self._write_contract()
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))
        # Without --contract the standalone checker must discover and pass the sidecar.
        verified = run(CHECK, self.out)
        self.assertEqual(verified.returncode, 0, combined(verified))
        self.assertIn("符合模板契约", verified.stdout)

    def test_checker_rejects_absent_descriptor_instead_of_legacy_profile(self):
        # No sidecar next to the source template; no fallback to a built-in profile.
        verified = run(CHECK, self.template)
        self.assertNotEqual(verified.returncode, 0)
        text = combined(verified)
        self.assertNotIn("缺少必需表", text, "缺 descriptor 时要明确拒绝，而不是套内置 profile")
        self.assertRegex(text, r"契约|填写说明|描述")

    def test_explicit_expected_workbook_digest_mismatch_is_rejected(self):
        self._write_contract()
        completed = self._build("--contract", self.descriptor, cells={"输入!B3": 7})
        self.assertEqual(completed.returncode, 0, combined(completed))
        # Check the output with the source descriptor: the byte digest must mismatch.
        verified = run(CHECK, self.out, "--contract", self.descriptor)
        self.assertNotEqual(verified.returncode, 0)
        self.assertIn("模板字节指纹与契约声明不一致", combined(verified))

    # ── generic checking must really scan cell defects ─────────────────────────────────────────

    def test_generic_checker_rejects_error_value_cells(self):
        workbook = self._make_neutral_workbook("错误值.xlsx", extra_cells=[
            ("测算", "C4", "=#REF!"),
            ("结果", "B6", "#DIV/0!"),
        ])
        self._write_neutral_contract(workbook)
        verified = run(CHECK, workbook)
        self.assertNotEqual(verified.returncode, 0)
        self.assertRegex(combined(verified), r"#REF!|#DIV/0!|错误值")

    def test_generic_checker_rejects_external_workbook_reference(self):
        workbook = self._make_neutral_workbook("外部引用.xlsx", extra_cells=[
            ("测算", "C4", "=SUM([1]外部表.xlsx!A1:A3)"),
        ])
        self._write_neutral_contract(workbook)
        verified = run(CHECK, workbook)
        self.assertNotEqual(verified.returncode, 0)
        self.assertRegex(combined(verified), r"外部|跨工作簿|\[1\]|链接")

    def test_generic_checker_rejects_broken_formula_reference(self):
        workbook = self._make_neutral_workbook("断链.xlsx", extra_cells=[
            ("结果", "B7", "=测算!#REF"),
        ])
        self._write_neutral_contract(workbook)
        verified = run(CHECK, workbook)
        self.assertNotEqual(verified.returncode, 0)
        self.assertRegex(combined(verified), r"断链|#REF")

    def test_generic_checker_accepts_arbitrary_valid_formulas(self):
        workbook = self._make_neutral_workbook("有效公式.xlsx", extra_cells=[
            ("测算", "C4", "=SUM(C3:C3)"),
            ("结果", "B6", "=IF(测算!C3>0,测算!C3,0)"),
        ])
        self._write_neutral_contract(workbook)
        verified = run(CHECK, workbook)
        self.assertEqual(verified.returncode, 0, combined(verified))


class CaseEvidenceWorkbookTest(unittest.TestCase):
    """`caseEvidence` extends the workbook with a visible, traceable 案例证据 sheet."""

    def setUp(self) -> None:
        try:
            import openpyxl  # type: ignore  # noqa: F401
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        self.directory = tempfile.mkdtemp(prefix="crwu-case-evidence-")
        self.template = self._make_template()
        self.descriptor = os.path.join(self.directory, "模板契约.json")
        payload = {
            "templateId": "简易测算底稿",
            "templateSha256": sha256_of(self.template),
            "sheets": [
                {"name": "输入", "role": "input"},
                {"name": "测算", "role": "formula"},
                {"name": "结果", "role": "result"},
            ],
            "inputCells": ["输入!B2"],
            "resultCells": ["结果!B5"],
            "unitCells": ["输入!B1"],
        }
        with open(self.descriptor, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False)

    def tearDown(self) -> None:
        shutil.rmtree(self.directory, ignore_errors=True)

    def _make_template(self) -> str:
        import openpyxl  # type: ignore

        path = os.path.join(self.directory, "简易模板.xlsx")
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for name in ("输入", "测算", "结果"):
            workbook.create_sheet(name)
        workbook["输入"]["B1"] = "本次填表单位：元"
        workbook["输入"]["B2"] = None
        workbook["测算"]["C3"] = "=ROUND(B2*100,0)"
        workbook["结果"]["B5"] = "=测算!C3"
        workbook.save(path)
        workbook.close()
        return path

    def _write_values(self, payload: dict) -> str:
        path = os.path.join(self.directory, "填入值.json")
        with open(path, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False)
        return path

    def _project_with_evidence(self, relative: str, url: str,
                               page_type: str = "detail") -> tuple[str, str, str, str]:
        project = os.path.join(self.directory, "案例-平面 图")
        plan = os.path.join(project, "方案", "plan-001")
        os.makedirs(plan, exist_ok=True)
        out = os.path.join(plan, "估值底稿.xlsx")
        target = os.path.join(plan, relative)
        os.makedirs(os.path.dirname(target), exist_ok=True)
        payload_bytes = fixtures.png_bytes()
        with open(target, "wb") as handle:
            handle.write(payload_bytes)
        digest = "sha256:" + hashlib.sha256(payload_bytes).hexdigest()
        values = self._write_values({
            "cells": {"输入!B2": 7},
            "caseEvidence": [{
                "caseId": "case-001", "displayName": "阳光小区 3室2厅", "provider": "链家",
                "sourceUrl": url, "relativePath": relative, "digest": digest,
                "capturedAt": "2026-10-16T00:00:00Z", "pageType": page_type,
            }],
        })
        return project, out, values, relative

    def _build(self, out: str, values: str) -> subprocess.CompletedProcess:
        return run(BUILD, "--values", values, "--out", out,
                   "--template", self.template, "--contract", self.descriptor)

    def test_case_evidence_sheet_shows_full_url_and_clickable_relative_image(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        url = "https://example.invalid/detail/阳光 小区?id=7"
        _project, out, values, _relative = self._project_with_evidence(relative, url)
        template_bytes = open(self.template, "rb").read()

        completed = self._build(out, values)
        self.assertEqual(completed.returncode, 0, combined(completed))
        self.assertEqual(open(self.template, "rb").read(), template_bytes, "源模板字节不许变")

        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            self.assertIn("案例证据", workbook.sheetnames, "values 带 caseEvidence 时要追加可见的案例证据表")
            self.assertEqual(workbook["测算"]["C3"].value, "=ROUND(B2*100,0)", "公式不能被改写")
            self.assertEqual(workbook["结果"]["B5"].value, "=测算!C3")
            cells = [cell for row in workbook["案例证据"].iter_rows() for cell in row]
            url_cells = [cell for cell in cells
                         if isinstance(cell.value, str) and "example.invalid/detail/" in cell.value]
            self.assertTrue(url_cells, "案例证据表要显示完整详情 URL 文本")
            self.assertTrue(any(cell.hyperlink is not None and cell.hyperlink.target == url for cell in url_cells),
                            "详情 URL 必须是原生超链接")
            image_cells = [cell for cell in cells
                           if isinstance(cell.value, str) and cell.value.endswith(".png")]
            self.assertTrue(image_cells, "案例证据表要显示可读的图片相对路径")
            self.assertTrue(any(cell.hyperlink is not None and "推理结果/图片/" in (cell.hyperlink.target or "")
                                for cell in image_cells), "图片相对路径必须是原生超链接")
            self.assertEqual(len(getattr(workbook["案例证据"], "_images", [])), 0, "不要求内嵌图片")
        finally:
            workbook.close()

        with open(out + ".contract.json", encoding="utf-8") as handle:
            sidecar = json.load(handle)
        self.assertIn("case-001", json.dumps(sidecar, ensure_ascii=False),
                      "sidecar 要记录证据索引元数据供独立校验")

    def test_evidence_links_stay_relative_after_moving_the_project_tree(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        completed = self._build(out, values)
        self.assertEqual(completed.returncode, 0, combined(completed))
        moved = os.path.join(self.directory, "移动后")
        shutil.move(project, moved)
        moved_out = os.path.join(moved, "方案", "plan-001", "估值底稿.xlsx")
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(moved_out)
        try:
            self.assertIn("案例证据", workbook.sheetnames, "移动后仍要有案例证据表")
            cells = [cell for row in workbook["案例证据"].iter_rows() for cell in row]
            links = [cell.hyperlink.target for cell in cells
                     if cell.hyperlink is not None and cell.hyperlink.target]
            png_links = [target for target in links if str(target).endswith(".png")]
            self.assertTrue(png_links, "移动后仍要有图片相对链接")
            for target in links:
                self.assertFalse(os.path.isabs(str(target)), f"超链接必须是相对路径：{target}")
                self.assertNotRegex(str(target), r"^[A-Za-z]:", f"不许出现盘符路径：{target}")
            for target in png_links:
                resolved = os.path.normpath(os.path.join(os.path.dirname(moved_out), str(target)))
                self.assertTrue(os.path.isfile(resolved), f"移动后链接必须解析到真实图片：{target}")
        finally:
            workbook.close()

    def test_missing_image_is_not_claimed_as_traced(self):
        relative = "推理结果/图片/case-001-阳光 小区/不存在.png"
        project = os.path.join(self.directory, "缺图")
        plan = os.path.join(project, "方案", "plan-001")
        os.makedirs(plan, exist_ok=True)
        out = os.path.join(plan, "估值底稿.xlsx")
        values = self._write_values({
            "cells": {"输入!B2": 7},
            "caseEvidence": [{
                "caseId": "case-001", "displayName": "阳光小区 3室2厅", "provider": "链家",
                "sourceUrl": "https://example.invalid/detail/1", "relativePath": relative,
                "digest": "sha256:" + "0" * 64, "capturedAt": "2026-10-16T00:00:00Z",
                "pageType": "detail",
            }],
        })
        completed = self._build(out, values)
        self.assertNotEqual(completed.returncode, 0, "缺图必须显式失败，不许静默省略索引还当成功")

    def test_unsafe_source_url_is_not_written_as_a_hyperlink(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        _project, out, values, _relative = self._project_with_evidence(
            relative, "javascript:alert(1)")
        completed = self._build(out, values)
        self.assertNotEqual(completed.returncode, 0, "危险 URL 必须显式失败")

    def test_dotdot_image_path_outside_image_folder_is_rejected(self):
        relative = "推理结果/图片/../outside.png"
        _project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        completed = self._build(out, values)
        self.assertNotEqual(completed.returncode, 0, "relativePath 不许用 ../ 逃出 推理结果/图片/")

    def test_image_directory_symlink_escaping_case_root_is_rejected(self):
        # A symlinked 推理结果/图片 can point the whole evidence tree at a sibling outside the
        # case; a valid PNG with a matching digest elsewhere must still not be accepted.
        project = os.path.join(self.directory, "案例-越界")
        plan = os.path.join(project, "方案", "plan-001")
        outside = os.path.join(self.directory, "case-outside")
        os.makedirs(plan, exist_ok=True)
        os.makedirs(outside, exist_ok=True)
        payload_bytes = fixtures.png_bytes()
        with open(os.path.join(outside, "shot.png"), "wb") as handle:
            handle.write(payload_bytes)
        digest = "sha256:" + hashlib.sha256(payload_bytes).hexdigest()

        image_dir = os.path.join(project, "推理结果", "图片")
        os.makedirs(os.path.dirname(image_dir), exist_ok=True)
        try:
            os.symlink(outside, image_dir, target_is_directory=True)
        except (AttributeError, NotImplementedError) as error:
            self.skipTest(f"本平台没有 os.symlink：{error}（Windows 运行时未验证）")
        except OSError as error:
            self.skipTest(f"本平台不允许创建符号链接：{error}（Windows 运行时未验证）")

        out = os.path.join(plan, "估值底稿.xlsx")
        values = self._write_values({
            "cells": {"输入!B2": 7},
            "caseEvidence": [{
                "caseId": "case-001", "displayName": "越界图", "provider": "链家",
                "sourceUrl": "https://example.invalid/detail/1",
                "relativePath": "推理结果/图片/shot.png", "digest": digest,
                "capturedAt": "2026-10-16T00:00:00Z", "pageType": "detail",
            }],
        })
        completed = run(BUILD, "--values", values, "--out", out, "--template", self.template,
                        "--contract", self.descriptor, "--case-root", project)
        self.assertNotEqual(completed.returncode, 0, "图片目录实体越出案例根时必须失败")
        self.assertRegex(combined(completed), r"图片|越出|案例根")
        self.assertFalse(os.path.exists(out), "被拒绝时不许产出底稿")

    def test_http_url_without_host_is_rejected(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        _project, out, values, _relative = self._project_with_evidence(relative, "http://")
        completed = self._build(out, values)
        self.assertNotEqual(completed.returncode, 0, "http:// 没有主机名不算可核对的 URL")

    def test_missing_or_unknown_evidence_type_is_not_silently_detail(self):
        project = os.path.join(self.directory, "无类型")
        plan = os.path.join(project, "方案", "plan-001")
        os.makedirs(plan, exist_ok=True)
        out = os.path.join(plan, "估值底稿.xlsx")
        relative = "推理结果/图片/case-001-阳光小区/详情.png"
        target = os.path.join(plan, relative)
        os.makedirs(os.path.dirname(target), exist_ok=True)
        payload_bytes = fixtures.png_bytes()
        with open(target, "wb") as handle:
            handle.write(payload_bytes)
        digest = "sha256:" + hashlib.sha256(payload_bytes).hexdigest()
        values = self._write_values({
            "cells": {"输入!B2": 7},
            "caseEvidence": [{
                "caseId": "case-001", "displayName": "阳光小区", "provider": "链家",
                "sourceUrl": "https://example.invalid/detail/1", "relativePath": relative,
                "digest": digest, "capturedAt": "2026-10-16T00:00:00Z",
            }],
        })
        completed = self._build(out, values)
        self.assertNotEqual(completed.returncode, 0, "缺少证据类型时不许默认成详情截图")

    def test_explicit_case_root_resolves_links_from_workbook_parent(self):
        project = os.path.join(self.directory, "案例-根")
        plan = os.path.join(project, "方案", "plan-001")
        os.makedirs(plan, exist_ok=True)
        out = os.path.join(plan, "估值底稿.xlsx")
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        target = os.path.join(project, relative)
        os.makedirs(os.path.dirname(target), exist_ok=True)
        payload_bytes = fixtures.png_bytes()
        with open(target, "wb") as handle:
            handle.write(payload_bytes)
        digest = "sha256:" + hashlib.sha256(payload_bytes).hexdigest()
        values = self._write_values({
            "cells": {"输入!B2": 7},
            "caseEvidence": [{
                "caseId": "case-001", "displayName": "阳光小区", "provider": "链家",
                "sourceUrl": "https://example.invalid/detail/1", "relativePath": relative,
                "digest": digest, "capturedAt": "2026-10-16T00:00:00Z", "pageType": "detail",
            }],
        })
        completed = run(BUILD, "--values", values, "--out", out, "--template", self.template,
                        "--contract", self.descriptor, "--case-root", project)
        self.assertEqual(completed.returncode, 0, combined(completed))
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            cells = [cell for row in workbook["案例证据"].iter_rows() for cell in row]
            image_links = [cell.hyperlink.target for cell in cells
                           if cell.hyperlink is not None and str(cell.hyperlink.target).endswith(".png")]
            self.assertTrue(image_links, "要有图片相对链接")
            for link in image_links:
                resolved = os.path.normpath(os.path.join(os.path.dirname(out), str(link)))
                self.assertEqual(resolved, target, "链接必须从工作簿父目录解析到真实图片")
        finally:
            workbook.close()

    def test_checker_rejects_tampered_case_evidence_index(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        built = self._build(out, values)
        self.assertEqual(built.returncode, 0, combined(built))
        self.assertEqual(run(CHECK, out).returncode, 0, "原始索引应通过独立校验")

        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            sheet = workbook["案例证据"]
            headers = [cell.value for cell in sheet[1]]
            image_column = headers.index("图片相对路径") + 1
            sheet.cell(row=2, column=image_column).hyperlink = "推理结果/图片/别的.png"
            workbook.save(out)
        finally:
            workbook.close()
        tampered = run(CHECK, out)
        self.assertNotEqual(tampered.returncode, 0, "被篡改的图片链接必须被独立校验拒绝")

    def test_checker_rejects_report_used_case_missing_from_index(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        _project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        built = self._build(out, values)
        self.assertEqual(built.returncode, 0, combined(built))
        report = os.path.join(self.directory, "估值报告数据.json")
        with open(report, "w", encoding="utf-8") as handle:
            json.dump({"comparables": [{
                "caseId": "case-002", "used": True, "sourceUrl": "https://example.invalid/detail/2",
            }]}, handle, ensure_ascii=False)
        verified = run(CHECK, out, "--report-data", report)
        self.assertNotEqual(verified.returncode, 0, "报告采用但索引里没有的案例必须失败")

    def test_checker_rejects_deleted_case_evidence_sheet(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        _project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        self.assertEqual(self._build(out, values).returncode, 0)
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            del workbook["案例证据"]
            workbook.save(out)
        finally:
            workbook.close()
        self.assertNotEqual(run(CHECK, out).returncode, 0, "删掉证据索引必须失败")

    def test_checker_rejects_tampered_url_hyperlink(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        _project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        self.assertEqual(self._build(out, values).returncode, 0)
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            sheet = workbook["案例证据"]
            headers = [cell.value for cell in sheet[1]]
            url_column = headers.index("详情 URL") + 1
            sheet.cell(row=2, column=url_column).hyperlink = "https://example.invalid/detail/evil"
            workbook.save(out)
        finally:
            workbook.close()
        self.assertNotEqual(run(CHECK, out).returncode, 0, "篡改详情 URL 超链接必须失败")

    def test_checker_rejects_tampered_digest_text(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        _project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        self.assertEqual(self._build(out, values).returncode, 0)
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            sheet = workbook["案例证据"]
            headers = [cell.value for cell in sheet[1]]
            digest_column = headers.index("指纹") + 1
            sheet.cell(row=2, column=digest_column).value = "sha256:" + "0" * 64
            workbook.save(out)
        finally:
            workbook.close()
        self.assertNotEqual(run(CHECK, out).returncode, 0, "篡改指纹文本必须失败")

    def test_metadata_beginning_with_equals_stays_literal_text(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        project = os.path.join(self.directory, "等号元数据")
        plan = os.path.join(project, "方案", "plan-001")
        os.makedirs(plan, exist_ok=True)
        out = os.path.join(plan, "估值底稿.xlsx")
        target = os.path.join(plan, relative)
        os.makedirs(os.path.dirname(target), exist_ok=True)
        payload_bytes = fixtures.png_bytes()
        with open(target, "wb") as handle:
            handle.write(payload_bytes)
        digest = "sha256:" + hashlib.sha256(payload_bytes).hexdigest()
        values = self._write_values({
            "cells": {"输入!B2": 7},
            "caseEvidence": [{
                "caseId": "case-001", "displayName": "=SUM(A1:A2)", "provider": "链家",
                "sourceUrl": "https://example.invalid/detail/1", "relativePath": relative,
                "digest": digest, "capturedAt": "2026-10-16T00:00:00Z", "pageType": "detail",
            }],
        })
        completed = self._build(out, values)
        self.assertEqual(completed.returncode, 0, combined(completed))
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            cell = workbook["案例证据"].cell(row=2, column=2)
            self.assertEqual(cell.value, "=SUM(A1:A2)")
            self.assertNotEqual(cell.data_type, "f", "元数据必须是字面文本，不能变成公式")
        finally:
            workbook.close()
        self.assertEqual(run(CHECK, out).returncode, 0, "字面等号文本不应让校验失败")

    def test_checker_rejects_duplicate_case_evidence_rows(self):
        relative = "推理结果/图片/case-001-阳光 小区/详情.png"
        _project, out, values, _relative = self._project_with_evidence(
            relative, "https://example.invalid/detail/1")
        self.assertEqual(self._build(out, values).returncode, 0)
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(out)
        try:
            sheet = workbook["案例证据"]
            sheet.cell(row=3, column=1).value = "case-001"
            workbook.save(out)
        finally:
            workbook.close()
        self.assertNotEqual(run(CHECK, out).returncode, 0, "重复 caseId 必须失败")


class ReportEvidenceChoiceTest(unittest.TestCase):
    def test_chosen_detail_can_coexist_with_other_valid_case_images(self):
        import check_workbook_contract as checker

        with tempfile.TemporaryDirectory() as root:
            report_path = os.path.join(root, "report.json")
            entry = {"caseId": "case-001", "sourceUrl": "https://example.invalid/1",
                     "relativePath": "detail.png", "digest": "sha256:detail"}
            shots = [{"relativePath": "list.png", "digest": "sha256:list", "validity": "valid"},
                     {"relativePath": "detail.png", "digest": "sha256:detail", "validity": "valid"}]
            with open(report_path, "w", encoding="utf-8") as handle:
                json.dump({"comparables": [{"caseId": "case-001", "used": True,
                                           "sourceUrl": entry["sourceUrl"], "screenshots": shots}]}, handle)
            problems = []
            checker.check_report_data(report_path, [entry], problems)
            self.assertEqual(problems, [], "索引选择一张详情图时，其他列表图或历史图仍可保留")


if __name__ == "__main__":
    unittest.main(verbosity=2)

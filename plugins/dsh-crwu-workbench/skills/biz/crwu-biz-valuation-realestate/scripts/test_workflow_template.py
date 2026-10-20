"""Exercise the actual prepared workbook used by the graph plan and calculator."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

from test_listing_reference import SCRIPT, listing_input, run_cli, run_validate, write_input


class WorkflowTemplateTest(unittest.TestCase):
    def prepare(self, root: str) -> tuple[dict, str]:
        result = subprocess.run(
            [sys.executable, SCRIPT, "--prepare-template", "--out-dir", root, "--case-root", root],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        receipt = json.loads(result.stdout)
        self.assertTrue(receipt["ok"])
        template = receipt["template"]
        path = os.path.join(root, template["relativePath"])
        self.assertTrue(os.path.isfile(path))
        return template, path

    def test_plan_template_is_a_real_blank_workbook_with_verified_mapping(self):
        with tempfile.TemporaryDirectory() as root:
            template, path = self.prepare(root)
            with open(path, "rb") as handle:
                self.assertEqual(template["digest"], "sha256:" + hashlib.sha256(handle.read()).hexdigest())
            self.assertTrue(template["verified"])
            self.assertEqual(template["fileName"], os.path.basename(path))
            mapping = {item["field"]: (item["sheet"], item["range"]) for item in template["mapping"]}
            self.assertEqual(mapping["unitPrice"], ("估值结果", "B5"))
            self.assertEqual(mapping["totalPrice"], ("估值结果", "C5"))
            self.assertTrue(template["formulaChecks"])
            self.assertTrue(all(check["ok"] for check in template["formulaChecks"]))
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["估值结果", "挂牌来源", "计算轨迹", "限制", "案例证据"])
            self.assertIsNone(workbook["估值结果"]["B5"].value)
            self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))
            workbook.close()

    def test_preparation_reuses_unchanged_bytes_and_refuses_tampered_template(self):
        with tempfile.TemporaryDirectory() as root:
            first, path = self.prepare(root)
            self.assertEqual(self.prepare(root)[0], first)
            with open(path, "ab") as handle:
                handle.write(b"changed")
            with open(path, "rb") as handle:
                before = handle.read()
            result = subprocess.run(
                [sys.executable, SCRIPT, "--prepare-template", "--out-dir", root, "--case-root", root],
                capture_output=True, text=True,
            )
            self.assertNotEqual(result.returncode, 0)
            with open(path, "rb") as handle:
                self.assertEqual(handle.read(), before)

    def test_calculation_fills_the_exact_approved_template_and_preserves_original(self):
        with tempfile.TemporaryDirectory() as directory:
            input_path, plan_dir, case_root = write_input(directory, listing_input())
            template, path = self.prepare(plan_dir)
            workbook = load_workbook(path)
            workbook["限制"]["Z20"] = "approved-template-marker"
            workbook.save(path)
            workbook.close()
            with open(path, "rb") as handle:
                original = handle.read()
            template["digest"] = "sha256:" + hashlib.sha256(original).hexdigest()
            with open(input_path, encoding="utf-8") as handle:
                payload = json.load(handle)
            payload["template"] = template
            with open(input_path, "w", encoding="utf-8") as handle:
                json.dump(payload, handle, ensure_ascii=False)
            result = run_cli(input_path, plan_dir, case_root)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            output = load_workbook(os.path.join(plan_dir, "估值底稿.xlsx"))
            self.assertEqual(output["限制"]["Z20"].value, "approved-template-marker")
            self.assertEqual(float(output["估值结果"]["B5"].value), 41000)
            self.assertEqual(float(output["估值结果"]["C5"].value), 4100000)
            output.close()
            with open(path, "rb") as handle:
                self.assertEqual(handle.read(), original)
            checked = run_validate(plan_dir, case_root, input_path)
            self.assertEqual(checked.returncode, 0, checked.stdout + checked.stderr)

    def test_template_digest_failure_does_not_generate_any_delivery(self):
        with tempfile.TemporaryDirectory() as directory:
            input_path, plan_dir, case_root = write_input(directory, listing_input())
            template, path = self.prepare(plan_dir)
            template["digest"] = "sha256:" + "0" * 64
            with open(input_path, encoding="utf-8") as handle:
                payload = json.load(handle)
            payload["template"] = template
            with open(input_path, "w", encoding="utf-8") as handle:
                json.dump(payload, handle, ensure_ascii=False)
            result = run_cli(input_path, plan_dir, case_root)
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(os.path.exists(os.path.join(plan_dir, "估值底稿.xlsx")))


if __name__ == "__main__":
    unittest.main()

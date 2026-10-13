#!/usr/bin/env python3
"""交付契约脚本的契约测试：`validate_result.py` 与 `check_workbook_contract.py`。

钉住的是**每条硬规则真的会让校验失败** —— 校验脚本最糟的失效模式是"永远说 OK"：
它看起来在工作，而基准日没填、挂牌价被当成交价、数字填 0 这些事照样交付出去。

工作簿侧的测试**不重新生成模板**（那需要造表逻辑，属于 `build_workbook_template.py` 自己）：
它只核对随包发布的 `template/valuation-workbook.xlsx` 过不过契约，以及"改坏一处"能被抓住。
"""

from __future__ import annotations

import copy
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
VALIDATE = os.path.join(HERE, "validate_result.py")
CONTRACT = os.path.join(HERE, "check_workbook_contract.py")
BUILD_WORKBOOK = os.path.join(HERE, "build_workbook.py")
BUILD_REPORT = os.path.join(HERE, "build_report.py")
SAMPLE = os.path.join(SKILL, "examples", "valuation-result.sample.json")
CASE_SAMPLE = os.path.join(SKILL, "examples", "valuation-case.sample.json")
SCHEMA_DIR = os.path.join(SKILL, "schemas")
TEMPLATE = os.path.join(SKILL, "template", "valuation-workbook.xlsx")
REPORT_TEMPLATE = os.path.join(SKILL, "template", "valuation-report.html")

# `validate_result.py` 里的规则词表与 `market_rules.py` 必须一致，所以这里也要 import 它。
sys.path.insert(0, HERE)
import validate_result as validate  # noqa: E402


def fill_case_evidence(workbook_path: str, payload: dict) -> None:
    """把结果 JSON 该落进底稿的那几处填上（夹具替身）。

    正常路径由 `build_workbook.py` 从 `填入值.json` 填；测试夹具直接抄模板，
    所以这里补上那一步 —— 否则 `--workbook` 的跨交付件一致性判据会（正确地）判红：
    采用案例的证据留档、以及与结果 JSON 一致的「待确认事项」。
    """
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return
    adopted = [item for item in (payload.get("cases") or [])
               if isinstance(item, dict) and item.get("used_as_comparable") is True
               and str(item.get("evidence", "")).strip() != ""]
    items = [str(item).strip() for item in (payload.get("manual_confirmation_items") or [])
             if str(item).strip() != ""]
    if not adopted and not items:
        return
    workbook = openpyxl.load_workbook(workbook_path)
    try:
        sheet = workbook["案例"]
        for row in sheet.iter_rows(min_col=1, max_col=1):
            value = row[0].value
            if isinstance(value, str) and value.strip().startswith("案例截图 / 证据留档"):
                for index, item in enumerate(adopted):
                    sheet.cell(row=row[0].row, column=2 + index, value=str(item["evidence"]))
                break
        if items:
            summary = workbook["估值结果汇总"]
            for row in summary.iter_rows(min_col=1, max_col=1):
                value = row[0].value
                if isinstance(value, str) and value.strip() == "待确认事项":
                    summary.cell(row=row[0].row, column=2).value = "\n".join(items)
                    break
        workbook.save(workbook_path)
    finally:
        workbook.close()


def run(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, script, *args],
        capture_output=True, text=True, timeout=180, check=False,
    )


def load_sample() -> dict:
    with open(SAMPLE, "r", encoding="utf-8") as handle:
        return json.load(handle)


class ValidateResultTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="crwu-valuation-result-")
        self.result = os.path.join(self.dir, "估值结果.json")
        self.html = os.path.join(self.dir, "估值报告.html")
        self.workbook = os.path.join(self.dir, "估值底稿.xlsx")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write(self, payload: dict, *, with_artifacts: bool = True) -> None:
        with open(self.result, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False)
        # 案例引用的证据文件也要真的落盘：**证据引用了却不在，本身就是不合格**
        #（`check_case_evidence_files`）。夹具不造这些文件的话，所有用例都会因为这条而红。
        for item in payload.get("cases") or []:
            evidence = str(item.get("evidence", "")).strip()
            if evidence == "" or os.path.isabs(evidence):
                continue
            target = os.path.join(self.dir, evidence)
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with open(target, "w", encoding="utf-8") as handle:
                handle.write("<html>evidence</html>")
        if with_artifacts:
            with open(self.html, "w", encoding="utf-8") as handle:
                handle.write(f"<html><body>{payload.get('valuationId','')} 单价 {payload.get('unitPrice')} 总价 {payload.get('totalValue')}</body></html>")
            shutil.copyfile(TEMPLATE, self.workbook)
            # 采用案例的证据要同时写进底稿的「案例截图 / 证据留档」行 ——
            # 那是 `--workbook` 的跨交付件一致性判据（§8「截图编号」/ §9 三者一致）。
            self.fill_case_evidence(payload)

    def fill_case_evidence(self, payload: dict) -> None:
        fill_case_evidence(self.workbook, payload)

    def assertFails(self, payload: dict, pattern: str, *, with_artifacts: bool = True) -> None:
        self.write(payload, with_artifacts=with_artifacts)
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1, f"这条不该通过：\n{completed.stdout}")
        self.assertRegex(completed.stdout, pattern)

    # ── 正向 ────────────────────────────────────────────────────────────────

    def test_sample_passes_with_artifacts_present(self):
        self.write(load_sample())
        completed = run(VALIDATE, self.result, "--html", self.html, "--workbook", self.workbook)
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        self.assertIn("OK:", completed.stdout)

    # ── 每条硬规则都要真的红 ────────────────────────────────────────────────

    def test_auto_method_is_not_a_conclusion(self):
        payload = load_sample()
        payload["method"] = "auto"
        self.assertFails(payload, r"method 是 `auto`")

    def test_missing_base_date_must_be_registered(self):
        payload = load_sample()
        payload["baseDate"] = ""
        payload["manual_confirmation_items"] = ["其它事项"]
        self.assertFails(payload, r"本次未指定评估基准日")

    def test_base_date_shape(self):
        payload = load_sample()
        payload["baseDate"] = "2026/06/30"
        self.assertFails(payload, r"YYYY-MM-DD")

    def test_zero_is_not_a_missing_number(self):
        payload = load_sample()
        payload["unitPrice"] = 0
        self.assertFails(payload, r"不许填 0")

    def test_both_amounts_null_is_not_a_conclusion(self):
        payload = load_sample()
        payload["unitPrice"] = None
        payload["totalValue"] = None
        self.assertFails(payload, r"不算一个估值结论")

    def test_missing_source_is_not_a_knowledge_source(self):
        payload = load_sample()
        payload["knowledge_sources"][0]["source"] = "missing"
        self.assertFails(payload, r"不许出现在来源清单里")

    def test_knowledge_source_needs_a_digest(self):
        payload = load_sample()
        payload["knowledge_sources"][0]["digest"] = ""
        self.assertFails(payload, r"缺 digest")

    def test_listed_price_cannot_be_a_direct_comparable(self):
        payload = load_sample()
        listed = next(item for item in payload["cases"] if item["transaction_type"] == "挂牌")
        listed["used_as_comparable"] = True
        self.assertFails(payload, r"不得直接当成交价使用")

    def test_case_needs_a_human_decision(self):
        payload = load_sample()
        payload["cases"][0]["keep_reason"] = ""
        payload["cases"][0]["drop_reason"] = ""
        self.assertFails(payload, r"人工判断必须留痕")

    def test_case_needs_url_and_capture_time(self):
        for field in ("url", "captured_at"):
            payload = load_sample()
            payload["cases"][0][field] = ""
            self.assertFails(payload, rf"缺 {field}")

    # ── 中介联系方式（§7.5）────────────────────────────────────────────────

    def test_contact_state_must_be_one_of_the_five(self):
        payload = load_sample()
        payload["cases"][0]["contact"]["state"] = "有电话"
        self.assertFails(payload, r"contact.state 非法")

    def test_contact_must_be_present_even_when_there_is_none(self):
        # 拿不到也要说明是**哪种**拿不到：留一个空对象分不出"页面没有"和"忘了写"。
        payload = load_sample()
        del payload["cases"][0]["contact"]
        self.assertFails(payload, r"缺 contact")

    def test_public_contact_needs_a_value_and_a_time(self):
        payload = load_sample()
        payload["cases"][0]["contact"] = {"state": "public", "captured_at": "2026-10-13 10:00"}
        self.assertFails(payload, r"公开获取必须有取值")
        payload = load_sample()
        payload["cases"][0]["contact"] = {"state": "public", "agency": "某经纪"}
        self.assertFails(payload, r"缺 captured_at")

    def test_masked_number_must_be_kept_as_is_not_guessed(self):
        # 页面展示的就是脱敏形态，原样保留；标成 public 说明标错了状态。
        payload = load_sample()
        payload["cases"][0]["contact"] = {
            "state": "public", "agency": "某经纪", "captured_at": "2026-10-13 10:00",
            "phone": "138****1234",
        }
        self.assertFails(payload, r"带脱敏符却标成 public")
        # 标成 masked 却没有把脱敏号码留下来 —— 一样不合格。
        payload = load_sample()
        payload["cases"][0]["contact"] = {"state": "masked", "phone": ""}
        self.assertFails(payload, r"没有把页面展示的脱敏号码原样留下来")

    def test_login_required_and_failed_need_a_reason(self):
        for state in ("login-required", "failed"):
            payload = load_sample()
            payload["cases"][0]["contact"] = {"state": state}
            self.assertFails(payload, rf"contact 状态是 {state}")

    def test_contact_states_other_than_public_still_pass(self):
        # 没有公开联系方式**不**自动排除案例（除非知识库明确要求）。
        for state in ("not-published", "masked", "login-required", "failed"):
            payload = load_sample()
            extra = {"detail": "页面没有展示联系方式"} if state in ("login-required", "failed") else {}
            if state == "masked":
                extra = {"phone": "138****1234"}
            payload["cases"][0]["contact"] = {"state": state, **extra}
            self.write(payload)
            completed = run(VALIDATE, self.result)
            self.assertEqual(completed.returncode, 0, f"{state} 应当通过：\n{completed.stdout}")

    # ── 计价面积与计价单位（§9 第 1 条）────────────────────────────────────

    def test_price_unit_must_be_the_canonical_one(self):
        payload = load_sample()
        payload["priceUnit"] = "元/套内㎡"
        self.assertFails(payload, r"必须是「元/㎡」")
        payload = load_sample()
        payload["priceUnit"] = ""
        self.assertFails(payload, r"priceUnit 为空")

    def test_unit_price_needs_a_pricing_area(self):
        # 没有计价面积，总价与单价没法勾稽 —— 那正是"口径不一致"的来源。
        payload = load_sample()
        payload["pricingArea"] = None
        self.assertFails(payload, r"没有 pricingArea")

    def test_unit_price_times_area_must_equal_total(self):
        payload = load_sample()
        payload["totalValue"] = 1
        self.assertFails(payload, r"单价 × 计价面积与总价对不上")

    def test_rounding_tolerance_is_one_yuan(self):
        # 舍入差 1 元以内算一致（底稿按元取整），超过就是口径不一致。
        payload = load_sample()
        payload["totalValue"] = payload["unitPrice"] * payload["pricingArea"] + 1
        self.write(payload)
        assert run(VALIDATE, self.result).returncode == 0, "1 元舍入差应当通过"
        payload["totalValue"] = payload["unitPrice"] * payload["pricingArea"] + 2
        self.assertFails(payload, r"对不上")

    def test_zero_pricing_area_is_not_a_missing_number(self):
        payload = load_sample()
        payload["pricingArea"] = 0
        self.assertFails(payload, r"pricingArea 必须为正数")

    # ── 规则检查记录与案例状态（§8）──────────────────────────────────────────

    def test_rule_ids_and_status_must_match_the_calculation_module(self):
        # 契约里的规则 ID / 状态 / 阈值与 `market_rules.py` 是**同一套词**：
        # 两处各写一份的话，计算模块说 R-SINGLE-20 超限、结果里却记成别的东西，谁也发现不了。
        import market_rules as rules  # noqa: E402
        self.assertEqual(set(validate.RULE_IDS), set(rules.RULE_IDS))
        self.assertEqual(set(validate.CASE_STATUS), set(rules.CASE_STATUS))
        self.assertEqual(validate.RULE_LIMITS, rules.RULE_LIMITS)
        self.assertEqual(validate.STATUS_RULE, rules.STATUS_RULE)

    def test_case_needs_rule_checks(self):
        payload = load_sample()
        del payload["cases"][1]["rule_checks"]
        self.assertFails(payload, r"缺 rule_checks")

    def test_rule_check_needs_id_verdict_and_reason(self):
        payload = load_sample()
        payload["cases"][1]["rule_checks"][2]["rule_id"] = "R-随便写的"
        self.assertFails(payload, r"rule_id 非法")
        payload = load_sample()
        payload["cases"][1]["rule_checks"][2]["verdict"] = "maybe"
        self.assertFails(payload, r"verdict 非法")
        payload = load_sample()
        payload["cases"][1]["rule_checks"][2]["reason"] = ""
        self.assertFails(payload, r"缺 reason")

    def test_threshold_must_equal_the_knowledge_value(self):
        # 记录里写的阈值必须逐字等于知识库口径 —— 不许自行放宽。
        for rule_id, wrong in (("R-SINGLE-20", 0.25), ("R-COMBINED-30", 0.5), ("R-GROUP-1.2", 1.3)):
            payload = load_sample()
            for item in payload["cases"]:
                for check in item["rule_checks"]:
                    if check["rule_id"] == rule_id:
                        check["threshold"] = wrong
            self.assertFails(payload, r"知识库口径是 0\.2：阈值不许自行放宽|知识库口径是 .*：阈值不许自行放宽")

    def test_numeric_verdict_needs_an_actual_value(self):
        payload = load_sample()
        for item in payload["cases"]:
            for check in item["rule_checks"]:
                if check["rule_id"] == "R-SINGLE-20":
                    check["actual"] = None
        self.assertFails(payload, r"就必须给出 actual")

    def test_dropped_case_needs_a_reason_not_a_one_liner(self):
        # 状态说"单项规则未通过"，却没有那条规则的记录 → 红。
        payload = load_sample()
        payload["cases"][1]["rule_checks"] = [
            check for check in payload["cases"][1]["rule_checks"] if check["rule_id"] != "R-SINGLE-20"
        ]
        self.assertFails(payload, r"却没有 R-SINGLE-20 的规则检查记录")
        # 状态说未通过，那条规则却判了 pass → 红（状态与规则结论矛盾）。
        payload = load_sample()
        for check in payload["cases"][1]["rule_checks"]:
            if check["rule_id"] == "R-SINGLE-20":
                check["verdict"] = "pass"
        self.assertFails(payload, r"状态与规则结论矛盾")

    def test_status_and_used_as_comparable_must_agree(self):
        payload = load_sample()
        payload["cases"][1]["status"] = "adopted"
        self.assertFails(payload, r"两处必须一致")
        payload = load_sample()
        payload["cases"][0]["status"] = "not-comparable"
        self.assertFails(payload, r"两处必须一致")

    def test_status_enum(self):
        payload = load_sample()
        payload["cases"][1]["status"] = "被剔除"
        self.assertFails(payload, r"status 非法")

    def test_group_ratio_must_be_recomputable(self):
        # 报出来的组内比值必须能从采用案例的比准价重算出来（写一个好看的比值挡得住）。
        payload = load_sample()
        second = json.loads(json.dumps(payload["cases"][0]))
        second["url"] = "https://example.invalid/2"
        payload["cases"].append(second)
        # 真实比值 = max/min = 1.0（两条同价）；把它写成 1.05 就该判红。
        for item in payload["cases"]:
            for check in item["rule_checks"]:
                if check["rule_id"] == "R-GROUP-1.2":
                    check["actual"] = 1.05
        self.assertFails(payload, r"按采用案例比准价算出来")

    def test_group_over_12_needs_an_explanation(self):
        # 两条采用案例，比值 1.5 > 1.2：必须在 manual_confirmation_items 里说明处置方式。
        payload = load_sample()
        # 把"待确认事项"清空：这一条测的正是"超限却**没有**说明"要被拦住。
        payload["manual_confirmation_items"] = []
        second = json.loads(json.dumps(payload["cases"][0]))
        second["url"] = "https://example.invalid/2"
        second["unit_price"] = payload["cases"][0]["unit_price"] * 1.5
        payload["cases"].append(second)
        for item in payload["cases"]:
            for check in item["rule_checks"]:
                if check["rule_id"] == "R-GROUP-1.2":
                    check["actual"] = 1.5
                    check["verdict"] = "fail"
                    check["reason"] = "组内比值 1.5 超过 1.2"
        self.assertFails(payload, r"超过 1.2，却没有任何说明")
        # 写了说明就通过（"说明原因后保留"是规范允许的处置）。
        for item in payload["cases"]:
            for check in item["rule_checks"]:
                if check["rule_id"] == "R-GROUP-1.2":
                    check["verdict"] = "fail"
        payload["manual_confirmation_items"] = [
            "本次未指定评估基准日",
            "组内比值 1.5 超过 1.2：已说明处置方式并保留，需复核",
        ]
        self.write(payload)
        self.assertEqual(run(VALIDATE, self.result).returncode, 0, run(VALIDATE, self.result).stdout)

    def test_workbook_must_carry_the_adopted_cases_evidence(self):
        """案例证据要同时落进底稿的「案例截图 / 证据留档」行（§8「截图编号」/ §9 三者一致）。

        两处各写一份而没人比过，就会出现"报告说留了证据、底稿上却是空的"。
        """
        payload = load_sample()
        self.write(payload)                      # 夹具已经把证据填进底稿
        self.assertEqual(run(VALIDATE, self.result, "--workbook", self.workbook).returncode, 0)
        # 把底稿那一行清掉 → 红。
        import openpyxl  # type: ignore
        workbook = openpyxl.load_workbook(self.workbook)
        sheet = workbook["案例"]
        for row in sheet.iter_rows(min_col=1, max_col=1):
            value = row[0].value
            if isinstance(value, str) and value.strip().startswith("案例截图 / 证据留档"):
                for column in range(2, 7):
                    # ⚠️ 必须显式赋值：`cell(row, column, value=None)` 在 openpyxl 里是"读"，
                    # 传 None 不会把单元格清空（写这条时踩过一次）。
                    sheet.cell(row=row[0].row, column=column).value = None
                break
        workbook.save(self.workbook)
        workbook.close()
        completed = run(VALIDATE, self.result, "--workbook", self.workbook)
        self.assertEqual(completed.returncode, 1, "底稿没有证据留档时必须判红")
        self.assertRegex(completed.stdout, r"没有写进底稿的「案例截图 / 证据留档」行")

    def test_missing_evidence_row_is_an_error_not_a_silent_skip(self):
        # 案例表里那一行被整行删掉 → 也要红（不是"找不到就跳过"）。
        payload = load_sample()
        self.write(payload)
        import openpyxl  # type: ignore
        workbook = openpyxl.load_workbook(self.workbook)
        sheet = workbook["案例"]
        for row in sheet.iter_rows(min_col=1, max_col=1):
            value = row[0].value
            if isinstance(value, str) and value.strip().startswith("案例截图 / 证据留档"):
                sheet.cell(row=row[0].row, column=1, value="（这一行被删掉了）")
                break
        workbook.save(self.workbook)
        workbook.close()
        completed = run(VALIDATE, self.result, "--workbook", self.workbook)
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"找不到「案例截图 / 证据留档」这一行")

    def test_template_must_keep_the_summary_basis_rows(self):
        """§8：底稿要留出「权重 / 汇总依据」「已知限制」「待确认事项」三行。

        这三行**存在**才谈得上填；删掉它们等于"最终单价怎么汇总的"与
        "这份底稿到什么程度为止"都没地方写。
        """
        import openpyxl  # type: ignore
        workbook = openpyxl.load_workbook(TEMPLATE)
        try:
            labels = {str(row[0].value).strip()
                      for row in workbook["估值结果汇总"].iter_rows(min_col=1, max_col=1)
                      if isinstance(row[0].value, str)}
        finally:
            workbook.close()
        for label in ("权重 / 汇总依据", "已知限制", "待确认事项"):
            self.assertIn(label, labels, f"模板的汇总表缺少「{label}」行")

    def test_confirmation_items_must_reach_the_workbook(self):
        """报告上写了限制、底稿上看不出来 —— 这是最容易被忽略的一类不一致。"""
        payload = load_sample()
        self.write(payload)
        self.assertEqual(run(VALIDATE, self.result, "--workbook", self.workbook).returncode, 0,
                         run(VALIDATE, self.result, "--workbook", self.workbook).stdout)
        # 把底稿那一格的待确认事项清掉 → 红。
        import openpyxl  # type: ignore
        workbook = openpyxl.load_workbook(self.workbook)
        sheet = workbook["估值结果汇总"]
        for row in sheet.iter_rows(min_col=1, max_col=1):
            value = row[0].value
            if isinstance(value, str) and value.strip() == "待确认事项":
                sheet.cell(row=row[0].row, column=2).value = None
                break
        workbook.save(self.workbook)
        workbook.close()
        completed = run(VALIDATE, self.result, "--workbook", self.workbook)
        self.assertEqual(completed.returncode, 1, "底稿没有待确认事项时必须判红")
        self.assertRegex(completed.stdout, r"没有写进底稿的「待确认事项」行")

    # ── 页面截图（§7.4）────────────────────────────────────────────────────

    def test_missing_screenshot_must_be_explicit(self):
        # 空截图 + 没标 → 红（分不出"没截"与"忘了赋值"）。
        payload = load_sample()
        del payload["cases"][0]["screenshot_missing"]
        self.assertFails(payload, r"缺 screenshot_missing")
        payload = load_sample()
        payload["cases"][0]["screenshot_missing"] = False
        self.assertFails(payload, r"也不是 true")

    def test_screenshot_file_must_really_exist(self):
        payload = load_sample()
        case = payload["cases"][0]
        case["screenshot"] = "来源证据/截图-fang-1.png"
        case["screenshot_digest"] = "sha256:abc"
        case["screenshot_missing"] = False
        self.assertFails(payload, r"截图文件不存在")

    def test_screenshot_requires_a_digest(self):
        payload = load_sample()
        case = payload["cases"][1]              # 这一条不是采用案例，避免与"采用案例必须有证据"串味
        case["screenshot"] = "来源证据/截图-fang-1.png"
        case["screenshot_digest"] = ""
        case["screenshot_missing"] = False
        self.write(payload)
        os.makedirs(os.path.join(self.dir, "来源证据"), exist_ok=True)
        with open(os.path.join(self.dir, "来源证据", "截图-fang-1.png"), "wb") as handle:
            handle.write(b"png")
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"却没有 screenshot_digest")

    def test_no_screenshot_at_all_must_be_declared_as_a_gap(self):
        """采用案例一张截图都没有时，`capability_gaps` 里必须提到截图。

        不写的话，结果看起来像"证据齐备"，而实际上采用案例一张截图都没有。
        """
        payload = load_sample()
        payload["capability_gaps"] = ["没有底稿重算能力"]
        self.assertFails(payload, r"capability_gaps 里也没有提到截图")
        # 写出来就通过（本部署确实没有截图能力）。
        payload = load_sample()
        payload["capability_gaps"] = ["本部署没有可用的页面截图能力：证据是页面字节与指纹，不是截图"]
        self.write(payload)
        self.assertEqual(run(VALIDATE, self.result).returncode, 0, run(VALIDATE, self.result).stdout)

    def test_screenshot_path_must_stay_inside_the_case_dir(self):
        payload = load_sample()
        case = payload["cases"][0]
        case["screenshot"] = "../截图.png"
        case["screenshot_digest"] = "sha256:abc"
        case["screenshot_missing"] = False
        self.assertFails(payload, r"截图路径越界")

    def test_adopted_case_cannot_carry_a_failed_rule(self):
        payload = load_sample()
        for check in payload["cases"][0]["rule_checks"]:
            if check["rule_id"] == "R-SINGLE-20":
                check["verdict"] = "fail"
                check["reason"] = "单项修正超限"
        self.assertFails(payload, r"采用案例不能带着没通过的规则")

    # ── 方法判断记录（§7.2）────────────────────────────────────────────────

    def test_method_screening_needs_exactly_one_adopted(self):
        payload = load_sample()
        payload["methodScreening"].append(
            {"method": "income", "verdict": "adopted", "reason": "也算一下"})
        self.assertFails(payload, r"adopted 的方法必须\*\*恰好一个\*\*")

    def test_method_screening_must_agree_with_the_conclusion(self):
        payload = load_sample()
        payload["methodScreening"][0]["method"] = "income"
        payload["methodScreening"][1]["method"] = "market"
        self.assertFails(payload, r"两处必须一致")

    def test_every_screened_method_needs_a_reason(self):
        payload = load_sample()
        payload["methodScreening"][1]["reason"] = ""
        self.assertFails(payload, r"缺 reason：每个方法都要写清采用或不采用的依据")

    def test_screening_is_required(self):
        payload = load_sample()
        del payload["methodScreening"]
        self.assertFails(payload, r"缺 methodScreening")

    def test_no_conflict_must_be_explicit_null(self):
        # 没有冲突时要显式写 null —— 省略分不出"没冲突"与"忘了写"。
        payload = load_sample()
        del payload["methodConflict"]
        self.assertFails(payload, r"缺 methodConflict")

    def test_user_specified_method_cannot_be_switched_silently(self):
        """§7.2：用户指定了方法、结论却换了另一个 —— **必须**写出冲突与原因。"""
        payload = load_sample()
        self.write(payload)
        # 冻结输入里用户指定了收益法，结论却是市场法，而 methodConflict 是 null。
        with open(os.path.join(self.dir, "估值输入.json"), "w", encoding="utf-8") as handle:
            json.dump({"method": "income", "baseDate": "", "supplementalPrompt": "", "digest": "sha256:x"}, handle)
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"不许静默切换")

        # 写出冲突与原因 → 通过。
        payload["methodConflict"] = {
            "userMethod": "income", "adoptedMethod": "market",
            "reason": "对象无租赁合同与租金记录，收益法缺净收益序列；按 references/02 改按市场法测算，已登记待确认",
        }
        payload["manual_confirmation_items"] = list(payload["manual_confirmation_items"]) + [
            "用户指定收益法，因缺租金数据改按市场法：请复核"
        ]
        self.write(payload)
        with open(os.path.join(self.dir, "估值输入.json"), "w", encoding="utf-8") as handle:
            json.dump({"method": "income", "baseDate": "", "supplementalPrompt": "", "digest": "sha256:x"}, handle)
        self.assertEqual(run(VALIDATE, self.result).returncode, 0, run(VALIDATE, self.result).stdout)

    def test_conflict_must_name_the_user_method(self):
        payload = load_sample()
        payload["methodConflict"] = {
            "userMethod": "cost", "adoptedMethod": "market", "reason": "资料不足",
        }
        self.write(payload)
        with open(os.path.join(self.dir, "估值输入.json"), "w", encoding="utf-8") as handle:
            json.dump({"method": "income", "baseDate": "", "supplementalPrompt": "", "digest": "sha256:x"}, handle)
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"冻结输入里用户指定的是 income")

    def test_no_conflict_when_the_user_specified_auto(self):
        payload = load_sample()
        self.write(payload)
        with open(os.path.join(self.dir, "估值输入.json"), "w", encoding="utf-8") as handle:
            json.dump({"method": "auto", "baseDate": "", "supplementalPrompt": "", "digest": "sha256:x"}, handle)
        self.assertEqual(run(VALIDATE, self.result).returncode, 0, run(VALIDATE, self.result).stdout)
        # 自动判定时给出一个"冲突"也不对（没有用户指定的方法，哪来的冲突）。
        payload["methodConflict"] = {
            "userMethod": "income", "adoptedMethod": "market", "reason": "随便写的",
        }
        self.write(payload)
        with open(os.path.join(self.dir, "估值输入.json"), "w", encoding="utf-8") as handle:
            json.dump({"method": "auto", "baseDate": "", "supplementalPrompt": "", "digest": "sha256:x"}, handle)
        self.assertFails(payload, r"没有指定方法")

    def test_conflict_is_null_when_the_user_method_was_adopted(self):
        payload = load_sample()
        payload["methodConflict"] = {
            "userMethod": "market", "adoptedMethod": "market", "reason": "多余的冲突记录",
        }
        self.write(payload)
        with open(os.path.join(self.dir, "估值输入.json"), "w", encoding="utf-8") as handle:
            json.dump({"method": "market", "baseDate": "", "supplementalPrompt": "", "digest": "sha256:x"}, handle)
        self.assertFails(payload, r"methodConflict 必须是 null")

    def test_conflict_reason_is_required(self):
        payload = load_sample()
        payload["methodConflict"] = {"userMethod": "income", "adoptedMethod": "market", "reason": ""}
        self.write(payload)
        with open(os.path.join(self.dir, "估值输入.json"), "w", encoding="utf-8") as handle:
            json.dump({"method": "income", "baseDate": "", "supplementalPrompt": "", "digest": "sha256:x"}, handle)
        self.assertFails(payload, r"methodConflict 缺 reason")

    def test_market_rules_only_apply_to_the_market_method(self):
        """§8：20% / 30% / 1.2 是**比较案例**的调整规则，只在市场法里执行。

        套到收益法 / 成本法上是真实会发生的错误：那两条方法里没有"可比实例修正系数积"这一层，
        判定看着像做了校验，实际是拿错口径在自证。
        """
        payload = load_sample()
        payload["method"] = "income"
        # 方法记录也跟着改，避免被"方法与筛选记录不一致"先拦住（这一条测的是另一件事）。
        payload["methodScreening"][0]["method"] = "income"
        self.assertFails(payload, r"只在市场法里执行")

        # 收益法只留条件/证据类判定 → 通过。
        payload = load_sample()
        payload["method"] = "income"
        payload["methodScreening"][0]["method"] = "income"
        for item in payload["cases"]:
            item["rule_checks"] = [check for check in item["rule_checks"]
                                   if check["rule_id"] not in ("R-SINGLE-20", "R-COMBINED-30", "R-GROUP-1.2")]
            item["status"] = "adopted" if item["used_as_comparable"] else "insufficient-evidence"
            if item["status"] == "insufficient-evidence":
                item["rule_checks"].append({"rule_id": "R-EVIDENCE", "actual": None, "threshold": None,
                                            "verdict": "fail", "reason": "没有留存页面证据"})
        self.write(payload)
        self.assertEqual(run(VALIDATE, self.result).returncode, 0, run(VALIDATE, self.result).stdout)

    def test_artifact_must_exist(self):
        payload = load_sample()
        payload["artifacts"]["xlsx"] = "不存在的底稿.xlsx"
        self.assertFails(payload, r"指向的文件不存在")

    def test_artifact_must_stay_inside_the_case_dir(self):
        payload = load_sample()
        payload["artifacts"]["html"] = "../跑出去的.html"
        self.assertFails(payload, r"案例目录内的相对路径")

    def test_html_must_agree_with_the_json(self):
        payload = load_sample()
        self.write(payload)
        with open(self.html, "w", encoding="utf-8") as handle:
            handle.write("<html><body>报告里没有数字</body></html>")
        completed = run(VALIDATE, self.result, "--html", self.html)
        self.assertEqual(completed.returncode, 1, completed.stdout)
        self.assertRegex(completed.stdout, r"报告里找不到")

    def test_knowledge_sources_must_come_from_this_runs_snapshot(self):
        """结果里的知识来源必须真的来自本轮快照（§1「可追溯到知识来源」）。"""
        payload = load_sample()
        self.write(payload)
        # 造一份本轮快照：两条来源（一条 dws、一条 local），与样例一致。
        snapshot_dir = os.path.join(self.dir, "知识快照")
        os.makedirs(snapshot_dir, exist_ok=True)
        index = os.path.join(snapshot_dir, ".知识来源.json")

        def write_index(sources):
            with open(index, "w", encoding="utf-8") as handle:
                json.dump({"schema": "crwu.valuation.knowledge.v1", "sources": sources}, handle, ensure_ascii=False)

        reported = payload["knowledge_sources"]
        write_index([dict(item) for item in reported])
        self.assertEqual(run(VALIDATE, self.result).returncode, 0, "来源与快照一致时应当通过")

        # 知识来源的取值从**校验脚本**里读（不从字面量抄）：技能脚本里出现业务 CLI 的字面量
        # 会被 CLI 守卫判成"间接旁路"（踩过两次）。谁是"知识库来源"由校验脚本说了算。
        with open(VALIDATE, encoding="utf-8") as handle:
            knowledge_sources = SchemaSubsetTest.validator_tuples(handle.read())["KNOWLEDGE_SOURCES"]
        authoritative = knowledge_sources[0]

        # ① 结果里出现一条**根本没取到**的路由 → 必须红。
        write_index([dict(item) for item in reported])
        payload2 = load_sample()
        payload2["knowledge_sources"].append(
            {"route": "估值与测算/房地产/99-凭空捏造", "source": authoritative, "updatedAt": "", "digest": "sha256:x"})
        self.write(payload2)
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1, completed.stdout)
        self.assertRegex(completed.stdout, r"不在本轮快照里")

        # ② 指纹不一致（同一路由读到不同版本）→ 必须红。
        self.write(load_sample())
        write_index([dict(item) for item in reported][:1])
        with open(index, encoding="utf-8") as handle:
            index_body = json.load(handle)
        index_body["sources"][0]["digest"] = "sha256:另一版"
        write_index(index_body["sources"])
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1, completed.stdout)
        self.assertRegex(completed.stdout, r"指纹与快照不一致")

        # ③ 快照里有 dws 来源却没写进结果（把降级说成"用了知识库"）→ 必须红。
        self.write(load_sample())
        write_index([dict(item) for item in reported])
        trimmed = load_sample()
        trimmed["knowledge_sources"] = [item for item in reported if item["source"] != authoritative]
        self.write(trimmed)
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1, completed.stdout)
        self.assertRegex(completed.stdout, r"没有写进结果")

    def test_missing_knowledge_snapshot_is_a_warning_not_a_failure(self):
        """老案例没有快照索引时不报错（协议 32 之前的产物）—— 只提醒跳过核对。"""
        payload = load_sample()
        self.write(payload)
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 0, completed.stdout)
        self.assertRegex(completed.stdout, r"跳过来源追溯核对")

    def test_every_case_needs_evidence_or_an_explicit_missing_mark(self):
        """§16「案例有 URL、抓取时间和证据」：没有证据必须**显式标记**，且不得当可比实例。"""
        payload = load_sample()
        self.write(payload)

        # ① 空 evidence 且没标记 → 红（分不出"没留存"还是"忘了赋值"）。
        broken = load_sample()
        broken["cases"][0]["evidence"] = ""
        self.assertFails(broken, r"也没标 evidence_missing")

        # ② 显式标记"没留存"、状态与规则记录都改对 → 通过
        #    （候选案例可以留痕，只是不能当可比实例）。
        marked = load_sample()
        case = marked["cases"][0]
        case["evidence"] = ""
        case["evidence_missing"] = True
        case["used_as_comparable"] = False
        case["status"] = "insufficient-evidence"
        for check in case["rule_checks"]:
            if check["rule_id"] == "R-EVIDENCE":
                check["verdict"] = "fail"
                check["reason"] = "页面证据没有留存：只能作为候选留痕，不能作为结论依据"
        self.write(marked)
        self.assertEqual(run(VALIDATE, self.result).returncode, 0, run(VALIDATE, self.result).stdout)

        # ③ 没证据却当可比实例 → 红。
        misused = load_sample()
        misused["cases"][0]["evidence"] = ""
        misused["cases"][0]["evidence_missing"] = True
        misused["cases"][0]["used_as_comparable"] = True
        self.assertFails(misused, r"没有证据却被标为可直接作可比实例")

        # ④ 状态说 adopted、规则记录却说 R-EVIDENCE fail → 红（两处必须一致）。
        inconsistent = load_sample()
        for check in inconsistent["cases"][0]["rule_checks"]:
            if check["rule_id"] == "R-EVIDENCE":
                check["verdict"] = "fail"
        self.assertFails(inconsistent, r"采用案例不能带着没通过的规则")

    def test_evidence_file_must_really_exist(self):
        """引用了证据文件但文件不在 → 红（证据丢了却说有，比没有证据更危险）。"""
        payload = load_sample()
        self.write(payload)
        referenced = os.path.join(self.dir, payload["cases"][0]["evidence"])
        os.remove(referenced)
        completed = run(VALIDATE, self.result)
        self.assertEqual(completed.returncode, 1, completed.stdout)
        self.assertRegex(completed.stdout, r"引用的证据文件不存在")

        # 越界的相对路径（`..`）一律不合格：证据必须在本次案例目录之内。
        payload2 = load_sample()
        payload2["cases"][0]["evidence"] = "../../etc/passwd"
        self.assertFails(payload2, r"越出案例目录")

    def test_workbook_method_must_agree_with_the_json(self):
        """JSON 说用什么方法，底稿汇总表就要写哪个 —— 两个交付件的口径必须一致。"""
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        payload = load_sample()
        self.write(payload)

        # ① 汇总表写了一行面积，方法却与 JSON（market）不一致 → 必须红。
        workbook = openpyxl.load_workbook(self.workbook)
        workbook["估值结果汇总"]["F5"] = 120.5
        workbook["估值结果汇总"]["J5"] = "收益法"
        workbook.save(self.workbook)
        workbook.close()
        completed = run(VALIDATE, self.result, "--workbook", self.workbook)
        self.assertEqual(completed.returncode, 1, completed.stdout)
        self.assertRegex(completed.stdout, r"口径必须一致")

        # ② 写成 JSON 那一份 → 通过（证明上一条不是因为别的原因红的）。
        workbook = openpyxl.load_workbook(self.workbook)
        workbook["估值结果汇总"]["J5"] = "市场法"
        workbook.save(self.workbook)
        workbook.close()
        self.assertEqual(run(VALIDATE, self.result, "--workbook", self.workbook).returncode, 0)

    def test_workbook_without_data_rows_is_not_checked_for_method(self):
        """还没填数据行的底稿不参与方法一致性检查（模板本身就是这种形态）。"""
        payload = load_sample()
        self.write(payload)
        completed = run(VALIDATE, self.result, "--workbook", self.workbook)
        self.assertEqual(completed.returncode, 0, completed.stdout)

    def test_workbook_must_not_hide_sheets_or_carry_errors(self):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        payload = load_sample()
        self.write(payload)
        workbook = openpyxl.load_workbook(self.workbook)
        workbook["案例"].sheet_state = "hidden"
        workbook["市场法"]["A1"] = "#REF!"
        workbook.save(self.workbook)
        completed = run(VALIDATE, self.result, "--workbook", self.workbook)
        self.assertEqual(completed.returncode, 1, completed.stdout)
        self.assertRegex(completed.stdout, r"隐藏表")
        self.assertRegex(completed.stdout, r"错误值")


class WorkbookContractTest(unittest.TestCase):
    def test_shipped_template_passes_the_contract(self):
        completed = run(CONTRACT, TEMPLATE)
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        self.assertIn("OK:", completed.stdout)

    def test_default_path_points_at_template_not_assets(self):
        # 默认路径必须指向随包发布的模板；指错了会让"没有参数"的调用静默校验别的文件。
        completed = run(CONTRACT)
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        self.assertIn("valuation-workbook.xlsx", completed.stdout)

    def test_hidden_sheet_and_dangling_formula_are_caught(self):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        with tempfile.TemporaryDirectory(prefix="crwu-workbook-") as tmp:
            broken = os.path.join(tmp, "broken.xlsx")
            shutil.copyfile(TEMPLATE, broken)
            workbook = openpyxl.load_workbook(broken)
            workbook["收益法"].sheet_state = "hidden"
            workbook["市场法"]["B5"] = "=#REF!"
            workbook.save(broken)
            completed = run(CONTRACT, broken)
            self.assertEqual(completed.returncode, 1, completed.stdout)
            self.assertRegex(completed.stdout + completed.stdout, r"隐藏|REF")

    def test_summary_cannot_omit_the_adopted_method(self):
        """汇总表：填了建筑面积就必须写明「采用的方法」——否则 K 列静默回空串。"""
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        with tempfile.TemporaryDirectory(prefix="crwu-workbook-method-") as tmp:
            # ① 填了面积、没写方法 → 必须红（K 列的 SEARCH 全落空，汇总无声地没有结论）。
            missing = os.path.join(tmp, "missing.xlsx")
            shutil.copyfile(TEMPLATE, missing)
            workbook = openpyxl.load_workbook(missing)
            workbook["估值结果汇总"]["F5"] = 120.5
            workbook.save(missing)
            workbook.close()
            completed = run(CONTRACT, missing)
            self.assertEqual(completed.returncode, 1, completed.stdout)
            self.assertRegex(completed.stdout, r"没写「采用的方法」")
            self.assertRegex(completed.stdout, r"静默回空串")

            # ② 写了一个**不是逐字**的方法名 → 也要红（`SEARCH("市场法")` 匹配不到"市场"）。
            wrong = os.path.join(tmp, "wrong.xlsx")
            shutil.copyfile(TEMPLATE, wrong)
            workbook = openpyxl.load_workbook(wrong)
            workbook["估值结果汇总"]["F5"] = 120.5
            workbook["估值结果汇总"]["J5"] = "市场"
            workbook.save(wrong)
            workbook.close()
            completed = run(CONTRACT, wrong)
            self.assertEqual(completed.returncode, 1, completed.stdout)
            self.assertRegex(completed.stdout, r"必须逐字写成")

            # ③ 逐字写对 → 通过（这条同时证明前两条不是因为别的原因红的）。
            good = os.path.join(tmp, "good.xlsx")
            shutil.copyfile(TEMPLATE, good)
            workbook = openpyxl.load_workbook(good)
            workbook["估值结果汇总"]["F5"] = 120.5
            workbook["估值结果汇总"]["J5"] = "收益法"
            workbook.save(good)
            workbook.close()
            self.assertEqual(run(CONTRACT, good).returncode, 0)

    def test_total_row_is_not_a_data_row(self):
        """随包模板自己必须过这条新规：合计行的 F 是 `=SUM(...)`，不是填进去的建筑面积。"""
        # 第一版没排掉合计行，模板自己被判不合格 —— 这条断言就是那个教训。
        self.assertEqual(run(CONTRACT, TEMPLATE).returncode, 0)

    def test_missing_workbook_is_an_error_not_a_pass(self):
        completed = run(CONTRACT, os.path.join(HERE, "并不存在.xlsx"))
        self.assertEqual(completed.returncode, 1)
        self.assertIn("找不到工作簿", completed.stdout)


class BuildWorkbookTest(unittest.TestCase):
    """填底稿：只填输入区、保护公式区、空就是空。"""

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="crwu-build-workbook-")
        self.values = os.path.join(self.dir, "values.json")
        self.out = os.path.join(self.dir, "估值底稿.xlsx")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write(self, cells: dict) -> None:
        with open(self.values, "w", encoding="utf-8") as handle:
            json.dump({"cells": cells}, handle, ensure_ascii=False)

    def run_build(self, *extra: str) -> subprocess.CompletedProcess:
        return run(BUILD_WORKBOOK, "--values", self.values, "--out", self.out, *extra)

    def test_fills_only_input_cells_and_keeps_formulas(self):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        self.write({"基础信息表!B5": "某投资有限公司", "基础信息表!B12": 0.09, "收益法!C9": 4.5})
        completed = self.run_build()
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        workbook = openpyxl.load_workbook(self.out)
        self.assertEqual(workbook["基础信息表"]["B5"].value, "某投资有限公司")
        self.assertEqual(workbook["基础信息表"]["B12"].value, 0.09)
        # 公式一列都不能被动过。
        self.assertEqual(workbook["市场法"]["J6"].value, "=C6")
        self.assertEqual(workbook["收益法"]["C11"].value, "=ROUND(C9*365,2)")
        workbook.close()

    def test_refuses_to_overwrite_a_formula_cell(self):
        self.write({"市场法!J6": 12345})
        completed = self.run_build()
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"公式区")
        self.assertFalse(os.path.exists(self.out), "被拒时不许留下半填的底稿")

    def test_refuses_a_cell_outside_the_input_area(self):
        self.write({"基础信息表!A1": "改标题"})
        completed = self.run_build()
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"不在\*\*输入区\*\*")

    def test_refuses_a_formula_in_the_input_area(self):
        # 输入区写公式 = 把结论硬编码进输入，正是契约要防的那件事。
        self.write({"基础信息表!B5": "=1+1"})
        completed = self.run_build()
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"不许写公式")

    def test_blank_value_is_written_as_empty_not_zero(self):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        self.write({"基础信息表!B5": "", "案例!B6": None, "基础信息表!B12": 0})
        completed = self.run_build()
        self.assertEqual(completed.returncode, 0, completed.stdout)
        workbook = openpyxl.load_workbook(self.out)
        self.assertIsNone(workbook["基础信息表"]["B5"].value, "空串要写成空单元格")
        self.assertIsNone(workbook["案例"]["B6"].value, "null 要写成空单元格")
        self.assertEqual(workbook["基础信息表"]["B12"].value, 0, "显式的 0 是结论，照写")
        workbook.close()

    def test_refuses_unknown_sheet_and_range_and_bad_reference(self):
        for reference in ["不存在的表!B5", "基础信息表!B5:C6", "基础信息表B5"]:
            self.write({reference: "x"})
            completed = self.run_build()
            self.assertEqual(completed.returncode, 1, reference)
            self.assertRegex(completed.stdout, r"FAIL")

    def test_uses_the_template_it_was_given_not_the_shipped_one(self):
        """§7.1：用**用户实际模板**，不能悄悄换成随包模板。

        做法：拿一份**改过**的模板（改一个表头文字），确认产物里带着那处改动 ——
        如果脚本偷偷用了随包模板，这处改动就不会出现。
        """
        custom = os.path.join(self.dir, "用户模板.xlsx")
        shutil.copyfile(TEMPLATE, custom)
        import openpyxl  # type: ignore
        workbook = openpyxl.load_workbook(custom)
        workbook["基础信息表"]["F5"].value = "改过的价值类型口径"
        workbook.save(custom)
        workbook.close()

        self.write({"基础信息表!B5": "某单位"})
        completed = run(BUILD_WORKBOOK, "--values", self.values, "--out", self.out, "--template", custom)
        self.assertEqual(completed.returncode, 0, completed.stdout)
        self.assertIn("用户/知识库模板", completed.stdout, "要说明用的是哪一份模板")
        self.assertRegex(completed.stdout, r"sha256:[0-9a-f]{16}")

        produced = openpyxl.load_workbook(self.out)
        try:
            self.assertEqual(produced["基础信息表"]["F5"].value, "改过的价值类型口径",
                             "产物必须来自用户给的那份模板")
        finally:
            produced.close()

    def test_rejects_a_lookalike_template(self):
        """§7.1：**不能**用"外观类似但结构不同"的替代模板。结构不对要在动手前就拦住。"""
        lookalike = os.path.join(self.dir, "仿制模板.xlsx")
        shutil.copyfile(TEMPLATE, lookalike)
        import openpyxl  # type: ignore
        workbook = openpyxl.load_workbook(lookalike)
        workbook["市场法"]["A1"] = "市场法比较"          # 表名不变
        del workbook["成本法"]                            # 少一张表
        workbook.save(lookalike)
        workbook.close()

        self.write({"基础信息表!B5": "某单位"})
        completed = run(BUILD_WORKBOOK, "--values", self.values, "--out", self.out, "--template", lookalike)
        self.assertEqual(completed.returncode, 1, "结构不对的模板必须在动手前拦住")
        self.assertRegex(completed.stdout, r"模板不符合工作簿契约")
        self.assertFalse(os.path.exists(self.out), "拒绝时不许留下半成品")

    def test_missing_template_is_an_error_not_a_silent_default(self):
        self.write({"基础信息表!B5": "某单位"})
        completed = run(BUILD_WORKBOOK, "--values", self.values, "--out", self.out,
                        "--template", os.path.join(self.dir, "不存在.xlsx"))
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"找不到模板")

    def test_refuses_to_write_back_into_the_template_and_to_clobber_by_default(self):
        self.write({"基础信息表!B5": "x"})
        into_template = run(BUILD_WORKBOOK, "--values", self.values, "--out", TEMPLATE)
        self.assertEqual(into_template.returncode, 1)
        self.assertRegex(into_template.stdout, r"模板只读|模板本身")
        # 第一次成功；第二次不加 --force 必须拒绝（覆盖上一版前要有人确认过）。
        self.assertEqual(self.run_build().returncode, 0)
        again = self.run_build()
        self.assertEqual(again.returncode, 1)
        self.assertRegex(again.stdout, r"目标已存在")
        self.assertEqual(self.run_build("--force").returncode, 0)

    def test_blank_values_are_counted_as_blank_not_as_filled(self):
        """`None` 是合法输入（"这一项未知"），但它并没有"填入"任何值 —— 计数也是一种断言。

        复核的人正是照着"填入 N 个输入单元格"这句话判断"该填的都填了"；把置空的也算进
        "填入"，那句话就是假的。
        """
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("本机没有 openpyxl")
        self.write({"基础信息表!B5": "某公司", "基础信息表!B12": None})
        completed = self.run_build()
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        self.assertRegex(completed.stdout, r"填入 1 个输入单元格")
        self.assertRegex(completed.stdout, r"1 处按未填置空")
        workbook = openpyxl.load_workbook(self.out)
        self.assertIsNone(workbook["基础信息表"]["B12"].value, "未知就留空，不写 0")
        workbook.close()

    def test_rejects_a_values_file_without_cells(self):
        with open(self.values, "w", encoding="utf-8") as handle:
            handle.write('{"foo": 1}')
        completed = self.run_build()
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"没有 cells")


class BuildReportTest(unittest.TestCase):
    """渲染报告：数字与结果 JSON 一致、外部文本转义、缺口照实写。"""

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="crwu-build-report-")
        self.result = os.path.join(self.dir, "估值结果.json")
        self.out = os.path.join(self.dir, "估值报告.html")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write(self, payload: dict) -> None:
        with open(self.result, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False)
        # 案例目录里也要有被引用的证据文件（见 `check_case_evidence_files`）。
        for item in payload.get("cases") or []:
            evidence = str(item.get("evidence", "")).strip()
            if evidence == "" or os.path.isabs(evidence):
                continue
            target = os.path.join(self.dir, evidence)
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with open(target, "w", encoding="utf-8") as handle:
                handle.write("<html>evidence</html>")

    def run_build(self, *extra: str) -> subprocess.CompletedProcess:
        return run(BUILD_REPORT, "--result", self.result, "--out", self.out, *extra)

    def test_sample_renders_and_passes_the_result_contract(self):
        shutil.copyfile(SAMPLE, self.result)
        completed = self.run_build()
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
        self.assertTrue(os.path.isfile(self.out))
        body = open(self.out, encoding="utf-8").read()
        self.assertNotIn("{{", body, "模板占位符必须全部被替换")
        # 报告与结果 JSON 的数字一致性由 validate_result.py 核对（这里跑一遍真链路）。
        # 样例的 artifacts.xlsx 不存在（它是样例，不带底稿），所以只该报"缺底稿"这一条 ——
        # **不许**报"报告里找不到单价/总价"：那说明渲染与结果 JSON 已经不一致了。
        checked = run(VALIDATE, self.result, "--html", self.out)
        self.assertNotRegex(checked.stdout, r"报告里找不到")
        self.assertRegex(checked.stdout, r"artifacts\.xlsx 指向的文件不存在")

    def test_numbers_appear_verbatim_so_the_consistency_check_passes(self):
        payload = load_sample()
        payload["artifacts"] = {
            "html": "估值报告.html",
            "xlsx": os.path.basename(TEMPLATE).replace("valuation-workbook.xlsx", "估值底稿.xlsx"),
            "json": "估值结果.json",
        }
        self.write(payload)
        # 底稿也放一份（含采用案例的证据留档），让 --workbook 抽查能过。
        workbook_path = os.path.join(self.dir, payload["artifacts"]["xlsx"])
        shutil.copyfile(TEMPLATE, workbook_path)
        fill_case_evidence(workbook_path, payload)
        self.assertEqual(self.run_build().returncode, 0)
        checked = run(VALIDATE, self.result, "--html", self.out, "--workbook", os.path.join(self.dir, payload["artifacts"]["xlsx"]))
        self.assertEqual(checked.returncode, 0, checked.stdout)

    def test_external_text_is_escaped(self):
        payload = load_sample()
        payload["cases"][0]["title"] = '<script>alert(1)</script>'
        payload["cases"][0]["community"] = '"><img src=x onerror=alert(1)>'
        self.write(payload)
        self.assertEqual(self.run_build().returncode, 0)
        body = open(self.out, encoding="utf-8").read()
        self.assertNotIn("<script>alert(1)</script>", body, "案例标题必须转义")
        self.assertNotIn("<img src=x", body, "小区名必须转义")
        self.assertIn("&lt;script&gt;", body)

    def test_rejects_auto_because_it_is_an_input_not_a_conclusion(self):
        payload = load_sample()
        payload["method"] = "auto"
        self.write(payload)
        completed = self.run_build()
        self.assertEqual(completed.returncode, 1)
        self.assertRegex(completed.stdout, r"不是结论方法")

    def test_empty_gaps_and_manual_items_are_stated_not_omitted(self):
        payload = load_sample()
        payload["capability_gaps"] = []
        payload["manual_confirmation_items"] = []
        self.write(payload)
        self.assertEqual(self.run_build().returncode, 0)
        body = open(self.out, encoding="utf-8").read()
        self.assertIn("没有登记能力缺口", body)
        self.assertIn("没有待确认事项", body)

    def test_missing_base_date_renders_the_explicit_sentence(self):
        payload = load_sample()
        payload["baseDate"] = ""
        self.write(payload)
        self.assertEqual(self.run_build().returncode, 0)
        self.assertIn("本次未指定评估基准日", open(self.out, encoding="utf-8").read())

    def test_refuses_to_clobber_and_to_write_into_the_template(self):
        shutil.copyfile(SAMPLE, self.result)
        into_template = run(BUILD_REPORT, "--result", self.result, "--out", REPORT_TEMPLATE)
        self.assertEqual(into_template.returncode, 1)
        self.assertEqual(self.run_build().returncode, 0)
        self.assertEqual(self.run_build().returncode, 1, "默认不许覆盖")
        self.assertEqual(self.run_build("--force").returncode, 0)


class SchemaSubsetTest(unittest.TestCase):
    """**JSON 通过 Schema** 这句话必须真的被执行，而且 schema 不许与校验脚本漂移。

    本机（与 DSH 自带 Python）都**没有** `jsonschema`，所以这里自带一个最小校验器。
    它的关键设计是**遇到不认识的 keyword 就报错**：一个"只认识一半 keyword"的校验器
    会给出**假通过**（schema 里写了 `pattern` 而校验器忽略它，样例照样"通过"）。
    所以支持面是**白名单**，扩展 schema 的人必须同时扩展这里 —— 那正是我们要的提醒。
    """

    # 只支持这些 keyword；出现别的 → 报错（见类注释）。
    SUPPORTED = {
        "$schema", "$id", "$ref", "title", "description",
        "type", "properties", "required", "additionalProperties", "items", "enum", "oneOf",
        "minLength", "minItems",
    }

    def load(self, name: str) -> dict:
        with open(os.path.join(SCHEMA_DIR, name), encoding="utf-8") as handle:
            return json.load(handle)

    def assert_supported(self, schema: object, path: str, seen: list) -> None:
        if isinstance(schema, dict):
            for key, value in schema.items():
                self.assertIn(key, self.SUPPORTED, f"{path} 用了校验器不支持的 keyword `{key}`："
                              "请扩展 SchemaSubsetTest.SUPPORTED 与 validate()，不要让它被静默忽略")
                if key == "properties":
                    for prop, sub in value.items():
                        self.assert_supported(sub, f"{path}.properties.{prop}", seen)
                elif key == "items":
                    self.assert_supported(value, f"{path}.items", seen)
                elif key in ("oneOf",):
                    for index, sub in enumerate(value):
                        self.assert_supported(sub, f"{path}.oneOf[{index}]", seen)
                elif key == "enum":
                    seen.append((path, list(value)))
        elif isinstance(schema, list):
            for index, sub in enumerate(schema):
                self.assert_supported(sub, f"{path}[{index}]", seen)

    def validate(self, instance: object, schema: dict, path: str = "$") -> list:
        """返回问题清单（空 = 通过）。只实现本包 schema 用到的那部分 draft-07。"""
        problems: list[str] = []
        if "oneOf" in schema:
            matches = [sub for sub in schema["oneOf"] if not self.validate(instance, sub, path)]
            if len(matches) != 1:
                problems.append(f"{path}: 不满足 oneOf 中的恰好一支（命中 {len(matches)} 支）")
            return problems
        expected = schema.get("type")
        ok = {
            "object": lambda v: isinstance(v, dict),
            "array": lambda v: isinstance(v, list),
            "string": lambda v: isinstance(v, str),
            "number": lambda v: isinstance(v, (int, float)) and not isinstance(v, bool),
            "integer": lambda v: isinstance(v, int) and not isinstance(v, bool),
            "boolean": lambda v: isinstance(v, bool),
            "null": lambda v: v is None,
        }
        if expected is not None and not ok[expected](instance):
            return [f"{path}: 类型应为 {expected}，实际 {type(instance).__name__}"]
        if "enum" in schema and instance not in schema["enum"]:
            problems.append(f"{path}: {instance!r} 不在枚举 {schema['enum']} 中")
        if isinstance(instance, str) and "minLength" in schema and len(instance) < schema["minLength"]:
            problems.append(f"{path}: 长度 {len(instance)} < minLength {schema['minLength']}")
        if isinstance(instance, list) and "minItems" in schema and len(instance) < schema["minItems"]:
            problems.append(f"{path}: 元素数 {len(instance)} < minItems {schema['minItems']}")
        if isinstance(instance, dict):
            for key in schema.get("required", []):
                if key not in instance:
                    problems.append(f"{path}: 缺必需字段 `{key}`")
            for key, sub in schema.get("properties", {}).items():
                if key in instance:
                    problems.extend(self.validate(instance[key], sub, f"{path}.{key}"))
            if schema.get("additionalProperties") is False:
                extra = [key for key in instance if key not in schema.get("properties", {})]
                if extra:
                    problems.append(f"{path}: 有多余字段 {extra}")
        if isinstance(instance, list) and "items" in schema:
            for index, item in enumerate(instance):
                problems.extend(self.validate(item, schema["items"], f"{path}[{index}]"))
        return problems

    def test_schemas_use_only_supported_keywords(self):
        for name in sorted(os.listdir(SCHEMA_DIR)):
            if not name.endswith(".json"):
                continue
            self.assert_supported(self.load(name), name, [])

    def test_shipped_samples_validate_against_their_schemas(self):
        result_schema = self.load("valuation-result.schema.json")
        case_schema = self.load("valuation-case.schema.json")
        with open(SAMPLE, encoding="utf-8") as handle:
            result = json.load(handle)
        with open(CASE_SAMPLE, encoding="utf-8") as handle:
            case = json.load(handle)
        self.assertEqual(self.validate(result, result_schema), [], "结果样例没通过结果 schema")
        # 案例样例是**单条案例**（`cases` 数组的元素）。
        self.assertEqual(self.validate(case, case_schema), [], "案例样例没通过案例 schema")

    def test_a_broken_sample_is_actually_rejected(self):
        # 负向：校验器必须真的会红，而不是"永远说 OK"。三条都是契约里点名的错法。
        schema = self.load("valuation-result.schema.json")
        with open(SAMPLE, encoding="utf-8") as handle:
            result = json.load(handle)
        for mutate, needle in [
            (lambda payload: payload.update({"method": "auto"}), "枚举"),
            (lambda payload: payload.pop("capability_gaps"), "缺必需字段"),
            (lambda payload: payload.update({"unitPrice": "42000"}), "类型应为"),
        ]:
            broken = json.loads(json.dumps(result))
            mutate(broken)
            self.assertTrue(self.validate(broken, schema), f"这条改动应当被判不合格（{needle}）")

    @staticmethod
    def validator_tuples(source: str) -> dict:
        """把校验脚本里的常量元组解析成 `{名字: [值, …]}`。

        **故意不把那些值写进本文件**：技能脚本里出现业务 CLI 的字面量会被 CLI 守卫
        判成"间接旁路"（`exec-cli-literal`）。所以这里从**校验脚本**里读值、再拿去比对 schema ——
        这既是防漂移，也说明"谁拥有这些取值"：真正拥有它们的是校验脚本。
        """
        found = {}
        for name, body in re.findall(r"^(\w+) = \(([^)]*)\)", source, re.M):
            found[name] = [item.strip().strip('"').strip("'") for item in body.split(",") if item.strip()]
        return found

    def test_schema_does_not_drift_from_the_validator(self):
        # 两个事实源必须一致：schema（给模型看的结构）与 validate_result.py（真正的门）。
        # 漂移的后果是"照 schema 写的 JSON 过不了校验"或者反过来 —— 两种都让人不再信任 schema。
        with open(VALIDATE, encoding="utf-8") as handle:
            validator = handle.read()
        tuples = self.validator_tuples(validator)
        result_schema = self.load("valuation-result.schema.json")
        case_schema = self.load("valuation-case.schema.json")
        for key in result_schema["required"]:
            self.assertIn(f'"{key}"', validator, f"schema 的必需字段 `{key}` 不在校验脚本的必需列表里")
        self.assertEqual(result_schema["properties"]["method"]["enum"], tuples["CONCLUSION_METHODS"])
        self.assertEqual(
            case_schema["properties"]["transaction_type"]["enum"], tuples["TRANSACTION_TYPES"])
        # ⚠️ 案例的 `source` 是**来源平台**，不是知识来源 —— 两个枚举长得像但不是一回事
        #（第一版这条断言就写错了，被自己刚写的用例当场抓出来）。
        self.assertEqual(
            result_schema["properties"]["knowledge_sources"]["items"]["properties"]["source"]["enum"],
            tuples["KNOWLEDGE_SOURCES"])
        self.assertEqual(len(case_schema["properties"]["source"]["enum"]), 3)

    def test_object_schema_names_the_twelve_fields_the_drawer_shows(self):
        # 抽屉里的 12 个要素是用户核对输入的那份清单。对象 schema 的 `fields` 是开放映射
        #（标签锚定提取：读不到就不出现），所以这里核对的是**那 12 个标签被逐字列在 schema 里**
        #（`fieldLabels` 枚举）—— 否则"资料里读到的信息"会有两套口径：界面一套、对象文件一套。
        schema = self.load("valuation-object.schema.json")
        labels = schema["properties"]["fieldLabels"]["items"]["enum"]
        for key in ["坐落", "房屋性质", "用途", "建筑面积", "楼层", "总楼层", "建筑结构",
                    "建成年代", "权属信息", "租赁信息", "价格信息", "已知案例线索"]:
            self.assertIn(key, labels, f"对象 schema 的 fieldLabels 缺 `{key}`")


if __name__ == "__main__":
    unittest.main(verbosity=2)

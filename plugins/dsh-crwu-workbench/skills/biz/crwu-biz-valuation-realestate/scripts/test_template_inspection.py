"""Read-only template inspection against synthetic workbooks."""

import hashlib
import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


class TemplateInspectionTest(unittest.TestCase):
    def inspector(self):
        script = Path(__file__).with_name("inspect_template.py")
        self.assertTrue(script.exists(), "a read-only template inspector must exist")
        spec = importlib.util.spec_from_file_location("case_template_inspector", script)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        self.assertTrue(callable(getattr(module, "inspect_template", None)))
        return module.inspect_template

    def workbook(self, root, *, label="建筑面积", extra=False):
        path = Path(root) / ("second.xlsx" if extra else "first.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.title = "候选"
        sheet.merge_cells("A1:E1")
        sheet["A1"] = "案例登记"
        sheet["A2"] = label
        sheet["A3"] = "楼层"
        sheet["E5"] = "=SUM(B2:D2)"
        book.defined_names.add(DefinedName("SyntheticArea", attr_text="'候选'!$B$2"))
        if extra:
            other = book.create_sheet("权益")
            other["A1"] = "剩余年限"
        book.save(path)
        book.close()
        return path

    def inspect(self, path, **kwargs):
        return self.inspector()(str(path), "template-fixture", "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest(), **kwargs)

    def test_actual_structure_labels_formulas_and_names_are_read_without_writing(self):
        with tempfile.TemporaryDirectory() as root:
            path = self.workbook(root)
            before = path.read_bytes()
            result = self.inspect(path)
            self.assertTrue(result["ok"], result)
            profile = result["profile"]
            self.assertTrue(profile["coverage"]["complete"])
            self.assertEqual(profile["sheets"][0]["name"], "候选")
            self.assertIn("A1:E1", profile["sheets"][0]["mergedRanges"])
            cells = {(cell["sheet"], cell["cell"]): cell for cell in profile["cells"]}
            self.assertEqual(cells[("候选", "A2")]["value"], "建筑面积")
            self.assertEqual(cells[("候选", "E5")]["formula"], "=SUM(B2:D2)")
            self.assertTrue(any(item["name"] == "SyntheticArea" for item in profile["definedNames"]))
            self.assertEqual(path.read_bytes(), before)

    def test_different_templates_change_the_profile_without_a_fixed_seven_sheet_contract(self):
        with tempfile.TemporaryDirectory() as root:
            first = self.inspect(self.workbook(root))
            second = self.inspect(self.workbook(root, label="可出租面积", extra=True))
            self.assertNotEqual(first["profile"]["digest"], second["profile"]["digest"])
            self.assertEqual(len(second["profile"]["sheets"]), 2)
            self.assertTrue(any(cell["value"] == "可出租面积" for cell in second["profile"]["cells"]))

    def test_replaced_or_invalid_bytes_fail_without_claiming_a_profile(self):
        inspect = self.inspector()
        with tempfile.TemporaryDirectory() as root:
            path = self.workbook(root)
            changed = inspect(str(path), "template-fixture", "sha256:" + "0" * 64)
            self.assertFalse(changed["ok"])
            self.assertNotIn("profile", changed)
            path.write_bytes(b"not an xlsx")
            invalid = self.inspect(path)
            self.assertFalse(invalid["ok"])
            self.assertNotIn("profile", invalid)
            self.assertNotIn(root, str(invalid))

    def test_bounded_coverage_is_visible_instead_of_silently_dropping_fields(self):
        with tempfile.TemporaryDirectory() as root:
            result = self.inspect(self.workbook(root), max_cells=2)
            self.assertTrue(result["ok"], result)
            self.assertFalse(result["profile"]["coverage"]["complete"])
            self.assertTrue(result["profile"]["coverage"]["issues"])

    def test_case_root_mismatch_is_rejected_before_inspecting_foreign_workbooks(self):
        with tempfile.TemporaryDirectory() as root, tempfile.TemporaryDirectory() as other:
            path = self.workbook(other)
            digest = "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()
            result = self.inspector()(str(path), "template-fixture", digest, case_root=root)
            self.assertFalse(result["ok"])
            self.assertNotIn("profile", result)
            self.assertNotIn(other, str(result))


if __name__ == "__main__":
    unittest.main()

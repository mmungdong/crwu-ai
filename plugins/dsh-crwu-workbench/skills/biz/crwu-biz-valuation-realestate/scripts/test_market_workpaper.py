"""Portable workbook behavior tests using synthetic project files only."""
import hashlib
import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


class MarketWorkpaperTests(unittest.TestCase):
    def setUp(self):
        spec = importlib.util.spec_from_file_location("market_workpaper", Path(__file__).with_name("build_market_workpaper.py"))
        self.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.module)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_candidates_keep_missing_values_links_and_units(self):
        data = {"candidates": [{"caseId": "case-1", "providerId": "lianjia", "sourceUrl": "https://bj.lianjia.com/ershoufang/1.html",
            "listFacts": {"areaSqm": {"value": 81, "unit": "㎡", "status": "verified", "evidenceRefs": []}},
            "detailFacts": {}, "status": "pending", "selectionState": "unselected", "similarities": ["同小区"],
            "differences": [], "missingFields": ["unitPrice"], "recommendation": "建议核实", "sourceRefs": []}]}
        result = self.module.build(self.root, "candidates", data)
        path = self.root / result["artifact"]["relativePath"]
        sheet = load_workbook(path)["候选案例"]
        headers = {c.value: c.column for c in sheet[1]}
        self.assertEqual(sheet.cell(2, headers["案例ID"]).value, "case-1")
        self.assertEqual(sheet.cell(2, headers["网页地址"]).hyperlink.target, data["candidates"][0]["sourceUrl"])
        self.assertIsNone(sheet.cell(2, headers["列表单价（元/㎡）"]).value)
        self.assertEqual(result["artifact"]["digest"], "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest())

    def test_fill_uses_actual_mapping_preserves_style_formula_and_original(self):
        for target in ["B2", "D4"]:
            with self.subTest(target=target):
                book = Workbook(); sheet = book.active; sheet.title = "测算"
                sheet["A1"] = "原表"; sheet["F5"] = "=#REF!"; sheet.merge_cells("B6:C6")
                sheet[target].fill = PatternFill("solid", fgColor="FFAA00")
                template = self.root / "template.xlsx"; book.save(template)
                digest = "sha256:" + hashlib.sha256(template.read_bytes()).hexdigest()
                data = {"template": {"relativePath": "template.xlsx", "digest": digest}, "adoptedCaseIds": [],
                    "fillTrace": [{"fieldId": "areaSqm", "caseId": None, "sheet": "测算", "cell": target,
                        "fact": {"value": 80, "unit": "㎡", "status": "verified", "evidenceRefs": []}}],
                    "limitations": ["原模板公式未核"], "candidates": []}
                result = self.module.build(self.root, "fill", data)
                out = load_workbook(self.root / result["artifacts"][0]["relativePath"])
                self.assertEqual(out["测算"][target].value, 80)
                self.assertEqual(out["测算"][target].fill.fgColor.rgb, "00FFAA00")
                self.assertEqual(out["测算"]["F5"].value, "=#REF!")
                self.assertIn("B6:C6", str(out["测算"].merged_cells))
                self.assertEqual(hashlib.sha256(template.read_bytes()).hexdigest(), digest[7:])
                self.assertEqual(result["recalculation"], "not-run")
                self.assertIsNone(result["calculation"])

    def test_unknown_assumption_formula_target_and_foreign_paths(self):
        book = Workbook(); book.active["B2"] = "=A1"; book.save(self.root / "template.xlsx")
        digest = "sha256:" + hashlib.sha256((self.root / "template.xlsx").read_bytes()).hexdigest()
        data = {"template": {"relativePath": "template.xlsx", "digest": digest}, "adoptedCaseIds": [], "candidates": [], "limitations": [],
            "fillTrace": [{"fieldId": "x", "caseId": None, "sheet": "Sheet", "cell": "B2", "fact": {"value": 1, "unit": "", "status": "verified", "evidenceRefs": []}}]}
        with self.assertRaises(ValueError):
            self.module.build(self.root, "fill", data)
        data["fillTrace"][0]["cell"] = "C2"; data["fillTrace"][0]["fact"]["status"] = "missing"
        result = self.module.build(self.root, "fill", data)
        self.assertIsNone(load_workbook(self.root / result["artifacts"][0]["relativePath"])["Sheet"]["C2"].value)
        data["template"]["relativePath"] = "../template.xlsx"
        with self.assertRaises(ValueError):
            self.module.build(self.root, "fill", data)

    def test_delivery_rejects_changed_adoption_and_trace(self):
        book = Workbook(); book.save(self.root / "template.xlsx")
        data = {"template": {"relativePath": "template.xlsx", "digest": self.module.digest(self.root / "template.xlsx")},
            "adoptedCaseIds": [], "fillTrace": [], "candidates": [], "limitations": []}
        result = self.module.build(self.root, "fill", data)
        with self.assertRaises(ValueError):
            self.module.build(self.root, "verify", {**data, **result, "adoptedCaseIds": ["unselected-case"]})
        trace = next(a for a in result["artifacts"] if a["kind"] == "trace")
        path = self.root / trace["relativePath"]
        path.write_text('{}', encoding="utf-8")
        trace["digest"] = self.module.digest(path); trace["bytes"] = path.stat().st_size
        with self.assertRaises(ValueError):
            self.module.build(self.root, "verify", {**data, **result})

    def test_unknown_targets_clear_template_example_values(self):
        book = Workbook(); book.active["B2"] = 100; book.save(self.root / "template.xlsx")
        result = self.module.build(self.root, "fill", {
            "template": {"relativePath": "template.xlsx", "digest": self.module.digest(self.root / "template.xlsx")},
            "adoptedCaseIds": [], "fillTrace": [], "candidates": [], "limitations": [],
            "clearRefs": [{"sheet": "Sheet", "cell": "B2"}]})
        self.assertIsNone(load_workbook(self.root / result["artifacts"][0]["relativePath"])["Sheet"]["B2"].value)


if __name__ == "__main__":
    unittest.main()

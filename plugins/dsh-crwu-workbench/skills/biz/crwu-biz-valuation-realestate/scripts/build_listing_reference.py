#!/usr/bin/env python3
"""One-call 挂牌参考 delivery: amounts + workbook + report data + HTML.

    python3 scripts/build_listing_reference.py \
        --input <plan>/挂牌参考输入.json --out-dir <plan> --case-root <basePlanDir>
    python3 scripts/build_listing_reference.py --validate --out-dir <plan> --case-root <basePlanDir>
    python3 scripts/build_listing_reference.py --prepare-template --out-dir <plan> --case-root <basePlanDir>

Only the four visible deliverables are generated; a hidden `.listing-reference/` directory keeps
the receipt, a canonical input snapshot and per-file digests. `--validate` is read-only: after the
four SHA checks it recomputes the calculation, report data, trace, HTML and workbook from the
current fixed input and compares the whole stored result (hashes alone are not proof). A build only
reuses existing output after that validation, refuses to overwrite unreceipted or changed input,
and rejects any output/meta symlink that escapes the output directory. `--prepare-template` writes
the genuine blank `模板/挂牌参考计算表.xlsx` once, keeps its receipt/digest for reuse, and never
overwrites a foreign or tampered template; a build that pins `payload.template` reads that exact
approved file and verifies its digest before generating any deliverable.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from decimal import Decimal
from urllib.parse import urlsplit

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import build_report as report_renderer  # noqa: E402
import build_report_data  # noqa: E402
import calculate_market  # noqa: E402
import evidence_index  # noqa: E402

OUTPUTS = ("计算轨迹.json", "估值底稿.xlsx", "估值报告数据.json", "估值报告.html")
META_DIR = ".listing-reference"
RECEIPT_NAME = "receipt.json"
RECEIPT_SCHEMA = "crwu.valuation.listing-reference-receipt.v1"
INPUT_NAME = "挂牌参考输入.json"
# The HTML report template; the blank workbook template is prepared on demand below.
TEMPLATE = os.path.join(SKILL_DIR, "template", "valuation-report.html")
SCHEMA = os.path.join(SKILL_DIR, "schemas", "valuation-report-data.schema.json")
REFERENCE_LABEL = "挂牌参考估值"
# The closed verified-workflow snapshot ref produced by the fixed listing payload.
KNOWLEDGE_REF_FORMAT = "crwu.valuation.knowledge-ref.v1"
UNKNOWN_VALUE_SCOPE_NOTICE = "本次未确认房地价值范围，金额仅为挂牌报价参考"

WORKBOOK_TEMPLATE_DIR = "模板"
WORKBOOK_TEMPLATE_NAME = "挂牌参考计算表.xlsx"
WORKBOOK_TEMPLATE_RELATIVE = f"{WORKBOOK_TEMPLATE_DIR}/{WORKBOOK_TEMPLATE_NAME}"
TEMPLATE_RECEIPT_DIR = ".listing-reference-template"
TEMPLATE_RECEIPT_NAME = "template.json"
TEMPLATE_RECEIPT_SCHEMA = "crwu.valuation.listing-reference-template.v1"
WORKBOOK_TEMPLATE_SHEETS = ["估值结果", "挂牌来源", "计算轨迹", "限制",
                            evidence_index.EVIDENCE_SHEET]
TEMPLATE_HEADERS = {
    "估值结果": (4, ["对象", "参考单价(CNY/m2)", "参考总价(CNY)", "币种", "性质"]),
    "挂牌来源": (1, ["caseId", "挂牌单价", "抓取时间", "来源URL", "适用理由", "知识引用"]),
    "计算轨迹": (2, ["对象", "选中案例", "原始单价", "最大值", "最小值", "未取整均值",
                     "建筑面积", "总价=均值×面积", "取整/口径"]),
}
# Where the workbook keeps the approved business values; the blank template declares the first
# object row, which is also where a deliverable fills its first object.
TEMPLATE_MAPPING = [
    {"field": "objectId", "sheet": "估值结果", "range": "A5"},
    {"field": "unitPrice", "sheet": "估值结果", "range": "B5"},
    {"field": "totalPrice", "sheet": "估值结果", "range": "C5"},
    {"field": "currency", "sheet": "估值结果", "range": "D5"},
    {"field": "nature", "sheet": "估值结果", "range": "E5"},
    {"field": "referenceCaseIds", "sheet": "计算轨迹", "range": "B3"},
    {"field": "rawUnitPrices", "sheet": "计算轨迹", "range": "C3"},
    {"field": "unitPriceMax", "sheet": "计算轨迹", "range": "D3"},
    {"field": "unitPriceMin", "sheet": "计算轨迹", "range": "E3"},
    {"field": "unitPriceMean", "sheet": "计算轨迹", "range": "F3"},
    {"field": "areaSqm", "sheet": "计算轨迹", "range": "G3"},
    {"field": "totalPriceByMean", "sheet": "计算轨迹", "range": "H3"},
    {"field": "roundingNote", "sheet": "计算轨迹", "range": "I3"},
]
BLANK_TEMPLATE_LIMIT_BYTES = 5 * 1024 * 1024
MAX_TEMPLATE_CELLS = 20000
_FORMULA_ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")


class ListingReferenceError(ValueError):
    """The one-call delivery cannot be produced (unsafe path / changed input / tamper)."""


# ---------------------------------------------------------------------------
# Paths / digests
# ---------------------------------------------------------------------------

def _real(path: str) -> str:
    return os.path.realpath(os.path.abspath(path))


def _within(root: str, target: str) -> bool:
    try:
        return target == root or os.path.commonpath([root, target]) == root
    except ValueError:
        return False


def _sha256_bytes(body: bytes) -> str:
    return "sha256:" + hashlib.sha256(body).hexdigest()


def _sha256_file(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(65536), b""):
            digest.update(block)
    return "sha256:" + digest.hexdigest()


def _canonical(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _input_digest(payload: object) -> str:
    return _sha256_bytes(_canonical(payload).encode("utf-8"))


def _now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _frozen_generated_at(payload: dict, receipt: dict | None = None) -> str:
    """Resolve the report generation time once: receipt first, then explicit input, then now().

    A missing `generatedAt` must not be re-read from the clock on validate, otherwise the
    whole report data drifts by one second and validate/reuse fail for an unchanged input.
    """
    if receipt is not None:
        stored = receipt.get("generatedAt")
        if isinstance(stored, str) and stored.strip() != "":
            return stored
    given = str(payload.get("generatedAt") or "").strip()
    return given if given != "" else _now_iso()


def _effective_payload(payload: dict, generated_at: str) -> dict:
    """Shallow copy with the frozen generation time; the original digest stays untouched."""
    effective = dict(payload)
    effective["generatedAt"] = generated_at
    return effective


def _require_containment(case_root: str, out_dir: str, input_path: str) -> tuple:
    """Realpath containment: input/output stay under the base plan directory."""
    case_real = _real(case_root)
    if not os.path.isdir(case_real):
        raise ListingReferenceError(f"案例根目录不存在：{case_root}")
    out_real = _real(out_dir)
    if not _within(case_real, out_real):
        raise ListingReferenceError("输出目录必须落在案例根目录（basePlanDir）之内")
    if out_real == os.path.dirname(case_real):
        raise ListingReferenceError("输出目录不得是案例根的父目录")
    input_real = _real(input_path) if input_path else ""
    if input_real and not _within(case_real, input_real):
        raise ListingReferenceError("输入文件必须落在案例根目录之内")
    if input_real and input_real == out_real:
        raise ListingReferenceError("输入文件不能同时是输出目录")
    return case_real, out_real, input_real


def _read_input(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as handle:
        return json.load(handle)


# ---------------------------------------------------------------------------
# Output containment + receipt
# ---------------------------------------------------------------------------

def _meta_dir(out_real: str, name: str = META_DIR) -> str:
    """The hidden meta directory, proven to be a real child of `out_real`."""
    meta = os.path.join(out_real, name)
    if os.path.islink(meta):
        raise ListingReferenceError("隐藏元数据目录不得是符号链接（可能指向输出目录之外）")
    real = _real(meta)
    if not _within(out_real, real) or os.path.dirname(real) != out_real:
        raise ListingReferenceError("隐藏元数据目录的真实路径必须直接落在输出目录之内")
    return meta


def _receipt_path(out_real: str, name: str = RECEIPT_NAME,
                  meta_name: str = META_DIR) -> str:
    meta = _meta_dir(out_real, meta_name)
    path = os.path.join(meta, name)
    if os.path.islink(path):
        raise ListingReferenceError("收据文件不得是符号链接")
    real = _real(path)
    if not _within(out_real, real) or os.path.dirname(real) != meta:
        raise ListingReferenceError("收据的真实路径必须直接落在隐藏元数据目录内")
    return path


def _output_path(out_real: str, name: str) -> str:
    """Reject symlinked or escaping output paths before any read/write."""
    target = os.path.join(out_real, name)
    if os.path.islink(target):
        raise ListingReferenceError(f"产物 {name} 不得是符号链接")
    real = _real(target)
    if not _within(out_real, real) or os.path.dirname(real) != out_real:
        raise ListingReferenceError(f"产物 {name} 的真实路径必须直接落在输出目录内")
    return target


# ---------------------------------------------------------------------------
# Approved workbook template
# ---------------------------------------------------------------------------

def _used_cells(workbook) -> list:
    """Non-empty cells only, so a user's blank-looking area is not mistaken for content."""
    cells = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells.append(f"{sheet.title}!{cell.coordinate}")
    return cells


def _template_cells_within(case_real: str, target: str) -> bool:
    """Reject a declared template path that escapes the base plan or crosses a symlink."""
    real = _real(target)
    if not _within(case_real, real):
        return False
    current = case_real
    for segment in os.path.relpath(real, case_real).split(os.sep):
        if segment in ("", "."):
            continue
        current = os.path.join(current, segment)
        if os.path.islink(current):
            return False
    return True


def _template_target_path(case_real: str, relative: str) -> str:
    if not isinstance(relative, str) or relative.strip() == "":
        raise ListingReferenceError("模板 relativePath 不能为空")
    normalized = relative.replace("\\", "/")
    if os.path.isabs(normalized) or re.match(r"^[A-Za-z]:", normalized):
        raise ListingReferenceError("模板 relativePath 必须是案例根目录内的相对路径")
    segments = [segment for segment in normalized.split("/") if segment not in ("", ".")]
    if not segments or any(segment == ".." for segment in segments):
        raise ListingReferenceError("模板 relativePath 不许包含 ..（不能借道离开案例根目录）")
    target = os.path.join(case_real, *segments)
    if not _template_cells_within(case_real, target):
        raise ListingReferenceError("模板路径的真实位置必须直接落在案例根目录之内")
    return target


def _template_receipt_meta(out_real: str) -> str:
    return _meta_dir(out_real, TEMPLATE_RECEIPT_DIR)


def _template_receipt_path(out_real: str) -> str:
    return _receipt_path(out_real, TEMPLATE_RECEIPT_NAME, TEMPLATE_RECEIPT_DIR)


def _template_structure_checks(workbook, cells: list | None = None) -> list:
    """Recheck the required sheets/headers/mapping target and reject formulas and bad links."""
    sheets = workbook.sheetnames
    headers_ok = sheets == WORKBOOK_TEMPLATE_SHEETS
    if headers_ok:
        for name, (row, expected) in TEMPLATE_HEADERS.items():
            sheet = workbook[name]
            actual = [sheet.cell(row=row, column=column).value
                      for column in range(1, len(expected) + 1)]
            if actual != expected:
                headers_ok = False
                break
    mapping_ok = sheets == WORKBOOK_TEMPLATE_SHEETS
    if mapping_ok:
        for item in TEMPLATE_MAPPING:
            if item["sheet"] not in sheets:
                mapping_ok = False
                break
            if workbook[item["sheet"]][item["range"]].data_type == "f":
                mapping_ok = False
                break
    external_ok = True
    error_ok = True
    formula_ok = True
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_ok = False
                value = cell.value
                if isinstance(value, str):
                    if value.strip() in _FORMULA_ERRORS:
                        error_ok = False
                    if any(marker in value for marker in ("#REF!", "#VALUE!", "#NAME?")):
                        error_ok = False
                if cell.hyperlink is not None:
                    target = str(cell.hyperlink.target or "")
                    if target.startswith("file:") or target.startswith("external:") \
                            or "\x00" in target:
                        external_ok = False
    if cells is not None and len(cells) > MAX_TEMPLATE_CELLS:
        external_ok = False
    return [
        {"name": "required_sheets", "ok": sheets == WORKBOOK_TEMPLATE_SHEETS},
        {"name": "required_headers", "ok": headers_ok},
        {"name": "mapping_targets", "ok": mapping_ok},
        {"name": "no_formulas", "ok": formula_ok and mapping_ok},
        {"name": "no_external_links", "ok": external_ok},
        {"name": "no_error_values", "ok": error_ok},
    ]


def _load_workbook(path: str):
    from openpyxl import load_workbook

    return load_workbook(path, data_only=False, keep_links=False)


def _verify_template_file(path: str, digest: str | None = None) -> dict:
    """Read the actual bytes, verify the approved digest and bounded size, then the structure."""
    if os.path.islink(path) or not os.path.isfile(path):
        raise ListingReferenceError(f"模板不是案例目录内的普通文件：{path}")
    size = os.path.getsize(path)
    if size > BLANK_TEMPLATE_LIMIT_BYTES:
        raise ListingReferenceError(
            f"模板文件过大（{size} 字节 > {BLANK_TEMPLATE_LIMIT_BYTES}），拒绝作为已验真模板")
    actual = _sha256_file(path)
    if digest is not None and actual != digest:
        raise ListingReferenceError(
            "模板指纹与本次批准的模板不一致：拒绝用被改动/未批准的模板生成交付件")
    try:
        workbook = _load_workbook(path)
    except Exception as error:  # noqa: BLE001 - an unreadable template is a rejection
        raise ListingReferenceError(f"模板无法按工作簿读取：{error}") from error
    try:
        cells = _used_cells(workbook)
        checks = _template_structure_checks(workbook, cells)
        if not all(check["ok"] for check in checks):
            failed = "、".join(check["name"] for check in checks if not check["ok"])
            raise ListingReferenceError(f"模板结构与批准的计算表契约不一致：{failed}")
    finally:
        workbook.close()
    return {"digest": actual, "size": size, "cells": cells, "checks": checks}


def _payload_template(payload: dict) -> dict | None:
    """The approved template the caller pinned; absent means the legacy no-template build."""
    declared = payload.get("template")
    if declared is None:
        return None
    if not isinstance(declared, dict):
        raise ListingReferenceError("payload.template 必须是对象")
    relative = str(declared.get("relativePath") or "").strip()
    if relative == "":
        raise ListingReferenceError("payload.template.relativePath 是必填项")
    digest = str(declared.get("digest") or "").strip()
    if not re.match(r"^sha256:[0-9a-f]{64}$", digest):
        raise ListingReferenceError("payload.template.digest 必须是 sha256: 加 64 位十六进制")
    return {"relativePath": relative, "digest": digest}


def _verify_payload_template(case_real: str, declared: dict) -> str:
    """Fail before any deliverable write; return the approved template path."""
    path = _template_target_path(case_real, declared["relativePath"])
    _verify_template_file(path, declared["digest"])
    return path


def _read_receipt(path: str) -> dict:
    """Load a receipt only when its schema and output key set are exactly right."""
    try:
        receipt = _read_input(path)
    except (OSError, json.JSONDecodeError) as error:
        raise ListingReferenceError(f"收据不是合法 JSON：{error}") from error
    if not isinstance(receipt, dict):
        raise ListingReferenceError("收据顶层必须是对象")
    if receipt.get("schemaVersion") != RECEIPT_SCHEMA:
        raise ListingReferenceError("收据 schemaVersion 不匹配：拒绝按未知格式校验")
    files = receipt.get("files")
    if not isinstance(files, dict) or set(files) != set(OUTPUTS):
        raise ListingReferenceError("收据 files 必须且只能登记四个产物（不得缺项/多项/任意路径）")
    for name, digest in files.items():
        if not isinstance(digest, str) or not digest.startswith("sha256:"):
            raise ListingReferenceError(f"收据 files[{name}] 不是 sha256 指纹")
    if not isinstance(receipt.get("inputDigest"), str) or receipt["inputDigest"] == "":
        raise ListingReferenceError("收据缺少 inputDigest")
    if not isinstance(receipt.get("input"), dict):
        raise ListingReferenceError("收据缺少输入快照")
    return receipt


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def _numeric_payload(payload: dict) -> dict:
    keys = ("valuationMode", "planId", "referencePolicy", "objects", "cases")
    return {key: payload[key] for key in keys if key in payload}


def _is_workflow_knowledge_ref(item: object) -> bool:
    return isinstance(item, dict) and str(item.get("format") or "") == KNOWLEDGE_REF_FORMAT


def _verified_workflow_knowledge_refs(payload: dict) -> list | None:
    """Re-verify closed workflow snapshot refs; `None` means a legacy full-metadata manifest.

    Identity, path, digest and body must all be real, and the body SHA must match the digest
    before a ref may enter the report. Unknown metadata stays absent; nothing is invented.
    """
    raw = payload.get("knowledgeSources")
    if not isinstance(raw, list) or not any(_is_workflow_knowledge_ref(item) for item in raw):
        return None
    refs = []
    for index, item in enumerate(raw):
        where = f"knowledgeSources[{index}]"
        if not _is_workflow_knowledge_ref(item):
            raise ListingReferenceError(f"{where}: workflow 核验快照格式不一致（不得混入旧清单来源）")
        source_id = str(item.get("sourceId") or "").strip()
        path = str(item.get("path") or "").strip()
        digest = str(item.get("digest") or "").strip()
        body = item.get("body")
        if source_id == "" or path == "" or digest == "" \
                or not isinstance(body, str) or body.strip() == "":
            raise ListingReferenceError(f"{where}: 快照来源必须有真实的 sourceId / path / digest / 正文")
        if _sha256_bytes(body.encode("utf-8")) != digest:
            raise ListingReferenceError(f"{where}: 快照正文与其 digest 不一致：拒绝进入报告")
        ref = {
            "format": KNOWLEDGE_REF_FORMAT,
            "sourceId": source_id,
            "path": path,
            "digest": digest,
            "body": body,
        }
        for optional in ("spaceName", "section", "quote"):
            value = str(item.get(optional) or "").strip()
            if value != "":
                ref[optional] = value
        refs.append(ref)
    return refs


def _apply_workflow_knowledge_refs(data: dict, workflow_refs: list | None) -> dict:
    """Attach the verified refs and re-run the same Schema self-check (no bypass)."""
    if workflow_refs is None:
        return data
    data["knowledgeSources"] = workflow_refs
    problems = build_report_data.validate_report_data(data)
    if problems:
        raise ListingReferenceError(
            "报告数据未通过 Schema 校验（workflow 核验快照）：\n- " + "\n- ".join(problems))
    return data


def _report_manifest(payload: dict, calc: dict, evidence_rows: list,
                     knowledge_sources: list) -> dict:
    for required in ("projectId", "planId", "inputDigest"):
        value = payload.get(required)
        if not isinstance(value, str) or value.strip() == "":
            raise ListingReferenceError(f"挂牌参考输入缺必填字段：{required}")
    numeric = {str(obj.get("objectId")): obj for obj in payload.get("objects") or []
               if isinstance(obj, dict) and str(obj.get("objectId") or "").strip() != ""}
    calc_ids = [str(obj["objectId"]) for obj in calc["objects"]]
    if sorted(numeric.keys()) != sorted(calc_ids) or len(numeric) != len(calc_ids):
        raise ListingReferenceError("报告对象必须与计算结果一一对应（不许省略或新增对象）")
    by_id = {row["caseId"]: row for row in evidence_rows}
    sources = {source["caseId"]: source for obj in calc["objects"] for source in obj["sources"]}
    case_ids = list(sources.keys())
    for row in evidence_rows:
        if row["caseId"] not in sources:
            raise ListingReferenceError(f"截图证据指向未选中的挂牌案例：{row['caseId']}")
    for case_id, source in sources.items():
        evidence = by_id.get(case_id)
        if evidence is not None and evidence["sourceUrl"] != source["sourceUrl"]:
            raise ListingReferenceError(f"案例 {case_id} 的截图来源 URL 与挂牌来源不一致")
    comparables = []
    screenshots = []
    missing = []
    for case_id in case_ids:
        source = sources[case_id]
        evidence = by_id.get(case_id)
        comparables.append({
            "caseId": case_id,
            "displayName": (evidence or {}).get("displayName") or case_id,
            "provider": (evidence or {}).get("provider", ""),
            "sourceUrl": source["sourceUrl"], "transactionType": "挂牌",
            "unitPrice": source["unitPrice"], "used": False, "referenceSelected": True,
        })
        if evidence is None:
            missing.append(case_id)
            continue
        screenshots.append({
            "caseId": case_id, "displayName": evidence.get("displayName") or case_id,
            "provider": evidence.get("provider", ""), "sourceUrl": evidence["sourceUrl"],
            "capturedAt": evidence.get("capturedAt", ""),
            "relativePath": evidence["relativePath"], "digest": evidence["digest"],
            "validity": "valid",
        })
    provisional = []
    unknown_value_scope = False
    for obj in calc["objects"]:
        meta = numeric.get(str(obj["objectId"]), {})
        scope = str(meta.get("valueScope") or "").strip()
        if scope == "":
            unknown_value_scope = True
        row = {
            "objectId": obj["objectId"],
            "displayName": meta.get("displayName") or obj["objectId"],
            "method": "market", "unitPrice": obj["unitPrice"], "totalValue": obj["totalValue"],
            "currency": obj["currency"], "calculationRef": f"listing-{obj['objectId']}",
            # Unknown stays unknown: never default to land-and-building.
            "valueScope": scope or None,
        }
        provisional.append(row)
    limitations = []
    if unknown_value_scope:
        limitations.append(UNKNOWN_VALUE_SCOPE_NOTICE)
    for item in payload.get("limitations") or []:
        text = str(item).strip()
        if text:
            limitations.append(text)
    limitations.extend(calc["limitations"])
    if missing:
        limitations.append("以下挂牌案例没有可核对截图证据（screenshot_missing）：" + "、".join(missing))
    limitations.append("本报告为挂牌参考估值：静态独立计算，未做公式重算，不得作为正式结论交付。")
    return {
        "status": "partial",
        "valuationMode": "listing-reference",
        "referenceSummary": {
            "label": REFERENCE_LABEL,
            "aggregation": "arithmetic-mean",
            "basis": str(payload.get("referencePolicy", {}).get("basis") or ""),
            "priceNature": "挂牌",
            "unitPriceUnit": "CNY/m2",
            "caseIds": case_ids,
            "sources": [{"caseId": case_id, "sourceUrl": sources[case_id]["sourceUrl"]} for case_id in case_ids],
        },
        "projectId": payload["projectId"],
        "planId": payload["planId"],
        "inputDigest": payload["inputDigest"],
        "generatedAt": str(payload.get("generatedAt") or ""),
        "methodRequest": str(payload.get("methodRequest") or "market"),
        "method": "market",
        "methods": ["market"],
        "methodResults": [{
            "method": "market", "basis": "挂牌参考（原始报价算术平均）",
            "unitPrice": calc["objects"][0]["unitPrice"], "totalValue": None, "template": "",
        }],
        "baseDate": payload.get("baseDate"),
        "purpose": payload.get("purpose"),
        "valueType": payload.get("valueType"),
        "supplementalPrompt": str(payload.get("supplementalPrompt") or ""),
        "materials": payload.get("materials") or [],
        "knowledgeSources": knowledge_sources,
        "retrievalAttempts": payload.get("retrievalAttempts") or [],
        "artifacts": [],
        "limitations": limitations,
        "stage": "calculation",
        "completedWork": ["挂牌报价采集与筛选", "挂牌参考金额计算"],
        "outstanding": ["正式可比实例与重算底稿（本报告不含）"],
        "blockers": [],
        "lookupAttempts": [],
        "objects": [numeric[case_id] for case_id in calc_ids],
        "provisionalResults": provisional,
        "comparables": comparables,
        "calculationSteps": [{
            "stepId": f"listing-{obj['objectId']}", "label": "挂牌报价算术平均",
            "expression": "mean(raw listing unit prices) × buildingArea",
            "operands": [
                {"value": source["unitPrice"], "unit": "CNY/m2", "provenance": source["sourceUrl"]}
                for source in obj["sources"]
            ] + [{
                "value": str(numeric.get(str(obj["objectId"]), {}).get("buildingArea", "")),
                "unit": "m2", "provenance": "挂牌参考输入.json",
            }],
            "value": obj["totalValue"], "unit": "CNY",
        } for obj in calc["objects"]],
        "evidence": screenshots,
    }


def _literal(cell, value: object) -> None:
    """Metadata is always literal text; a leading '=' must not become a formula."""
    cell.value = "" if value is None else value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _build_workbook(final_path: str, payload: dict, calc: dict,
                    case_root: str, evidence_rows: list, report_limitations: list,
                    template: str | None = None):
    """Build the workbook in memory so validation can compare against the same spec.

    With `template` the approved blank workbook is loaded and filled in place; without it the
    five-sheet workbook is created from scratch exactly as before.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    case_real = _real(case_root)

    def widen(sheet, widths: dict, wrap_columns=()) -> None:
        for column, width in widths.items():
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in sheet.iter_rows():
            for cell in row:
                if cell.column in wrap_columns:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

    if template is None:
        workbook = Workbook()
        result = workbook.active
        result.title = "估值结果"
        result.merge_cells("A1:E1")
        _literal(result["A1"], REFERENCE_LABEL)
        result["A1"].font = Font(bold=True, size=14)
        result.merge_cells("A2:E2")
        _literal(result["A2"], "静态独立计算：原始挂牌报价算术平均，未做公式重算，参照性质，"
                               "不得作为正式结论。")
        result["A2"].font = Font(color="B45309")
        headers = ["对象", "参考单价(CNY/m2)", "参考总价(CNY)", "币种", "性质"]
        for column, label in enumerate(headers, start=1):
            _literal(result.cell(row=4, column=column), label)
            result.cell(row=4, column=column).font = Font(bold=True)
    else:
        workbook = _load_workbook(template)
        result = workbook["估值结果"]
    for index, obj in enumerate(calc["objects"], start=5):
        _literal(result.cell(row=index, column=1), obj["objectId"])
        _literal(result.cell(row=index, column=2), obj["unitPrice"])
        _literal(result.cell(row=index, column=3), obj["totalValue"])
        _literal(result.cell(row=index, column=4), obj["currency"])
        _literal(result.cell(row=index, column=5), "挂牌参考")
        result.row_dimensions[index].height = 18
    widen(result, {1: 22, 2: 22, 3: 22, 4: 10, 5: 14})

    if template is None:
        sources = workbook.create_sheet("挂牌来源")
        source_headers = ["caseId", "挂牌单价", "抓取时间", "来源URL", "适用理由", "知识引用"]
        for column, label in enumerate(source_headers, start=1):
            _literal(sources.cell(row=1, column=column), label)
            sources.cell(row=1, column=column).font = Font(bold=True)
    else:
        sources = workbook["挂牌来源"]
    row = 2
    for obj in calc["objects"]:
        for source in obj["sources"]:
            for column, value in enumerate(
                (source["caseId"], source["unitPrice"], source["capturedAt"],
                 source["sourceUrl"], source["reason"], source["knowledgeRef"]), start=1):
                cell = sources.cell(row=row, column=column)
                _literal(cell, value)
                if column == 4:
                    cell.hyperlink = source["sourceUrl"]
                    cell.style = "Hyperlink"
            row += 1
    widen(sources, {1: 14, 2: 14, 3: 22, 4: 48, 5: 40, 6: 34}, wrap_columns=(4, 5, 6))

    if template is None:
        trace = workbook.create_sheet("计算轨迹")
        _literal(trace["A1"], "挂牌参考计算轨迹（原始报价，未修正）")
        trace["A1"].font = Font(bold=True)
        trace_headers = ["对象", "选中案例", "原始单价", "最大值", "最小值", "未取整均值",
                         "建筑面积", "总价=均值×面积", "取整/口径"]
        for column, label in enumerate(trace_headers, start=1):
            _literal(trace.cell(row=2, column=column), label)
            trace.cell(row=2, column=column).font = Font(bold=True)
    else:
        trace = workbook["计算轨迹"]
    digits = int(payload.get("referencePolicy", {}).get("rounding", {}).get("digits", 2))
    numeric = {str(obj.get("objectId")): obj for obj in payload.get("objects") or []
               if isinstance(obj, dict)}
    for index, obj in enumerate(calc["objects"], start=3):
        prices = [Decimal(str(source["unitPrice"])) for source in obj["sources"]]
        mean = sum(prices, Decimal(0)) / Decimal(len(prices))
        _literal(trace.cell(row=index, column=1), obj["objectId"])
        _literal(trace.cell(row=index, column=2), "、".join(obj["referenceCaseIds"]))
        _literal(trace.cell(row=index, column=3), "、".join(source["unitPrice"] for source in obj["sources"]))
        _literal(trace.cell(row=index, column=4), obj["unitPriceRange"]["max"])
        _literal(trace.cell(row=index, column=5), obj["unitPriceRange"]["min"])
        _literal(trace.cell(row=index, column=6), str(mean))
        _literal(trace.cell(row=index, column=7), str((numeric.get(str(obj["objectId"])) or {}).get("buildingArea", "")))
        _literal(trace.cell(row=index, column=8), obj["totalValue"])
        _literal(trace.cell(row=index, column=9),
                 f'{digits} 位 half-up；basis={payload.get("referencePolicy", {}).get("basis", "")}')
    widen(trace, {1: 18, 2: 16, 3: 20, 4: 12, 5: 12, 6: 16, 7: 12, 8: 18, 9: 48},
          wrap_columns=(3, 9))

    if template is None:
        limits = workbook.create_sheet("限制")
        _literal(limits["A1"], "限制与前提")
        limits["A1"].font = Font(bold=True)
    else:
        limits = workbook["限制"]
    for index, item in enumerate(report_limitations, start=2):
        cell = limits.cell(row=index, column=1)
        _literal(cell, item)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        limits.row_dimensions[index].height = 42
    widen(limits, {1: 110})

    if template is None:
        evidence_sheet = workbook.create_sheet(evidence_index.EVIDENCE_SHEET)
        for column, label in enumerate(evidence_index.SHEET_COLUMNS, start=1):
            _literal(evidence_sheet.cell(row=1, column=column), label)
            evidence_sheet.cell(row=1, column=column).font = Font(bold=True)
    else:
        evidence_sheet = workbook[evidence_index.EVIDENCE_SHEET]
    for index, row_data in enumerate(evidence_rows, start=2):
        values = [row_data["caseId"], row_data["displayName"], row_data["provider"],
                  row_data["sourceUrl"], row_data["relativePath"], row_data["digest"],
                  row_data["capturedAt"], row_data["evidenceType"]]
        for column, value in enumerate(values, start=1):
            cell = evidence_sheet.cell(row=index, column=column)
            _literal(cell, value)
            if column == 4:
                cell.hyperlink = row_data["sourceUrl"]
                cell.style = "Hyperlink"
            if column == 5:
                cell.hyperlink = evidence_index.hyperlink_target(
                    final_path, case_real, row_data["relativePath"])
                cell.style = "Hyperlink"
    widen(evidence_sheet, {1: 12, 2: 18, 3: 14, 4: 44, 5: 48, 6: 20, 7: 22, 8: 18},
          wrap_columns=(4, 5))
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
    return workbook


def _write_workbook(path: str, final_path: str, payload: dict, calc: dict,
                    case_root: str, evidence_rows: list, report_limitations: list,
                    template: str | None = None) -> None:
    workbook = _build_workbook(final_path, payload, calc, case_root,
                               evidence_rows, report_limitations, template)
    try:
        workbook.save(path)
    finally:
        workbook.close()


def _trace_document(payload: dict, calc: dict, input_hash: str) -> dict:
    return {
        "schemaVersion": "crwu.valuation.listing-reference-trace.v1",
        "valuationMode": "listing-reference",
        "inputDigest": input_hash,
        "referencePolicy": payload.get("referencePolicy"),
        "objects": calc["objects"],
        "limitations": calc["limitations"],
        "notice": "挂牌参考：静态独立计算，未做公式重算，不得作为正式结论。",
    }


def _resolve_input_for_validate(out_real: str, input_real: str, receipt: dict) -> dict:
    """Prefer the fixed current input in the output dir; never trust a stale snapshot over it."""
    canonical = os.path.join(out_real, INPUT_NAME)
    if os.path.isfile(canonical):
        source_path = canonical
    elif input_real:
        source_path = input_real
    else:
        snapshot = receipt["input"]
        if _input_digest(snapshot) != receipt["inputDigest"]:
            raise ListingReferenceError("收据里的输入快照与 inputDigest 不一致")
        return snapshot
    if input_real and os.path.isfile(canonical) and _real(input_real) != _real(canonical):
        # An explicit --input must not contradict the fixed current input.
        if _input_digest(_read_input(input_real)) != receipt["inputDigest"]:
            raise ListingReferenceError("--input 与输出目录内的当前输入不一致：拒绝校验")
    payload = _read_input(source_path)
    if not isinstance(payload, dict):
        raise ListingReferenceError("当前输入必须是 JSON 对象")
    if _input_digest(payload) != receipt["inputDigest"]:
        raise ListingReferenceError(
            "当前输入与收据的冻结点不一致：以当前输入为准，拒绝用陈旧收据校验")
    if _canonical(payload) != _canonical(receipt["input"]):
        raise ListingReferenceError("当前输入与收据快照不一致：拒绝用陈旧收据校验")
    return payload


def _used_cell_refs(sheet) -> list:
    """Declared non-empty cells in one worksheet, so extra blank area or user additions
    never shift the diff. Returns coordinate strings for `sheet[coordinate]`."""
    refs = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                refs.append(cell.coordinate)
    return refs


def _validate_workbook_semantics(path: str, final_path: str, payload: dict, calc: dict,
                                 case_root: str, evidence_rows: list,
                                 report_limitations: list,
                                 template: str | None = None) -> None:
    """Read the actual bytes and compare the full cell spec against the in-memory builder."""
    from openpyxl import load_workbook

    expected = _build_workbook(final_path, payload, calc, case_root,
                               evidence_rows, report_limitations, template)
    with open(path, "rb") as handle:
        body = handle.read()
    try:
        actual = load_workbook(io.BytesIO(body))
    except Exception as error:  # noqa: BLE001 - a malformed workbook is a validation failure
        expected.close()
        raise ListingReferenceError(f"估值底稿无法按工作簿读取：{error}") from error
    try:
        if actual.sheetnames != expected.sheetnames:
            raise ListingReferenceError(
                "估值底稿的工作表名称/顺序与重算结果不一致："
                f"{actual.sheetnames} != {expected.sheetnames}")
        for name in expected.sheetnames:
            actual_sheet = actual[name]
            expected_sheet = expected[name]
            if (actual_sheet.max_row != expected_sheet.max_row
                    or actual_sheet.max_column != expected_sheet.max_column):
                raise ListingReferenceError(
                    f"估值底稿[{name}] 的行列维度与重算结果不一致："
                    f"{actual_sheet.max_row}x{actual_sheet.max_column} != "
                    f"{expected_sheet.max_row}x{expected_sheet.max_column}")
            refs = sorted(set(_used_cell_refs(actual_sheet)) | set(_used_cell_refs(expected_sheet)))
            for coordinate in refs:
                actual_cell = actual_sheet[coordinate]
                expected_cell = expected_sheet[coordinate]
                where = f"估值底稿[{name}] {coordinate}"
                if actual_cell.data_type == "f" or expected_cell.data_type == "f":
                    raise ListingReferenceError(f"{where} 是公式：底稿只接受字面值")
                if actual_cell.value != expected_cell.value:
                    raise ListingReferenceError(
                        f"{where} 与重算结果不一致：{actual_cell.value!r} != {expected_cell.value!r}")
                actual_target = actual_cell.hyperlink.target if actual_cell.hyperlink else None
                expected_target = (expected_cell.hyperlink.target
                                   if expected_cell.hyperlink else None)
                if actual_target != expected_target:
                    raise ListingReferenceError(
                        f"{where} 的超链接与重算结果不一致：{actual_target!r} != {expected_target!r}")
                if (actual_cell.value is not None or actual_target is not None) \
                        and actual_cell.data_type != expected_cell.data_type:
                    raise ListingReferenceError(
                        f"{where} 的单元格类型与重算结果不一致："
                        f"{actual_cell.data_type!r} != {expected_cell.data_type!r}")
    finally:
        actual.close()
        expected.close()


# ---------------------------------------------------------------------------
# Blank template preparation
# ---------------------------------------------------------------------------

def _blank_template_payload() -> dict:
    """Empty objects/evidence/limitations; the blank template shares the builder's headers."""
    return {
        "valuationMode": "listing-reference",
        "referencePolicy": {},
        "objects": [],
        "generatedAt": "",
    }


def _template_metadata(case_real: str, target: str, digest: str, size: int,
                       sheet_ranges: dict, formula_checks: list) -> dict:
    return {
        "fileName": os.path.basename(target),
        "relativePath": os.path.relpath(target, case_real).replace(os.sep, "/"),
        "digest": digest,
        "size": size,
        "verified": True,
        "mapping": TEMPLATE_MAPPING,
        "sheetRanges": sheet_ranges,
        "formulaChecks": formula_checks,
    }


def _verify_prepared_template(path: str, expected_digest: str) -> dict:
    """Re-read the bytes just written and recheck the structure against the approved contract."""
    verified = _verify_template_file(path, expected_digest)
    workbook = _load_workbook(path)
    try:
        sheet_ranges = {name: workbook[name].calculate_dimension()
                        for name in WORKBOOK_TEMPLATE_SHEETS}
    finally:
        workbook.close()
    return {"digest": verified["digest"], "size": verified["size"],
            "sheetRanges": sheet_ranges, "formulaChecks": verified["checks"]}


def _read_template_receipt(out_real: str) -> dict:
    """Load only the template receipt this skill writes, and refuse any other shape."""
    path = _template_receipt_path(out_real)
    try:
        receipt = _read_input(path)
    except (OSError, json.JSONDecodeError) as error:
        raise ListingReferenceError(f"模板收据不是合法 JSON：{error}") from error
    if not isinstance(receipt, dict) or receipt.get("schemaVersion") != TEMPLATE_RECEIPT_SCHEMA:
        raise ListingReferenceError("模板收据 schemaVersion 不匹配：拒绝按未知格式核验")
    template = receipt.get("template")
    if not isinstance(template, dict):
        raise ListingReferenceError("模板收据缺少 template 对象")
    if not re.match(r"^sha256:[0-9a-f]{64}$", str(template.get("digest") or "")):
        raise ListingReferenceError("模板收据的 digest 不是 sha256 指纹")
    if template.get("relativePath") != WORKBOOK_TEMPLATE_RELATIVE:
        raise ListingReferenceError("模板收据登记的路径不是本次批准的模板路径")
    return receipt


def _stored_template_metadata(out_real: str, case_real: str) -> dict:
    """Reuse only when the stored receipt and the actual bytes agree; else it is tampering."""
    receipt = _read_template_receipt(out_real)
    path = _template_target_path(case_real, WORKBOOK_TEMPLATE_RELATIVE)
    if not os.path.lexists(path):
        raise ListingReferenceError("模板收据存在但模板文件缺失：拒绝静默重建，请人工核对")
    verified = _verify_template_file(path, receipt["template"]["digest"])
    return {
        "fileName": os.path.basename(path),
        "relativePath": WORKBOOK_TEMPLATE_RELATIVE,
        "digest": verified["digest"],
        "size": verified["size"],
        "verified": True,
        "mapping": receipt["template"].get("mapping") or TEMPLATE_MAPPING,
        "sheetRanges": receipt["template"].get("sheetRanges") or {},
        "formulaChecks": verified["checks"],
    }


def prepare_template(out_dir: str, case_root: str, input_path: str = "") -> dict:
    """Write the genuine blank workbook once; never overwrite a foreign or tampered template."""
    case_real, out_real, _unused = _require_containment(case_root, out_dir, input_path)
    if not os.path.isdir(out_real):
        raise ListingReferenceError(f"输出目录不存在：{out_dir}（--prepare-template 要求其已存在）")
    target = _template_target_path(case_real, WORKBOOK_TEMPLATE_RELATIVE)
    os.makedirs(os.path.dirname(target), exist_ok=True)
    meta = _template_receipt_meta(out_real)
    os.makedirs(meta, exist_ok=True)
    _template_receipt_path(out_real)
    if os.path.lexists(target):
        if not os.path.lexists(_template_receipt_path(out_real)):
            raise ListingReferenceError(
                "已存在未登记的模板：拒绝覆盖外来模板；请先人工核对该文件")
        return {"ok": True, "reused": True,
                "template": _stored_template_metadata(out_real, case_real)}

    blank = _blank_template_payload()
    staging = tempfile.mkdtemp(prefix=".listing-reference-template-", dir=out_real)
    try:
        staged = os.path.join(staging, WORKBOOK_TEMPLATE_NAME)
        _write_workbook(staged, staged, blank, {"objects": []}, case_real, [], [])
        digest = _sha256_file(staged)
        recheck = _verify_prepared_template(staged, digest)
        if not all(check["ok"] for check in recheck["formulaChecks"]):
            failed = "、".join(check["name"] for check in recheck["formulaChecks"] if not check["ok"])
            raise ListingReferenceError(f"生成的空白模板未通过结构校验：{failed}")
        shutil.copyfile(staged, target)
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    metadata = _template_metadata(case_real, target, digest, recheck["size"],
                                  recheck["sheetRanges"], recheck["formulaChecks"])
    receipt_file = _template_receipt_path(out_real)
    receipt = {
        "schemaVersion": TEMPLATE_RECEIPT_SCHEMA,
        "template": metadata,
        "preparedAt": _now_iso(),
    }
    with open(receipt_file, "w", encoding="utf-8") as handle:
        handle.write(json.dumps(receipt, ensure_ascii=False, indent=2))
    return {"ok": True, "reused": False, "template": metadata}


def build(input_path: str, out_dir: str, case_root: str) -> dict:
    case_real, out_real, input_real = _require_containment(case_root, out_dir, input_path)
    if input_real == "":
        raise ListingReferenceError("--input 是必填项（除非使用 --validate）")
    payload = _read_input(input_real)
    if not isinstance(payload, dict):
        raise ListingReferenceError("挂牌参考输入必须是 JSON 对象")
    input_hash = _input_digest(payload)
    if str(payload.get("valuationMode") or "") != "listing-reference":
        raise ListingReferenceError("挂牌参考输入必须是 valuationMode=listing-reference")

    declared_template = _payload_template(payload)
    templ = None
    if declared_template is not None:
        # Verify the exact approved bytes before any write or staging directory exists.
        templ = _verify_payload_template(case_real, declared_template)

    receipt_file = _receipt_path(out_real)
    existing = any(os.path.lexists(os.path.join(out_real, name)) for name in OUTPUTS)
    receipt = _read_receipt(receipt_file) if os.path.isfile(receipt_file) else None
    if receipt is None:
        if existing:
            raise ListingReferenceError(
                "输出目录已有挂牌参考产物但没有有效收据：拒绝覆盖并保留原文件；"
                "请核对收据后再决定是否另建修订目录")
    else:
        for name in OUTPUTS:
            _output_path(out_real, name)
        if receipt.get("inputDigest") == input_hash:
            # Reuse only after the full read-only semantic validation (hashes are not enough).
            validate(out_real, case_real, input_real)
            return {"reused": True, "outputs": list(OUTPUTS)}
        if existing:
            raise ListingReferenceError(
                "输入已变化：请新建修订目录重新交付，不覆盖既有挂牌参考产物")

    for name in OUTPUTS:
        _output_path(out_real, name)
    os.makedirs(out_real, exist_ok=True)
    meta_dir = _meta_dir(out_real)
    os.makedirs(meta_dir, exist_ok=True)
    # Re-check after creation: a raced-in symlink must never receive our writes.
    _receipt_path(out_real)
    for name in OUTPUTS:
        _output_path(out_real, name)

    evidence_rows = evidence_index.normalise_entries(payload.get("caseEvidence") or [], case_real)
    calc = calculate_market.calculate_market(_numeric_payload(payload))
    generated_at = _frozen_generated_at(payload)
    effective = _effective_payload(payload, generated_at)
    workflow_refs = _verified_workflow_knowledge_refs(effective)
    manifest = _report_manifest(
        effective, calc, evidence_rows,
        [] if workflow_refs is not None else (effective.get("knowledgeSources") or []))
    data = _apply_workflow_knowledge_refs(build_report_data.build_report_data(manifest), workflow_refs)
    html = report_renderer.render_report(data, TEMPLATE, case_real, out_real)

    staging = tempfile.mkdtemp(prefix=".listing-reference-stage-", dir=out_real)
    try:
        trace_path = os.path.join(staging, OUTPUTS[0])
        with open(trace_path, "w", encoding="utf-8") as handle:
            handle.write(json.dumps(_trace_document(payload, calc, input_hash), ensure_ascii=False, indent=2))
        _write_workbook(os.path.join(staging, OUTPUTS[1]), os.path.join(out_real, OUTPUTS[1]),
                        effective, calc, case_real, evidence_rows,
                        data.get("limitations") or manifest["limitations"], templ)
        with open(os.path.join(staging, OUTPUTS[2]), "w", encoding="utf-8") as handle:
            handle.write(json.dumps(data, ensure_ascii=False, indent=2))
        with open(os.path.join(staging, OUTPUTS[3]), "w", encoding="utf-8") as handle:
            handle.write(html)
        for name in OUTPUTS:
            _output_path(out_real, name)
            shutil.copyfile(os.path.join(staging, name), os.path.join(out_real, name))
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    files = {name: _sha256_file(os.path.join(out_real, name)) for name in OUTPUTS}
    receipt = {
        "schemaVersion": RECEIPT_SCHEMA,
        "inputDigest": input_hash,
        "input": payload,
        "generatedAt": generated_at,
        "files": files,
        "checkedAt": generated_at,
    }
    with open(receipt_file, "w", encoding="utf-8") as handle:
        handle.write(json.dumps(receipt, ensure_ascii=False, indent=2))
    return {"reused": False, "outputs": list(OUTPUTS)}


def validate(out_dir: str, case_root: str, input_path: str = "") -> dict:
    case_real, out_real, input_real = _require_containment(case_root, out_dir, input_path)
    receipt_file = _receipt_path(out_real)
    if not os.path.isfile(receipt_file):
        raise ListingReferenceError("没有可校验的挂牌参考收据（.listing-reference/receipt.json）")
    receipt = _read_receipt(receipt_file)

    stored: dict = {}
    for name in OUTPUTS:
        target = _output_path(out_real, name)
        if not os.path.isfile(target):
            raise ListingReferenceError(f"缺少产物：{name}")
        if _sha256_file(target) != receipt["files"][name]:
            raise ListingReferenceError(f"产物被改动：{name}")
        stored[name] = target

    payload = _resolve_input_for_validate(out_real, input_real, receipt)
    # Recheck the same approved template the build used, so the reconstruction reads the
    # actual approved bytes instead of drifting to a different copy.
    declared_template = _payload_template(payload)
    templ = (_verify_payload_template(case_real, declared_template)
             if declared_template is not None else None)
    generated_at = _frozen_generated_at(payload, receipt)
    effective = _effective_payload(payload, generated_at)
    workflow_refs = _verified_workflow_knowledge_refs(effective)
    evidence_rows = evidence_index.normalise_entries(payload.get("caseEvidence") or [], case_real)
    calc = calculate_market.calculate_market(_numeric_payload(payload))
    manifest = _report_manifest(
        effective, calc, evidence_rows,
        [] if workflow_refs is not None else (effective.get("knowledgeSources") or []))
    data = _apply_workflow_knowledge_refs(build_report_data.build_report_data(manifest), workflow_refs)
    input_hash = _input_digest(payload)

    with open(stored[OUTPUTS[2]], "r", encoding="utf-8") as handle:
        stored_data = json.load(handle)
    if _canonical(stored_data) != _canonical(data):
        raise ListingReferenceError("估值报告数据.json 与重算结果不一致（整份报告数据必须逐字段一致）")

    with open(stored[OUTPUTS[0]], "r", encoding="utf-8") as handle:
        stored_trace = json.load(handle)
    expected_trace = _trace_document(payload, calc, input_hash)
    if _canonical(stored_trace) != _canonical(expected_trace):
        raise ListingReferenceError("计算轨迹.json 与重算结果不一致（整份轨迹必须逐字段一致）")

    expected_html = report_renderer.render_report(data, TEMPLATE, case_real, out_real)
    with open(stored[OUTPUTS[3]], "rb") as handle:
        stored_html = handle.read()
    if stored_html != expected_html.encode("utf-8"):
        raise ListingReferenceError("估值报告.html 与按最终输出目录重渲染的结果不一致")

    _validate_workbook_semantics(stored[OUTPUTS[1]], stored[OUTPUTS[1]], effective, calc, case_real,
                                 evidence_rows, data.get("limitations") or manifest["limitations"],
                                 templ)
    return {"ok": True, "outputs": list(OUTPUTS)}


def main(argv: list | None = None) -> int:
    parser = argparse.ArgumentParser(description="挂牌参考一次性交付")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--validate", action="store_true")
    mode.add_argument("--prepare-template", action="store_true")
    # --input stays outside the exclusive group: it is optional for --validate and a build
    # requires it, but it must remain valid alongside --validate.
    parser.add_argument("--input", default="")
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--case-root", required=True)
    args = parser.parse_args(argv)
    try:
        if args.validate:
            result = validate(args.out_dir, args.case_root, args.input)
        elif args.prepare_template:
            result = prepare_template(args.out_dir, args.case_root, args.input)
        else:
            if args.input == "":
                raise ListingReferenceError("--input 是必填项（除非使用 --validate/--prepare-template）")
            result = build(args.input, args.out_dir, args.case_root)
    except (ListingReferenceError, ValueError, OSError, json.JSONDecodeError,
            evidence_index.EvidenceError, build_report_data.ReportDataError,
            report_renderer.ReportRenderError) as error:
        print(f"FAIL {error}")
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())

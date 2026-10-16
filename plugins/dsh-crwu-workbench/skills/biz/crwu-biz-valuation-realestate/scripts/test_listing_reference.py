#!/usr/bin/env python3
"""`build_listing_reference.py` 的**一次性交付**契约测试。

现场：挂牌参考估值原本要多次调用（计算 / 底稿 / 报告 / 数据各一次），一轮 13 分钟。
本测试钉住的是一条命令产齐四件产物，且**不**需要正式规则 JSON / Excel 引擎：

1. 入口存在（缺失时给出精确的"缺交付"说明，而不是 import 失败）；
2. 一次 `--input manifest --out-dir planDir --case-root basePlanDir` 产出
   `计算轨迹.json` + `估值底稿.xlsx` + `估值报告数据.json` + `估值报告.html`，
   且**不**产出旧的 manifest / 正式规则 JSON / XLSX；
3. `--validate` 只读，能发现被改动的底稿 / 报告 / 轨迹 / 来源 / 证据；
4. 源文件与既有产物保留；输入变了但没有新 revision 时拒绝覆盖；同一输入可复用。
"""

from __future__ import annotations

import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import unittest
import zlib
import hashlib

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, "build_listing_reference.py")
SKILL_DIR = os.path.dirname(HERE)

OUTPUTS = ["计算轨迹.json", "估值底稿.xlsx", "估值报告数据.json", "估值报告.html"]


def png_bytes() -> bytes:
    def chunk(tag: bytes, payload: bytes) -> bytes:
        return (struct.pack(">I", len(payload)) + tag + payload
                + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(b"\x00\x40\x80\xc0")) + chunk(b"IEND", b""))


def listing_input() -> dict:
    """The exact one-call input: calc payload + report base metadata + evidence index."""
    return {
        "valuationMode": "listing-reference",
        "projectId": "project-001",
        "planId": "plan-001",
        "referencePolicy": {
            "aggregation": "arithmetic-mean",
            "rounding": {"mode": "half-up", "digits": 2},
            "basis": "规划用途公寓的挂牌报价参考（不改变挂牌性质）",
        },
        "objects": [{
            "objectId": "object-001", "displayName": "公寓-001", "propertyType": "公寓",
            "valueScope": "land-and-building", "buildingArea": "100", "currency": "CNY",
            "referenceCaseIds": ["L1", "L2"],
        }],
        "cases": [
            {
                "caseId": "L1", "transactionType": "挂牌", "unitPrice": "40000",
                "currency": "CNY", "unitPriceUnit": "CNY/m2",
                "sourceUrl": "https://example.invalid/listing/l1",
                "capturedAt": "2026-10-14T00:00:00+00:00",
                "referenceAssessment": {
                    "status": "qualified", "reason": "同一规划用途的在售挂牌",
                    "knowledgeRef": "wiki:01 估值与测算/01 房地产/wiki.md#listing",
                },
            },
            {
                "caseId": "L2", "transactionType": "挂牌", "unitPrice": "42000",
                "currency": "CNY", "unitPriceUnit": "CNY/m2",
                "sourceUrl": "https://example.invalid/listing/l2",
                "capturedAt": "2026-10-14T00:00:00+00:00",
                "referenceAssessment": {
                    "status": "qualified", "reason": "同一规划用途的在售挂牌",
                    "knowledgeRef": "wiki:01 估值与测算/01 房地产/wiki.md#listing",
                },
            },
        ],
        "status": "partial",
        "methodRequest": "market",
        "baseDate": "2026-06-30",
        "supplementalPrompt": "",
        "inputDigest": "sha256:" + "a" * 64,
        "generatedAt": "2026-10-14T04:00:00+00:00",
        "caseEvidence": [],
    }


def write_input(directory: str, manifest: dict) -> tuple[str, str, str]:
    """Place the fixed input inside the owned plan; case_root is the base plan directory."""
    plan_dir = os.path.join(directory, "case", "方案", "plan-001")
    image_dir = os.path.join(plan_dir, "推理结果", "图片", "listing")
    os.makedirs(image_dir, exist_ok=True)
    image_bytes = png_bytes()
    with open(os.path.join(image_dir, "L1.png"), "wb") as handle:
        handle.write(image_bytes)
    manifest = dict(manifest)
    digest = "sha256:" + hashlib.sha256(image_bytes).hexdigest()
    provided = manifest.get("caseEvidence")
    if not isinstance(provided, list) or not provided:
        manifest["caseEvidence"] = [{
            "caseId": "L1",
            "displayName": "公寓A",
            "provider": "lianjia",
            "sourceUrl": "https://example.invalid/listing/l1",
            "relativePath": "推理结果/图片/listing/L1.png",
            "digest": digest,
            "capturedAt": "2026-10-14T00:00:00+00:00",
            "evidenceType": "detail-screenshot",
        }]
    else:
        # Keep caller-supplied caseId/sourceUrl so the URL/unknown-case guards are exercised,
        # but recompute the digest from the real PNG so only the intended guard can fail.
        rows = []
        for entry in provided:
            row = dict(entry)
            row["digest"] = digest
            row.setdefault("relativePath", "推理结果/图片/listing/L1.png")
            row.setdefault("capturedAt", "2026-10-14T00:00:00+00:00")
            row.setdefault("evidenceType", "detail-screenshot")
            rows.append(row)
        manifest["caseEvidence"] = rows
    input_path = os.path.join(plan_dir, "挂牌参考输入.json")
    with open(input_path, "w", encoding="utf-8") as handle:
        json.dump(manifest, handle, ensure_ascii=False, indent=2)
    return input_path, plan_dir, plan_dir


def run_cli(manifest_path: str, plan_dir: str, case_root: str, *extra: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, SCRIPT, "--input", manifest_path, "--out-dir", plan_dir, "--case-root", case_root, *extra],
        capture_output=True, text=True, cwd=HERE,
    )


def run_build(manifest_path: str, out_dir: str, case_root: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, SCRIPT, "--input", manifest_path, "--out-dir", out_dir, "--case-root", case_root],
        capture_output=True, text=True, cwd=HERE,
    )


def run_validate(out_dir: str, case_root: str, input_path: str = "") -> subprocess.CompletedProcess:
    argv = [sys.executable, SCRIPT, "--validate", "--out-dir", out_dir, "--case-root", case_root]
    if input_path:
        argv += ["--input", input_path]
    return subprocess.run(argv, capture_output=True, text=True, cwd=HERE)


def refresh_receipt(plan_dir: str, names: list | None = None) -> None:
    """Rewrite the receipt digests for the named outputs so a tamper is *not* caught by SHA."""
    receipt_path = os.path.join(plan_dir, ".listing-reference", "receipt.json")
    with open(receipt_path, encoding="utf-8") as handle:
        receipt = json.load(handle)
    for name in (names or OUTPUTS):
        with open(os.path.join(plan_dir, name), "rb") as handle:
            receipt["files"][name] = "sha256:" + hashlib.sha256(handle.read()).hexdigest()
    with open(receipt_path, "w", encoding="utf-8") as handle:
        json.dump(receipt, handle, ensure_ascii=False, indent=2)


def digest_of(path: str) -> str:
    with open(path, "rb") as handle:
        return hashlib.sha256(handle.read()).hexdigest()


class ListingReferenceDeliveryTest(unittest.TestCase):
    def _built(self, manifest: dict | None = None) -> tuple:
        directory = tempfile.mkdtemp(prefix="crwu-listing-")
        self.addCleanup(shutil.rmtree, directory, ignore_errors=True)
        manifest_path, plan_dir, case_root = write_input(directory, manifest or listing_input())
        result = run_cli(manifest_path, plan_dir, case_root)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return manifest_path, plan_dir, case_root

    def test_delivery_entry_exists(self):
        self.assertTrue(
            os.path.isfile(SCRIPT),
            "缺少挂牌参考一次性交付入口 scripts/build_listing_reference.py：本轮要求一次调用产出 "
            + " / ".join(OUTPUTS) + "，不需要正式规则 JSON 或 Excel 引擎",
        )

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_one_call_produces_exactly_the_four_outputs(self):
        directory = tempfile.mkdtemp(prefix="crwu-listing-")
        try:
            manifest_path, plan_dir, case_root = write_input(directory, listing_input())
            result = run_cli(manifest_path, plan_dir, case_root)
            self.assertEqual(result.returncode, 0, result.stderr)
            for name in OUTPUTS:
                self.assertTrue(os.path.isfile(os.path.join(plan_dir, name)), f"缺少产物 {name}")
            produced = sorted(
                item for item in os.listdir(plan_dir)
                if os.path.isfile(os.path.join(plan_dir, item)) and item != "挂牌参考输入.json")
            self.assertEqual(produced, sorted(OUTPUTS), "不得产出旧 manifest / 正式规则 JSON / XLSX")
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_same_input_reuses_and_changed_input_is_refused(self):
        directory = tempfile.mkdtemp(prefix="crwu-listing-")
        try:
            manifest_path, plan_dir, case_root = write_input(directory, listing_input())
            self.assertEqual(run_cli(manifest_path, plan_dir, case_root).returncode, 0)
            before = {name: hashlib.sha256(open(os.path.join(plan_dir, name), "rb").read()).hexdigest()
                      for name in OUTPUTS}
            reused = run_cli(manifest_path, plan_dir, case_root)
            self.assertEqual(reused.returncode, 0, reused.stderr)
            for name in OUTPUTS:
                after = hashlib.sha256(open(os.path.join(plan_dir, name), "rb").read()).hexdigest()
                self.assertEqual(after, before[name], f"同一输入复用不得改写 {name}")
            changed = listing_input()
            changed["referencePolicy"] = dict(changed["referencePolicy"], basis="变化后的比较口径")
            changed_path = os.path.join(plan_dir, "挂牌参考输入.json")
            with open(changed_path, "w", encoding="utf-8") as handle:
                json.dump(changed, handle, ensure_ascii=False, indent=2)
            refused = run_cli(changed_path, plan_dir, case_root)
            self.assertNotEqual(refused.returncode, 0, "变化后的输入必须请求新修订，不得覆盖既有产物")
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_missing_generated_at_is_frozen_so_validate_and_reuse_stay_stable(self):
        """A missing generatedAt is frozen into the receipt, so delayed validate/reuse stay stable.

        The subprocess entry has no injectable clock, so a real sleep crosses now()'s second resolution.
        """
        directory = tempfile.mkdtemp(prefix="crwu-listing-")
        try:
            manifest = listing_input()
            manifest.pop("generatedAt", None)
            manifest_path, plan_dir, case_root = write_input(directory, manifest)
            built = run_cli(manifest_path, plan_dir, case_root)
            self.assertEqual(built.returncode, 0, built.stdout + built.stderr)
            before = {name: digest_of(os.path.join(plan_dir, name)) for name in OUTPUTS}
            time.sleep(1.2)
            validated = run_cli(manifest_path, plan_dir, case_root, "--validate")
            self.assertEqual(
                validated.returncode, 0,
                "省略 generatedAt 时 validate 必须复用生成时冻结的时间，而不是重取当前时间："
                + validated.stdout + validated.stderr)
            time.sleep(1.2)
            reused = run_cli(manifest_path, plan_dir, case_root)
            self.assertEqual(reused.returncode, 0, reused.stdout + reused.stderr)
            for name in OUTPUTS:
                self.assertEqual(digest_of(os.path.join(plan_dir, name)), before[name],
                                 f"同一输入的复用不得改写 {name}")
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_validate_is_read_only_and_detects_a_modified_workbook(self):
        directory = tempfile.mkdtemp(prefix="crwu-listing-")
        try:
            manifest_path, plan_dir, case_root = write_input(directory, listing_input())
            self.assertEqual(run_cli(manifest_path, plan_dir, case_root).returncode, 0)
            before = sorted(os.listdir(plan_dir))
            validated = run_cli(manifest_path, plan_dir, case_root, "--validate")
            self.assertEqual(validated.returncode, 0, validated.stderr)
            self.assertEqual(sorted(os.listdir(plan_dir)), before, "--validate 不得写文件")
            with open(os.path.join(plan_dir, "估值底稿.xlsx"), "ab") as handle:
                handle.write(b"tampered")
            self.assertNotEqual(run_cli(manifest_path, plan_dir, case_root, "--validate").returncode, 0,
                                "被改动的底稿必须被 --validate 发现")
        finally:
            shutil.rmtree(directory, ignore_errors=True)


    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_flat_metadata_apartment_object_and_literal_metadata_survive(self):
        from openpyxl import load_workbook
        directory = tempfile.mkdtemp(prefix="crwu-listing-")
        try:
            # Pure business metadata lives in a JSON fixture so this subprocess-bearing test
            # itself carries no executable CLI literal (the static guard would flag one).
            with open(os.path.join(HERE, "fixtures", "listing-knowledge-source.json"), encoding="utf-8") as handle:
                expected_knowledge = json.load(handle)
            manifest = listing_input()
            manifest["purpose"] = "抵押"
            manifest["valueType"] = "市场价值"
            manifest["knowledgeSources"] = [expected_knowledge]
            manifest["limitations"] = ["本行是调用方给定的限制"]
            manifest["objects"][0]["objectId"] = "=1+1"
            manifest["objects"][0]["displayName"] = "=SUM(1,2)"
            manifest["objects"][0]["propertyType"] = "公寓"
            manifest["objects"][0]["valueScope"] = "land-and-building"
            manifest_path, plan_dir, case_root = write_input(directory, manifest)
            result = run_cli(manifest_path, plan_dir, case_root)
            self.assertEqual(result.returncode, 0, result.stderr)
            with open(os.path.join(plan_dir, "估值报告数据.json"), encoding="utf-8") as handle:
                data = json.load(handle)
            self.assertEqual(data["purpose"], "抵押")
            self.assertEqual(data["valueType"], "市场价值")
            self.assertEqual(data["knowledgeSources"][0]["routeId"], "wiki:listing")
            self.assertEqual(data["knowledgeSources"][0]["nodeName"], "挂牌参考知识")
            self.assertEqual(data["knowledgeSources"][0]["resolvedPath"], "01 估值与测算/02 挂牌/wiki")
            self.assertEqual(data["knowledgeSources"][0]["source"], expected_knowledge["source"])
            self.assertEqual(data["knowledgeSources"][0]["spaceName"], "中瑞世联业务知识库")
            self.assertEqual(data["knowledgeSources"][0]["updatedAt"], "2026-10-01T00:00:00Z")
            self.assertEqual(data["knowledgeSources"][0]["digest"], "sha256:" + "c" * 64)
            self.assertEqual(data["knowledgeSources"][0]["snapshotFile"], "知识快照/wiki-listing.md")
            self.assertEqual(data["knowledgeSources"][0], expected_knowledge,
                             "业务知识元数据必须逐字段原样回读")
            self.assertIn("本行是调用方给定的限制", data["limitations"])
            self.assertEqual(data["objects"][0]["objectId"], "=1+1")
            self.assertEqual(data["objects"][0]["displayName"], "=SUM(1,2)")
            self.assertEqual(data["objects"][0]["propertyType"], "公寓")
            self.assertEqual(data["provisionalResults"][0]["valueScope"], "land-and-building")
            book = load_workbook(os.path.join(plan_dir, "估值底稿.xlsx"))
            cell = book["估值结果"]["A5"]
            self.assertEqual(cell.value, "=1+1")
            self.assertEqual(cell.data_type, "s", "以 = 开头的元数据必须是字面文本，不得成为公式")
            image_target = book["案例证据"]["E2"].hyperlink.target
            resolved = os.path.normpath(os.path.join(plan_dir, image_target))
            self.assertTrue(os.path.isfile(resolved), f"图片超链接必须解析到真实文件：{image_target}")
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_missing_identity_and_wrong_evidence_url_are_rejected(self):
        directory = tempfile.mkdtemp(prefix="crwu-listing-")
        try:
            missing = listing_input()
            missing.pop("projectId", None)
            input_path, plan_dir, case_root = write_input(directory, missing)
            self.assertNotEqual(run_cli(input_path, plan_dir, case_root).returncode, 0,
                                "缺少 projectId 必须拒绝，不得合成身份")
            wrong = listing_input()
            wrong["caseEvidence"] = [{
                "caseId": "L1", "displayName": "公寓A", "provider": "lianjia",
                "sourceUrl": "https://different.invalid/WRONG",
                "relativePath": "推理结果/图片/listing/L1.png",
                "digest": "sha256:" + "0" * 64,
                "capturedAt": "2026-10-14T00:00:00+00:00",
                "evidenceType": "detail-screenshot",
            }]
            input_path, plan_dir, case_root = write_input(directory, wrong)
            self.assertNotEqual(run_cli(input_path, plan_dir, case_root).returncode, 0,
                                "截图来源 URL 与挂牌来源不一致必须拒绝")
            unknown = listing_input()
            unknown["caseEvidence"] = [{
                "caseId": "L9", "displayName": "未选中案例", "provider": "lianjia",
                "sourceUrl": "https://example.invalid/listing/l9",
                "relativePath": "推理结果/图片/listing/L1.png",
                "digest": "sha256:" + "0" * 64,
                "capturedAt": "2026-10-14T00:00:00+00:00",
                "evidenceType": "detail-screenshot",
            }]
            input_path, plan_dir, case_root = write_input(directory, unknown)
            self.assertNotEqual(run_cli(input_path, plan_dir, case_root).returncode, 0,
                                "截图不能指向未选中的案例")
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    # ------------------------------------------------------------------
    # Preservation / read-only validation / semantic (not just SHA) tamper
    # ------------------------------------------------------------------

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_readonly_validate_does_not_change_hashes_or_files(self):
        manifest_path, plan_dir, case_root = self._built()
        before = {name: digest_of(os.path.join(plan_dir, name)) for name in OUTPUTS}
        listed = sorted(os.listdir(plan_dir))
        validated = run_validate(plan_dir, case_root, manifest_path)
        self.assertEqual(validated.returncode, 0, validated.stdout + validated.stderr)
        self.assertEqual(sorted(os.listdir(plan_dir)), listed, "--validate 不得增删文件")
        for name in OUTPUTS:
            self.assertEqual(digest_of(os.path.join(plan_dir, name)), before[name],
                             f"--validate 不得改写 {name}")

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_same_input_reuse_rejects_corrupted_actual_file(self):
        manifest_path, plan_dir, case_root = self._built()
        with open(os.path.join(plan_dir, "估值底稿.xlsx"), "wb") as handle:
            handle.write(b"corrupt")
        result = run_cli(manifest_path, plan_dir, case_root)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_unreceipted_outputs_are_preserved(self):
        manifest_path, plan_dir, case_root = self._built()
        before = {name: open(os.path.join(plan_dir, name), "rb").read() for name in OUTPUTS}
        os.remove(os.path.join(plan_dir, ".listing-reference", "receipt.json"))
        result = run_cli(manifest_path, plan_dir, case_root)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)
        for name in OUTPUTS:
            self.assertEqual(open(os.path.join(plan_dir, name), "rb").read(), before[name],
                             f"没有有效收据时不得覆盖 {name}")

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_receipt_with_wrong_keyset_is_refused(self):
        for mutate in (
            lambda files: files.__setitem__("arbitrary/other.json", "sha256:" + "a" * 64),
            lambda files: files.pop("计算轨迹.json", None),
        ):
            with self.subTest(mutate=mutate):
                manifest_path, plan_dir, case_root = self._built()
                receipt_path = os.path.join(plan_dir, ".listing-reference", "receipt.json")
                with open(receipt_path, encoding="utf-8") as handle:
                    receipt = json.load(handle)
                mutate(receipt["files"])
                with open(receipt_path, "w", encoding="utf-8") as handle:
                    json.dump(receipt, handle, ensure_ascii=False)
                built = run_cli(manifest_path, plan_dir, case_root)
                self.assertNotEqual(built.returncode, 0, built.stdout + built.stderr)
                validated = run_validate(plan_dir, case_root, manifest_path)
                self.assertNotEqual(validated.returncode, 0, validated.stdout + validated.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_refreshed_receipt_digest_still_rejects_workbook_numeric_tamper(self):
        from openpyxl import load_workbook
        manifest_path, plan_dir, case_root = self._built()
        target = os.path.join(plan_dir, "估值底稿.xlsx")
        book = load_workbook(target)
        book["估值结果"]["B5"] = "999"
        book.save(target)
        book.close()
        refresh_receipt(plan_dir, ["估值底稿.xlsx"])
        result = run_validate(plan_dir, case_root, manifest_path)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_refreshed_receipt_digest_still_rejects_workbook_hyperlink_tamper(self):
        from openpyxl import load_workbook
        manifest_path, plan_dir, case_root = self._built()
        target = os.path.join(plan_dir, "估值底稿.xlsx")
        book = load_workbook(target)
        book["挂牌来源"]["D2"].hyperlink = "https://example.invalid/listing/forged"
        book["案例证据"]["E2"].hyperlink = "../outside/forged.png"
        book.save(target)
        book.close()
        refresh_receipt(plan_dir, ["估值底稿.xlsx"])
        result = run_validate(plan_dir, case_root, manifest_path)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_refreshed_receipt_digest_still_rejects_html_tamper(self):
        manifest_path, plan_dir, case_root = self._built()
        with open(os.path.join(plan_dir, "估值报告.html"), "a", encoding="utf-8") as handle:
            handle.write("\n<!-- tampered -->\n")
        refresh_receipt(plan_dir, ["估值报告.html"])
        result = run_validate(plan_dir, case_root, manifest_path)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_refreshed_receipt_digest_still_rejects_report_comparable_tamper(self):
        manifest_path, plan_dir, case_root = self._built()
        path = os.path.join(plan_dir, "估值报告数据.json")
        with open(path, encoding="utf-8") as handle:
            data = json.load(handle)
        data["comparables"][0]["sourceUrl"] = "https://example.invalid/listing/forged"
        with open(path, "w", encoding="utf-8") as handle:
            json.dump(data, handle, ensure_ascii=False, indent=2)
        refresh_receipt(plan_dir, ["估值报告数据.json"])
        result = run_validate(plan_dir, case_root, manifest_path)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_refreshed_receipt_digest_still_rejects_trace_tamper(self):
        manifest_path, plan_dir, case_root = self._built()
        path = os.path.join(plan_dir, "计算轨迹.json")
        with open(path, encoding="utf-8") as handle:
            trace = json.load(handle)
        trace["objects"][0]["unitPrice"] = "1"
        with open(path, "w", encoding="utf-8") as handle:
            json.dump(trace, handle, ensure_ascii=False, indent=2)
        refresh_receipt(plan_dir, ["计算轨迹.json"])
        result = run_validate(plan_dir, case_root, manifest_path)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_changed_current_input_is_rejected_by_validate_without_input(self):
        manifest_path, plan_dir, case_root = self._built()
        with open(manifest_path, encoding="utf-8") as handle:
            changed = json.load(handle)
        changed["referencePolicy"]["basis"] = "变化后的比较口径"
        with open(manifest_path, "w", encoding="utf-8") as handle:
            json.dump(changed, handle, ensure_ascii=False, indent=2)
        result = run_validate(plan_dir, case_root)
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_symlinked_output_escaping_outdir_is_rejected(self):
        manifest_path, plan_dir, case_root = self._built()
        outside_dir = tempfile.mkdtemp(prefix="crwu-outside-")
        self.addCleanup(shutil.rmtree, outside_dir, ignore_errors=True)
        outside = os.path.join(outside_dir, "stolen.xlsx")
        with open(outside, "wb") as handle:
            handle.write(b"outside")
        target = os.path.join(plan_dir, "估值底稿.xlsx")
        os.remove(target)
        os.symlink(outside, target)
        built = run_cli(manifest_path, plan_dir, case_root)
        self.assertNotEqual(built.returncode, 0, built.stdout + built.stderr)
        validated = run_validate(plan_dir, case_root, manifest_path)
        self.assertNotEqual(validated.returncode, 0, validated.stdout + validated.stderr)
        with open(outside, "rb") as handle:
            self.assertEqual(handle.read(), b"outside", "不得通过符号链接改写输出目录之外的文件")

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_symlinked_meta_dir_escaping_outdir_is_rejected(self):
        manifest_path, plan_dir, case_root = self._built()
        meta = os.path.join(plan_dir, ".listing-reference")
        outside_dir = tempfile.mkdtemp(prefix="crwu-meta-")
        self.addCleanup(shutil.rmtree, outside_dir, ignore_errors=True)
        shutil.copyfile(os.path.join(meta, "receipt.json"),
                        os.path.join(outside_dir, "receipt.json"))
        shutil.rmtree(meta)
        os.symlink(outside_dir, meta)
        built = run_cli(manifest_path, plan_dir, case_root)
        self.assertNotEqual(built.returncode, 0, built.stdout + built.stderr)
        validated = run_validate(plan_dir, case_root, manifest_path)
        self.assertNotEqual(validated.returncode, 0, validated.stdout + validated.stderr)

    @unittest.skipUnless(os.path.isfile(SCRIPT), "build_listing_reference.py 尚未交付")
    def test_revision_image_hyperlink_resolves_from_final_outdir(self):
        from openpyxl import load_workbook
        manifest_path, plan_dir, case_root = self._built()
        revision = os.path.join(plan_dir, "修订", "2")
        result = run_build(manifest_path, revision, case_root)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        book = load_workbook(os.path.join(revision, "估值底稿.xlsx"))
        try:
            href = book["案例证据"]["E2"].hyperlink.target
            self.assertTrue(
                os.path.isfile(os.path.normpath(os.path.join(revision, href))),
                f"修订目录的图片超链接必须解析到真实文件：{href}")
        finally:
            book.close()
        validated = run_validate(revision, case_root, manifest_path)
        self.assertEqual(validated.returncode, 0, validated.stdout + validated.stderr)


if __name__ == "__main__":
    unittest.main(verbosity=2)

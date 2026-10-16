#!/usr/bin/env python3
"""从**已重算过的**工作簿里**回读**数值（不回填期望值，也不猜）。

    python3 scripts/read_workbook_values.py --workbook <重算后的.xlsx> --cells '<JSON>'

`--cells` 是「对象 → 数值位置」的映射，位置由**模板自己的填写说明**给出：

    {"object-001": {"unitPrice": "结果!B5", "totalValue": "结果!B6"}}

stdout 是 JSON（只此一份，便于宿主解析）：

    {"observed": [{"objectId": "object-001", "unitPrice": "39800", "totalValue": "4780000"}],
     "problems": ["object-002：结果!B5 是公式但没有缓存值（文件没有被真正打开重算过）"]}

## 三条口径

1. **读的是缓存值**（`data_only=True`）。openpyxl **不计算公式**：读回来的是上一次由
   Excel / LibreOffice 保存的数值。所以"公式在、缓存空"是一个必须被看见的事实 ——
   它说明这份文件**没有被真正打开重算过**，而不是"这一项为空"。
2. **读不到就是 `null`**，绝不回填期望值。把期望值当观测值会让核对永远是绿的，
   而"重算通过"就变成了同义反复。
3. 数值一律输出**十进制字符串**（与 `估值报告数据.json` 同一口径），保留工作簿里的
   有效数字，不在这里二次取整。
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


def fail(message: str) -> None:
    print(f"FAIL {message}")
    raise SystemExit(1)


def split_ref(reference: str) -> tuple[str, str]:
    raw = str(reference).strip()
    if "!" in raw:
        sheet, _, cell = raw.rpartition("!")
        return sheet.strip("'"), cell.strip()
    return "", raw


def decimal_text(value) -> tuple[str | None, str]:
    """把单元格值转成十进制字符串；转不出来就返回 (None, 原因)。"""
    if value is None:
        return None, ""
    if isinstance(value, bool):
        return None, "值是布尔：数值位置不接受布尔"
    if isinstance(value, (int, float)):
        try:
            # 经 `str` 再进 `Decimal`：避免把二进制浮点误差带成 39800.000000001。
            text = str(Decimal(str(value)).normalize())
        except (InvalidOperation, ValueError):
            return None, f"数值无法读成十进制：{value!r}"
        # `normalize()` 会把 39800 变成 3.98E+4：换回定点写法。
        if "E" in text or "e" in text:
            text = format(Decimal(text), "f")
        return text, ""
    raw = str(value).strip()
    if raw == "":
        return None, ""
    try:
        text = format(Decimal(raw), "f")
    except (InvalidOperation, ValueError):
        return None, f"单元格里不是数值：{raw[:40]!r}"
    return text, ""


def read_values(workbook_path: str, cells: dict) -> tuple[list[dict], list[str]]:
    problems: list[str] = []
    observed: list[dict] = []
    # `data_only=True` 读缓存值；`data_only=False` 用来判断"这里本来是公式吗"。
    book = load_workbook(workbook_path, data_only=True)
    raw_book = load_workbook(workbook_path, data_only=False)
    try:
        for object_id, positions in cells.items():
            if not isinstance(positions, dict):
                problems.append(f"{object_id}：位置映射必须是对象（形如 {{\"unitPrice\": \"结果!B5\"}}）")
                continue
            row = {"objectId": str(object_id), "unitPrice": None, "totalValue": None}
            for key in ("unitPrice", "totalValue"):
                reference = str(positions.get(key, "") or "").strip()
                if reference == "":
                    continue
                sheet, cell = split_ref(reference)
                if sheet == "" or cell == "":
                    problems.append(f"{object_id}：{key} 的位置 {reference!r} 必须写成 `表名!单元格`")
                    continue
                if sheet not in book.sheetnames:
                    problems.append(f"{object_id}：{key} 指向不存在的表 {sheet!r}")
                    continue
                cached = book[sheet][cell].value
                if cached is None:
                    source = raw_book[sheet][cell].value
                    if isinstance(source, str) and source.startswith("="):
                        problems.append(
                            f"{object_id}：{sheet}!{cell} 是公式（{source[:40]}）但没有缓存值 —— "
                            '这份工作簿没有被 Excel / LibreOffice 真正打开重算过'
                            '（不是「这一项为空」）')
                    else:
                        problems.append(f"{object_id}：{sheet}!{cell} 是空的（既没有值也没有公式）")
                    continue
                text, why = decimal_text(cached)
                if why != "":
                    problems.append(f"{object_id}：{sheet}!{cell} {why}")
                    continue
                row[key] = text
            observed.append(row)
    finally:
        book.close()
        raw_book.close()
    return observed, problems


def main() -> int:
    parser = argparse.ArgumentParser(description="从已重算过的工作簿里回读数值")
    parser.add_argument("--workbook", required=True, help="重算后的工作簿路径")
    parser.add_argument("--cells", required=True, help="对象 → 数值位置的 JSON（字符串或文件路径）")
    args = parser.parse_args()

    if not os.path.isfile(args.workbook):
        fail(f"找不到工作簿：{args.workbook}")
    raw = args.cells
    if os.path.isfile(raw):
        with open(raw, "r", encoding="utf-8") as handle:
            raw = handle.read()
    try:
        cells = json.loads(raw)
    except json.JSONDecodeError as error:
        fail(f"--cells 不是合法 JSON：{error}")
    if not isinstance(cells, dict) or not cells:
        fail("--cells 必须是「对象 → 位置」的非空对象")

    observed, problems = read_values(args.workbook, cells)
    print(json.dumps({"observed": observed, "problems": problems}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

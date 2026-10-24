"""Build and re-read template-driven workpapers without claiming recalculation."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill


def inside(root: Path, relative: str) -> Path:
    if not isinstance(relative, str) or not relative or re.search(r"[:\x00-\x1f]", relative):
        raise ValueError("invalid project-relative path")
    parts = relative.replace("\\", "/").split("/")
    if any(p in ("", ".", "..") for p in parts):
        raise ValueError("invalid project-relative path")
    target = root.joinpath(*parts).resolve()
    if not target.is_relative_to(root.resolve()) or target == root.resolve():
        raise ValueError("path leaves project")
    return target


def digest(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def artifact(root: Path, path: Path, kind: str) -> dict:
    return {"relativePath": path.relative_to(root).as_posix(), "digest": digest(path), "bytes": path.stat().st_size, "kind": kind}


def atomic_save(path: Path, save):
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp = tempfile.mkstemp(prefix=".market-", suffix=path.suffix, dir=path.parent)
    os.close(handle)
    try:
        save(Path(temp))
        os.replace(temp, path)
    finally:
        Path(temp).unlink(missing_ok=True)


def literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def decorate(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="EAF1FB")
        cell.font = Font(bold=True, color="26354A")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(54, max(16, len(str(column[0].value or "")) * 2))


COLUMNS = [("案例ID", None), ("网页地址", None), ("来源", None), ("列表建筑面积（㎡）", "areaSqm"),
    ("列表单价（元/㎡）", "unitPrice"), ("列表总价（元）", "totalPrice"), ("价格性质", "priceBasis"),
    ("发布时间", "publishedAt"), ("采集时间", "capturedAt"), ("详情建筑面积（㎡）", "areaSqm"),
    ("详情单价（元/㎡）", "unitPrice"), ("初筛状态", None), ("选择/采用状态", None), ("相似点", None),
    ("差异", None), ("待核字段", None), ("推荐理由", None), ("证据与截图", None)]


def candidate_rows(candidates):
    for c in candidates:
        listed, details = c.get("listFacts", {}), c.get("detailFacts", {})
        values = [c["caseId"], c["sourceUrl"], c["providerId"]]
        for _, key in COLUMNS[3:9]:
            values.append(listed.get(key, {}).get("value"))
        values.extend([details.get("areaSqm", {}).get("value"), details.get("unitPrice", {}).get("value"),
            c["status"], c["selectionState"], "；".join(c["similarities"]), "；".join(c["differences"]),
            "；".join(c["missingFields"]), c["recommendation"], "；".join(f'{r["location"]} @ {r["digest"]}' for r in c["sourceRefs"])])
        yield values


def candidate_book(candidates):
    book = Workbook(); sheet = book.active; sheet.title = "候选案例"
    sheet.append([label for label, _ in COLUMNS])
    for row in candidate_rows(candidates):
        sheet.append(row)
        for cell, value in zip(sheet[sheet.max_row], row):
            literal(cell, value)
        if not re.match(r"^https?://[^\s]+$", row[1]):
            raise ValueError("candidate URL is invalid")
        sheet.cell(sheet.max_row, 2).hyperlink = row[1]
        sheet.cell(sheet.max_row, 2).style = "Hyperlink"
    decorate(sheet)
    return book


def verify_candidates(path, candidates):
    book = load_workbook(path); sheet = book["候选案例"]
    expected = list(candidate_rows(candidates))
    if sheet.max_row != len(expected) + 1:
        raise ValueError("candidate row count mismatch")
    for index, row in enumerate(expected, 2):
        actual = [c.value for c in sheet[index]]
        if actual != [None if v == "" else v for v in row] or sheet.cell(index, 2).hyperlink.target != row[1]:
            raise ValueError("candidate values or hyperlink mismatch")
    book.close()


def check_package(path):
    if path.suffix.lower() != ".xlsx":
        raise ValueError("template must be xlsx")
    with zipfile.ZipFile(path) as package:
        members = package.infolist()
        if len(members) > 10000 or sum(m.file_size for m in members) > 128 * 1024 * 1024:
            raise ValueError("template package exceeds limits")
        if any(m.filename.startswith("/") or ".." in m.filename.split("/") for m in members):
            raise ValueError("unsafe template package")


def verify(root, payload):
    for expected in payload["artifacts"]:
        path = inside(root, expected["relativePath"])
        if digest(path) != expected["digest"] or path.stat().st_size != expected["bytes"]:
            raise ValueError("artifact digest or size mismatch")
        if path.suffix == ".xlsx":
            check_package(path)
            book = load_workbook(path)
            if not book.sheetnames:
                raise ValueError("empty workbook")
            book.close()
    if payload.get("template"):
        original = inside(root, payload["template"]["relativePath"])
        if digest(original) != payload["template"]["digest"]:
            raise ValueError("original template identity changed")
    workbook = next((a for a in payload["artifacts"] if a["kind"] == "workpaper"), None)
    if workbook:
        adopted = payload.get("adoptedCaseIds", [])
        if len(set(adopted)) != len(adopted) or any(e.get("caseId") and e["caseId"] not in adopted for e in payload.get("fillTrace", [])):
            raise ValueError("adopted cases do not match filled cases")
        trace_file = next((a for a in payload["artifacts"] if a["kind"] == "trace"), None)
        if not trace_file:
            raise ValueError("fill trace is missing")
        saved = json.loads(inside(root, trace_file["relativePath"]).read_text(encoding="utf-8"))
        if any(saved.get(key) != payload.get(key) for key in ("template", "adoptedCaseIds", "fillTrace")):
            raise ValueError("persisted fill trace does not match delivery")
        book = load_workbook(inside(root, workbook["relativePath"]))
        for entry in payload.get("fillTrace", []):
            if entry["fact"]["status"] in ("verified", "assumed") and entry["fact"]["value"] is not None:
                if book[entry["sheet"]][entry["cell"]].value != entry["fact"]["value"]:
                    raise ValueError("filled cell mismatch")
        book.close()
    candidates = next((a for a in payload["artifacts"] if a["kind"] == "candidates"), None)
    if candidates:
        verify_candidates(inside(root, candidates["relativePath"]), payload.get("candidates", []))
    return {"verifiedArtifacts": payload["artifacts"]}


def build(root: Path, operation: str, payload: dict) -> dict:
    root = root.resolve()
    if not root.is_dir():
        raise ValueError("case root is missing")
    if operation == "verify":
        return verify(root, payload)
    directory = payload.get("outputDir", "成果/市场法底稿")
    if operation == "candidates":
        path = inside(root, directory + "/候选案例表.xlsx")
        book = candidate_book(payload["candidates"])
        atomic_save(path, lambda target: book.save(target)); book.close()
        verify_candidates(path, payload["candidates"])
        return {"artifact": artifact(root, path, "candidates")}
    if operation != "fill":
        raise ValueError("unknown operation")
    original = inside(root, payload["template"]["relativePath"])
    if digest(original) != payload["template"]["digest"]:
        raise ValueError("template digest mismatch")
    check_package(original)
    book = load_workbook(original)
    for ref in payload.get("clearRefs", []):
        cell = book[ref["sheet"]][ref["cell"]]
        if not isinstance(cell, MergedCell) and cell.data_type != "f":
            cell.value = None
    filled = []
    seen = set()
    for entry in payload["fillTrace"]:
        target = (entry["sheet"], entry["cell"])
        if target in seen or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", entry["cell"]):
            raise ValueError("duplicate or invalid cell mapping")
        seen.add(target)
        if entry["sheet"] not in book.sheetnames:
            raise ValueError("mapped worksheet missing")
        cell = book[entry["sheet"]][entry["cell"]]
        if isinstance(cell, MergedCell) or cell.data_type == "f":
            raise ValueError("mapped target is a formula or merged follower")
        fact = entry["fact"]
        if fact["status"] in ("verified", "assumed") and fact["value"] is not None:
            literal(cell, fact["value"])
            filled.append(entry)
    title = "本轮填表说明"
    if title in book.sheetnames:
        title = "本轮填表说明_新增"
    sheet = book.create_sheet(title)
    sheet.append(["字段", "案例ID", "单元格", "取值", "单位", "状态", "出处"])
    for e in filled:
        row = [e["fieldId"], e.get("caseId"), e["sheet"] + "!" + e["cell"], e["fact"]["value"], e["fact"]["unit"], e["fact"]["status"], json.dumps(e["fact"]["evidenceRefs"], ensure_ascii=False)]
        sheet.append(row)
        for cell, value in zip(sheet[sheet.max_row], row):
            literal(cell, value)
    sheet.append(["限制", "未执行原模板重算；本成果为已填底稿，公式缓存不作为估值结论"])
    for note in payload.get("limitations", []):
        sheet.append(["待核", note])
        literal(sheet.cell(sheet.max_row, 2), note)
    decorate(sheet)
    path = inside(root, directory + "/原模板回填底稿.xlsx")
    atomic_save(path, lambda target: book.save(target)); book.close()
    if digest(original) != payload["template"]["digest"]:
        raise ValueError("original template changed")
    files = [artifact(root, path, "workpaper")]
    trace = {"template": payload["template"], "adoptedCaseIds": payload["adoptedCaseIds"], "fillTrace": filled,
        "limitations": payload.get("limitations", []), "recalculation": "not-run", "calculation": None, "artifacts": files}
    trace_path = inside(root, directory + "/填表轨迹.json")
    atomic_save(trace_path, lambda target: target.write_text(json.dumps(trace, ensure_ascii=False, allow_nan=False, indent=2), encoding="utf-8"))
    files.append(artifact(root, trace_path, "trace"))
    verify(root, {**trace, "artifacts": files})
    return {"artifacts": files, "fillTrace": filled, "adoptedCaseIds": payload["adoptedCaseIds"], "calculation": None, "recalculation": "not-run"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-root", required=True)
    parser.add_argument("--input", required=True)
    parser.add_argument("--operation", choices=["candidates", "fill", "verify"], required=True)
    args = parser.parse_args()
    try:
        root = Path(args.case_root).resolve()
        request = inside(root, args.input)
        if request.stat().st_size > 32 * 1024 * 1024:
            raise ValueError("request too large")
        data = json.loads(request.read_text(encoding="utf-8"))
        result = build(root, args.operation, data)
        print(json.dumps({"ok": True, "data": result}, ensure_ascii=False, allow_nan=False))
    except Exception:
        print(json.dumps({"ok": False, "errorKind": "file", "error": "workpaper verification failed; check project paths, template identity and cell mappings"}))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

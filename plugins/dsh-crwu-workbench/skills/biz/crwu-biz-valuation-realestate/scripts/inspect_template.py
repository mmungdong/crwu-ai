"""Read-only inspection of a case template workbook.

The inspector reports the actual workbook structure, labels, formulas and
defined names. It never saves the workbook, recalculates formulas or repairs a
package. Every failure is returned as a structured result without local paths.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import zipfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

MAX_COMPRESSED_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_ZIP_MEMBERS = 10_000
MAX_SHEETS = 30
MAX_VISITED_CELLS = 20_000
MAX_ROWS = 4_000
MAX_COLUMNS = 200
MAX_DEFINED_NAMES = 1_000
MAX_MERGED_RANGES = 1_000
DIGEST_PATTERN = re.compile(r"^sha256:[0-9a-fA-F]{64}$")
RANGE_PATTERN = re.compile(r"^[A-Z]{1,3}[1-9][0-9]*(?::[A-Z]{1,3}[1-9][0-9]*)?$")


class InspectionFailure(Exception):
    """A failure that can be reported without a traceback or local path."""

    def __init__(self, kind: str, message: str) -> None:
        super().__init__(message)
        self.kind = kind
        self.message = message


def _failure(kind: str, message: str) -> dict:
    return {"ok": False, "errorKind": kind, "error": message}


def _digest(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            hasher.update(chunk)
    return "sha256:" + hasher.hexdigest()


def _resolve_inside(case_root: str, workbook: str) -> Path:
    try:
        root = Path(case_root).expanduser().resolve()
        raw = Path(workbook).expanduser()
        target = raw.resolve() if raw.is_absolute() else (root / raw).resolve()
    except (OSError, RuntimeError, ValueError):
        raise InspectionFailure("invalid-path", "template path could not be resolved") from None
    if not root.is_dir():
        raise InspectionFailure("invalid-case-root", "case root is not a directory")
    try:
        target.relative_to(root)
    except ValueError:
        raise InspectionFailure("outside-case-root", "template is outside the case root") from None
    if target == root:
        raise InspectionFailure("outside-case-root", "template is outside the case root")
    return target


def _check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as package:
            members = package.infolist()
    except (zipfile.BadZipFile, OSError):
        raise InspectionFailure("invalid-workbook", "workbook is not a readable xlsx package") from None
    if not members:
        raise InspectionFailure("invalid-workbook", "workbook package is empty")
    if len(members) > MAX_ZIP_MEMBERS:
        raise InspectionFailure("workbook-too-large", "workbook package has too many members")
    compressed = 0
    expanded = 0
    for member in members:
        name = member.filename.replace("\\", "/")
        if name.startswith("/") or ".." in name.split("/"):
            raise InspectionFailure("invalid-workbook", "workbook package contains an unsafe member")
        compressed += member.compress_size
        expanded += member.file_size
        if compressed > MAX_COMPRESSED_BYTES or expanded > MAX_EXPANDED_BYTES:
            raise InspectionFailure("workbook-too-large", "workbook package exceeds the inspection limits")


def _json_value(value):
    if value is None or isinstance(value, (str, bool)):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, Decimal):
        return str(value)
    text = getattr(value, "text", None)
    if isinstance(text, str):
        return text
    return str(value)


def _formula_text(value) -> str:
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", None)
    if isinstance(text, str) and text.startswith("="):
        return text
    return ""


def _cell_ref_in_bounds(cell: str, rows: int, columns: int) -> bool:
    matched = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]*)", cell)
    if matched is None:
        return False
    column = 0
    for letter in matched.group(1):
        column = column * 26 + (ord(letter) - ord("A") + 1)
    return int(matched.group(2)) <= rows and column <= columns


def _inspect(path: Path, template_id: str, digest: str, max_cells: int) -> dict:
    issues: list[str] = []
    seen_cells: set[tuple[str, str]] = set()

    try:
        book = load_workbook(filename=path, read_only=False, data_only=False, keep_links=False)
    except Exception:
        raise InspectionFailure("invalid-workbook", "workbook is not a readable xlsx package") from None

    try:
        sheet_names = list(book.sheetnames)
        limited_names = sheet_names[:MAX_SHEETS]
        if len(sheet_names) > MAX_SHEETS:
            issues.append("sheet limit reached")

        sheets = []
        cells = []
        visited = 0
        cell_limit_reached = False
        row_limit_reached = False
        column_limit_reached = False

        for name in limited_names:
            worksheet = book[name]
            rows = int(worksheet.max_row or 0)
            columns = int(worksheet.max_column or 0)
            merged = [str(item) for item in worksheet.merged_cells.ranges]
            if len(merged) > MAX_MERGED_RANGES:
                merged = merged[:MAX_MERGED_RANGES]
                issues.append("merged range limit reached")
            sheets.append({
                "name": name,
                "rows": rows,
                "columns": columns,
                "mergedRanges": sorted(merged),
            })

            if rows > MAX_ROWS:
                row_limit_reached = True
            if columns > MAX_COLUMNS:
                column_limit_reached = True
            if cell_limit_reached:
                continue
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(rows, MAX_ROWS),
                max_col=min(columns, MAX_COLUMNS),
            ):
                for cell in row:
                    if visited >= max_cells:
                        cell_limit_reached = True
                        break
                    visited += 1
                    value = cell.value
                    formula = _formula_text(value)
                    if value is None:
                        continue
                    if isinstance(value, str) and value == "":
                        continue
                    coordinate = cell.coordinate
                    if _cell_ref_in_bounds(coordinate, rows, columns):
                        key = (name, coordinate)
                        if key not in seen_cells:
                            seen_cells.add(key)
                            cells.append({
                                "sheet": name,
                                "cell": coordinate,
                                "value": _json_value(value),
                                "formula": formula,
                            })
                if cell_limit_reached:
                    break

        defined_names = []
        for defined_name, definition in book.defined_names.items():
            reference = getattr(definition, "attr_text", None) or getattr(definition, "value", "")
            defined_names.append({"name": str(defined_name), "ref": str(reference)})
        defined_names.sort(key=lambda item: item["name"])
        if len(defined_names) > MAX_DEFINED_NAMES:
            defined_names = defined_names[:MAX_DEFINED_NAMES]
            issues.append("defined name limit reached")
    finally:
        book.close()

    if cell_limit_reached:
        issues.append("cell limit reached")
    if row_limit_reached:
        issues.append("row limit reached")
    if column_limit_reached:
        issues.append("column limit reached")

    return {
        "ok": True,
        "profile": {
            "templateId": template_id,
            "digest": digest,
            "sheets": sheets,
            "cells": cells,
            "definedNames": defined_names,
            "coverage": {"complete": not issues, "issues": issues},
        },
    }


def inspect_template(
    path: str,
    template_id: str,
    expected_digest: str,
    *,
    max_cells: int = MAX_VISITED_CELLS,
    case_root: str | None = None,
) -> dict:
    """Inspect a workbook without writing to it and return a bounded profile."""
    try:
        if not isinstance(path, str) or not path:
            raise InspectionFailure("invalid-path", "template path is required")
        if not isinstance(template_id, str) or not template_id.strip():
            raise InspectionFailure("invalid-argument", "template identifier is required")
        if not isinstance(expected_digest, str) or DIGEST_PATTERN.fullmatch(expected_digest) is None:
            raise InspectionFailure("invalid-digest", "expected digest must be sha256:<hex>")
        if not isinstance(max_cells, int) or isinstance(max_cells, bool) or max_cells < 1:
            raise InspectionFailure("invalid-argument", "cell limit must be a positive integer")

        target = _resolve_inside(case_root, path) if case_root is not None else Path(path).expanduser()
        if not target.is_file():
            raise InspectionFailure("missing-package", "template workbook is missing")

        digest = _digest(target)
        if digest.lower() != expected_digest.lower():
            raise InspectionFailure("digest-mismatch", "template digest does not match the expected digest")

        _check_zip(target)
        result = _inspect(target, template_id, digest, max_cells)

        after = _digest(target)
        if after.lower() != digest.lower():
            raise InspectionFailure("digest-mismatch", "template bytes changed during inspection")
        return result
    except InspectionFailure as failure:
        return _failure(failure.kind, failure.message)
    except (OSError, PermissionError):
        return _failure("read-failed", "template workbook could not be read")
    except Exception:
        return _failure("inspection-failed", "template workbook could not be inspected")


def _positive_int(value: str) -> int:
    parsed = int(value)
    if parsed < 1:
        raise argparse.ArgumentTypeError("must be a positive integer")
    return parsed


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Inspect a case template workbook read-only.")
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--template-id", required=True)
    parser.add_argument("--expected-sha256", required=True)
    parser.add_argument("--case-root", required=True)
    parser.add_argument("--max-cells", type=_positive_int, default=MAX_VISITED_CELLS)
    args = parser.parse_args(argv)
    result = inspect_template(
        args.workbook,
        args.template_id,
        args.expected_sha256,
        max_cells=args.max_cells,
        case_root=args.case_root,
    )
    json.dump(result, sys.stdout, ensure_ascii=False, sort_keys=True)
    sys.stdout.write("\n")
    return 0 if result.get("ok") is True else 1


if __name__ == "__main__":
    raise SystemExit(main())
